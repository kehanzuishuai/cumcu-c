from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
ATTACHMENT1 = ROOT / "source_materials" / "C题" / "附件" / "附件1.xlsx"
ATTACHMENT2 = ROOT / "source_materials" / "C题" / "附件" / "附件2.xlsx"
DETAIL = ROOT / "results" / "q3_detail.csv"
DAILY = ROOT / "results" / "q3_daily.csv"
STORAGE = ROOT / "results" / "q3_storage_4h.csv"
EVENTS = ROOT / "results" / "q3_emergency_events.csv"
VERSIONS = ROOT / "results" / "q3_plan_versions.csv"
SUMMARY = ROOT / "results" / "q3_summary.json"
OUTPUT = ROOT / "validation" / "q3_independent_recompute.json"
TOL = 1e-6
POS_TOL = 1e-7
FLOW_MAX_KWH = 5000.0 / 6.0


def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value)).date()


def load_sources() -> tuple[list[date], np.ndarray, np.ndarray, np.ndarray]:
    wb1 = load_workbook(ATTACHMENT1, read_only=True, data_only=True)
    try:
        price = np.asarray([float(row[1]) for row in wb1.worksheets[0].iter_rows(min_row=2, values_only=True)])
    finally:
        wb1.close()
    wb2 = load_workbook(ATTACHMENT2, read_only=True, data_only=True)
    try:
        matrices = []
        dates = None
        for ws in wb2.worksheets:
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            current = [parse_date(row[0]) for row in rows]
            if dates is None:
                dates = current
            elif dates != current:
                raise ValueError("附件2日期不一致")
            matrices.append(np.asarray([[float(value) for value in row[1:]] for row in rows]))
    finally:
        wb2.close()
    assert dates is not None
    return dates, matrices[0], matrices[1], price


def fmt_minute(total: int) -> str:
    if total == 1440:
        return "24:00"
    return f"{total // 60:02d}:{total % 60:02d}"


def rebuild_events(detail: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for day, group in detail.groupby("date", sort=True):
        values = group.sort_values("slot")["emergency_purchase_kwh"].to_numpy(float)
        start = None
        event_id = 0
        for slot, value in enumerate(values):
            if value > POS_TOL and start is None:
                start = slot
            if start is not None and (value <= POS_TOL or slot == 143):
                end = slot if value <= POS_TOL else slot + 1
                event_id += 1
                rows.append({"date": day, "event_id": event_id,
                             "interval": f"{fmt_minute(start * 10)}-{fmt_minute(end * 10)}",
                             "purchase_kwh": float(np.sum(values[start:end]))})
                start = None
    return pd.DataFrame(rows, columns=["date", "event_id", "interval", "purchase_kwh"])


def main() -> None:
    source_dates, source_load, source_pv, price = load_sources()
    detail = pd.read_csv(DETAIL, encoding="utf-8-sig").sort_values(["date", "slot"]).reset_index(drop=True)
    daily = pd.read_csv(DAILY, encoding="utf-8-sig").sort_values("date").reset_index(drop=True)
    storage = pd.read_csv(STORAGE, encoding="utf-8-sig").sort_values(["date", "block"]).reset_index(drop=True)
    events = pd.read_csv(EVENTS, encoding="utf-8-sig")
    versions = pd.read_csv(VERSIONS, encoding="utf-8-sig")
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))

    expected_load = source_load[31:].reshape(-1)
    expected_pv = source_pv[31:].reshape(-1)
    source_difference = float(max(
        np.max(np.abs(detail["actual_load_kw"].to_numpy() - expected_load)),
        np.max(np.abs(detail["actual_pv_kw"].to_numpy() - expected_pv)),
        np.max(np.abs(detail["load_kwh"].to_numpy() - expected_load / 6.0)),
        np.max(np.abs(detail["pv_kwh"].to_numpy() - expected_pv / 6.0)),
    ))
    q0 = detail["original_purchase_kwh"].to_numpy(float)
    effective = detail["effective_purchase_kwh"].to_numpy(float)
    upward = np.maximum(effective - q0, 0.0)
    downward = np.maximum(q0 - effective, 0.0)
    tiled_price = np.tile(price, 334)
    base_cost = tiled_price * q0
    down_credit = 0.5 * tiled_price * downward
    up_cost = 1.5 * tiled_price * upward
    normal_cost = base_cost - down_credit + up_cost
    emergency_cost = 5.0 * tiled_price * detail["emergency_purchase_kwh"].to_numpy(float)
    total_cost = normal_cost + emergency_cost
    adjustment_difference = float(max(
        np.max(np.abs(upward - detail["upward_adjustment_kwh"].to_numpy(float))),
        np.max(np.abs(downward - detail["downward_adjustment_kwh"].to_numpy(float))),
        np.max(np.abs(base_cost - detail["base_plan_cost_yuan"].to_numpy(float))),
        np.max(np.abs(down_credit - detail["downward_credit_yuan"].to_numpy(float))),
        np.max(np.abs(up_cost - detail["upward_adjustment_cost_yuan"].to_numpy(float))),
        np.max(np.abs(normal_cost - detail["normal_settlement_cost_yuan"].to_numpy(float))),
        np.max(np.abs(emergency_cost - detail["emergency_cost_yuan"].to_numpy(float))),
        np.max(np.abs(total_cost - detail["total_cost_yuan"].to_numpy(float))),
    ))
    balance = effective + detail["pv_kwh"] + detail["discharge_bus_kwh"] + detail["emergency_purchase_kwh"] - detail["load_kwh"] - detail["charge_bus_kwh"] - detail["surplus_discard_kwh"]
    state = detail["soc_close_kwh"] - detail["soc_open_kwh"] - 0.9 * detail["charge_bus_kwh"] + detail["discharge_bus_kwh"] / 0.9

    recomputed_daily = detail.assign(
        _base=base_cost, _down=down_credit, _up=up_cost, _normal=normal_cost,
        _emergency_cost=emergency_cost, _total=total_cost,
    ).groupby("date", sort=True).agg(
        original_purchase_kwh=("original_purchase_kwh", "sum"),
        effective_purchase_kwh=("effective_purchase_kwh", "sum"),
        upward_adjustment_kwh=("upward_adjustment_kwh", "sum"),
        downward_adjustment_kwh=("downward_adjustment_kwh", "sum"),
        emergency_purchase_kwh=("emergency_purchase_kwh", "sum"),
        base_plan_cost_yuan=("_base", "sum"), downward_credit_yuan=("_down", "sum"),
        upward_adjustment_cost_yuan=("_up", "sum"), normal_settlement_cost_yuan=("_normal", "sum"),
        emergency_cost_yuan=("_emergency_cost", "sum"), total_cost_yuan=("_total", "sum"),
        charge_kwh=("charge_bus_kwh", "sum"), discharge_kwh=("discharge_bus_kwh", "sum"),
        surplus_discard_kwh=("surplus_discard_kwh", "sum"),
        soc_open_kwh=("soc_open_kwh", "first"), soc_close_kwh=("soc_close_kwh", "last"),
    ).reset_index()
    daily_numeric = [column for column in recomputed_daily.columns if column != "date"]
    max_daily_difference = float(max(np.max(np.abs(recomputed_daily[column].to_numpy() - daily[column].to_numpy())) for column in daily_numeric))

    temp = detail.assign(block=(detail["slot"].to_numpy(int) - 1) // 24 + 1)
    rebuilt_storage = temp.groupby(["date", "block"], sort=True).agg(
        charge_kwh=("charge_bus_kwh", "sum"), discharge_kwh=("discharge_bus_kwh", "sum")
    ).reset_index()
    max_storage_difference = float(max(
        np.max(np.abs(rebuilt_storage[column].to_numpy() - storage[column].to_numpy()))
        for column in ("charge_kwh", "discharge_kwh")
    ))

    rebuilt_events = rebuild_events(detail)
    events_sorted = events[["date", "event_id", "interval", "purchase_kwh"]].sort_values(["date", "event_id"]).reset_index(drop=True)
    event_keys_exact = rebuilt_events[["date", "event_id", "interval"]].equals(events_sorted[["date", "event_id", "interval"]])
    max_event_difference = float(np.max(np.abs(rebuilt_events["purchase_kwh"].to_numpy() - events_sorted["purchase_kwh"].to_numpy()))) if len(events_sorted) else 0.0

    allowed_issue_slots = {0, 36, 72, 108}
    version_issue_valid = bool(versions["issue_slot"].isin(allowed_issue_slots).all())
    past_lock_violations = int(np.sum((versions["issue_slot"] > 0) & (versions["target_slot"] <= versions["issue_slot"])))
    latest = versions.sort_values(["date", "target_slot", "issue_slot"]).groupby(["date", "target_slot"], as_index=False).tail(1).sort_values(["date", "target_slot"]).reset_index(drop=True)
    version_effective_difference = float(np.max(np.abs(latest["new_effective_plan_kwh"].to_numpy() - effective)))
    original_versions = versions[versions["issue_slot"] == 0].sort_values(["date", "target_slot"]).reset_index(drop=True)
    version_original_difference = float(np.max(np.abs(original_versions["original_plan_kwh"].to_numpy() - q0)))

    crossday_gap = float(np.max(np.abs(daily["soc_open_kwh"].to_numpy()[1:] - daily["soc_close_kwh"].to_numpy()[:-1])))
    simultaneous = int(np.sum((detail["charge_bus_kwh"] > POS_TOL) & (detail["discharge_bus_kwh"] > POS_TOL)))
    emergency_charge = int(np.sum((detail["emergency_purchase_kwh"] > POS_TOL) & (detail["charge_bus_kwh"] > POS_TOL)))
    max_charge = float(detail["charge_bus_kwh"].max())
    max_discharge = float(detail["discharge_bus_kwh"].max())
    information_violations = int(np.sum(detail["train_end_date"] >= detail["date"]))
    expected_dates = [value.isoformat() for value in source_dates[31:]]
    date_axis_exact = sorted(detail["date"].unique().tolist()) == expected_dates
    reported = summary["strategies"]["q3_all_main"]
    summary_differences = {
        "original_purchase_kwh": abs(float(recomputed_daily["original_purchase_kwh"].sum()) - reported["original_purchase_kwh"]),
        "effective_purchase_kwh": abs(float(recomputed_daily["effective_purchase_kwh"].sum()) - reported["effective_purchase_kwh"]),
        "emergency_purchase_kwh": abs(float(recomputed_daily["emergency_purchase_kwh"].sum()) - reported["emergency_purchase_kwh"]),
        "total_cost_yuan": abs(float(recomputed_daily["total_cost_yuan"].sum()) - reported["total_cost_yuan"]),
        "final_soc_kwh": abs(float(recomputed_daily.iloc[-1]["soc_close_kwh"]) - reported["soc_final_kwh"]),
    }
    metrics = {
        "detail_rows": int(len(detail)), "daily_rows": int(len(daily)), "storage_rows": int(len(storage)),
        "event_rows": int(len(events_sorted)), "version_rows": int(len(versions)),
        "date_axis_exact": date_axis_exact, "max_source_difference": source_difference,
        "max_adjustment_and_cost_slot_difference": adjustment_difference,
        "max_balance_residual_kwh": float(np.max(np.abs(balance))),
        "max_state_residual_kwh": float(np.max(np.abs(state))),
        "soc_min_kwh": float(min(detail["soc_open_kwh"].min(), detail["soc_close_kwh"].min())),
        "soc_max_kwh": float(max(detail["soc_open_kwh"].max(), detail["soc_close_kwh"].max())),
        "max_crossday_soc_gap_kwh": crossday_gap, "max_daily_recompute_difference": max_daily_difference,
        "max_storage_recompute_difference_kwh": max_storage_difference, "event_keys_exact": bool(event_keys_exact),
        "max_event_amount_difference_kwh": max_event_difference, "version_issue_valid": version_issue_valid,
        "past_lock_violations": past_lock_violations, "max_version_effective_difference_kwh": version_effective_difference,
        "max_version_original_difference_kwh": version_original_difference, "simultaneous_actual_slots": simultaneous,
        "emergency_charge_overlap_slots": emergency_charge, "information_cutoff_violations": information_violations,
        "storage_power_limit_kw": 5000.0, "storage_flow_limit_kwh_per_slot": FLOW_MAX_KWH,
        "max_charge_kwh_per_slot": max_charge, "max_discharge_kwh_per_slot": max_discharge,
        "summary_differences": summary_differences,
    }
    pass_flags = {
        "row_counts": len(detail) == 334 * 144 and len(daily) == 334 and len(storage) == 334 * 6,
        "date_axis": date_axis_exact, "source": source_difference <= 1e-9,
        "adjustment_cost": adjustment_difference <= TOL, "balance": metrics["max_balance_residual_kwh"] <= TOL,
        "state": metrics["max_state_residual_kwh"] <= TOL,
        "storage_power": max_charge <= FLOW_MAX_KWH + TOL and max_discharge <= FLOW_MAX_KWH + TOL,
        "soc_bounds": metrics["soc_min_kwh"] >= 1200 - TOL and metrics["soc_max_kwh"] <= 10800 + TOL,
        "crossday": crossday_gap <= TOL, "daily": max_daily_difference <= TOL,
        "storage": max_storage_difference <= TOL, "events": event_keys_exact and max_event_difference <= TOL,
        "versions": version_issue_valid and past_lock_violations == 0 and version_effective_difference <= TOL and version_original_difference <= TOL,
        "mutual_exclusion": simultaneous == 0, "emergency_not_charging": emergency_charge == 0,
        "information_cutoff": information_violations == 0, "summary": max(summary_differences.values()) <= TOL,
    }
    payload = {"schema_version": 1, "question": "Q3", "status": "PASS" if all(pass_flags.values()) else "FAIL",
               "pass_flags": pass_flags, "metrics": metrics,
               "independence": "Reads attachments 1-2 and Q3 frozen CSV/JSON outputs directly; does not import common_data, dispatch_core, or q3_solve."}
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": payload["status"], "metrics": metrics}, ensure_ascii=False))
    if payload["status"] != "PASS":
        raise RuntimeError("Q3独立复算失败")


if __name__ == "__main__":
    main()
