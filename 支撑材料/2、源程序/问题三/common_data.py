from __future__ import annotations

import hashlib
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ATTACHMENT_DIR = PROJECT_ROOT / "source_materials" / "C题" / "附件"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace(".", "-").replace("/", "-")
    return datetime.strptime(text, "%Y-%m-%d").date()


def minute_label(value: Any) -> int:
    if isinstance(value, datetime):
        return value.hour * 60 + value.minute
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    text = str(value).strip()
    plus_day = "+1" in text
    text = text.replace("+1", "")
    hour, minute = (int(part) for part in text.split(":")[:2])
    return hour * 60 + minute + (1440 if plus_day else 0)


def natural_slot_label(slot: int) -> tuple[str, str, str]:
    start = slot * 10
    end = (slot + 1) * 10

    def fmt(total: int) -> str:
        if total == 1440:
            return "24:00"
        return f"{total // 60:02d}:{total % 60:02d}"

    start_text = fmt(start)
    end_text = fmt(end)
    return start_text, end_text, f"{start_text}-{end_text}"


def load_attachment1() -> dict[str, Any]:
    path = ATTACHMENT_DIR / "附件1.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()
    endpoint_minutes = np.asarray([minute_label(row[0]) for row in rows], dtype=int)
    if len(rows) != 144 or endpoint_minutes.tolist() != list(range(10, 1450, 10)):
        raise ValueError("附件1时间轴不是已冻结的0:10至24:00共144个右端点")
    return {
        "path": path,
        "sha256": sha256(path),
        "raw_endpoint_labels": [str(row[0]) for row in rows],
        "endpoint_minutes": endpoint_minutes,
        "price_yuan_per_kwh": np.asarray([float(row[1]) for row in rows], dtype=float),
        "load_kw": np.asarray([float(row[2]) for row in rows], dtype=float),
        "pv_kw": np.asarray([float(row[3]) for row in rows], dtype=float),
    }


def load_attachment2() -> dict[str, Any]:
    path = ATTACHMENT_DIR / "附件2.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(wb.worksheets) != 2:
            raise ValueError("附件2应包含且仅包含负载与光伏两张工作表")
        parsed: list[dict[str, Any]] = []
        for ws in wb.worksheets:
            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            endpoint_minutes = np.asarray([minute_label(value) for value in header[1:]], dtype=int)
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            parsed.append({
                "dates": [parse_date(row[0]) for row in rows],
                "values_kw": np.asarray([[float(value) for value in row[1:]] for row in rows], dtype=float),
                "raw_endpoint_labels": [str(value) for value in header[1:]],
                "endpoint_minutes": endpoint_minutes,
            })
    finally:
        wb.close()

    load_data, pv_data = parsed
    if load_data["dates"] != pv_data["dates"]:
        raise ValueError("附件2两张表的日期轴不一致")
    if load_data["raw_endpoint_labels"] != pv_data["raw_endpoint_labels"]:
        raise ValueError("附件2两张表的时间标签不一致")
    if len(load_data["dates"]) != 365:
        raise ValueError(f"附件2日期数应为365，实际为{len(load_data['dates'])}")
    if load_data["endpoint_minutes"].tolist() != list(range(10, 1450, 10)):
        raise ValueError("附件2时间轴不是已冻结的0:10至24:00共144个右端点")
    if load_data["values_kw"].shape != (365, 144) or pv_data["values_kw"].shape != (365, 144):
        raise ValueError("附件2负载/光伏矩阵形状不是365×144")
    if not np.isfinite(load_data["values_kw"]).all() or not np.isfinite(pv_data["values_kw"]).all():
        raise ValueError("附件2包含非有限数值")
    if np.min(load_data["values_kw"]) < 0 or np.min(pv_data["values_kw"]) < 0:
        raise ValueError("附件2包含负功率")
    return {
        "path": path,
        "sha256": sha256(path),
        "dates": load_data["dates"],
        "raw_endpoint_labels": load_data["raw_endpoint_labels"],
        "endpoint_minutes": load_data["endpoint_minutes"],
        "load_kw": load_data["values_kw"],
        "pv_kw": pv_data["values_kw"],
    }


def load_attachment3() -> dict[str, Any]:
    path = ATTACHMENT_DIR / "附件3.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        rows = list(ws.iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()
    if len(header) != 26 or len(rows) != 365 * 4:
        raise ValueError("附件3应为日期、发布时间和24个提前小时，共1460条记录")

    records: list[tuple[date, int, np.ndarray]] = []
    current_date: date | None = None
    for row in rows:
        if row[0] is not None and str(row[0]).strip() != "":
            current_date = parse_date(row[0])
        if current_date is None:
            raise ValueError("附件3首条记录缺少日期")
        issue_minute = minute_label(row[1])
        values = np.asarray([float(value) for value in row[2:]], dtype=float)
        records.append((current_date, issue_minute, values))

    dates = sorted({record[0] for record in records})
    if len(dates) != 365:
        raise ValueError(f"附件3唯一日期数应为365，实际为{len(dates)}")
    date_index = {value: idx for idx, value in enumerate(dates)}
    issue_minutes = [0, 360, 720, 1080]
    issue_index = {value: idx for idx, value in enumerate(issue_minutes)}
    forecast_kw = np.full((365, 4, 24), np.nan, dtype=float)
    for record_date, issue_minute, values in records:
        if issue_minute not in issue_index:
            raise ValueError(f"附件3出现非0/6/12/18点发布时间: {issue_minute}")
        target = forecast_kw[date_index[record_date], issue_index[issue_minute]]
        if np.isfinite(target).any():
            raise ValueError(f"附件3重复键: {record_date} {issue_minute}")
        target[:] = values
    if not np.isfinite(forecast_kw).all() or np.min(forecast_kw) < 0:
        raise ValueError("附件3包含缺失、非有限或负光伏预测")
    return {
        "path": path,
        "sha256": sha256(path),
        "dates": dates,
        "issue_minutes": issue_minutes,
        "lead_hours": list(range(1, 25)),
        "forecast_kw": forecast_kw,
        "headers": [str(value) for value in header],
    }


def load_attachment4() -> dict[str, Any]:
    path = ATTACHMENT_DIR / "附件4.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        rows = list(ws.iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()
    dates = [parse_date(row[0]) for row in rows]
    endpoint_minutes = np.asarray([minute_label(value) for value in header[1:]], dtype=int)
    price = np.asarray([[float(value) for value in row[1:]] for row in rows], dtype=float)
    if len(dates) != 365 or len(set(dates)) != 365:
        raise ValueError("附件4日期轴不是365个唯一自然日")
    if endpoint_minutes.tolist() != list(range(10, 1450, 10)):
        raise ValueError("附件4时间轴不是已冻结的0:10至24:00共144个右端点")
    if price.shape != (365, 144) or not np.isfinite(price).all() or np.min(price) < 0:
        raise ValueError("附件4价格矩阵不是365×144非负有限数")
    return {
        "path": path,
        "sha256": sha256(path),
        "dates": dates,
        "raw_endpoint_labels": [str(value) for value in header[1:]],
        "endpoint_minutes": endpoint_minutes,
        "price_yuan_per_kwh": price,
    }
