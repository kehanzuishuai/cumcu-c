from __future__ import annotations

import json
import os
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(os.environ.get("CUMCM_PROJECT_ROOT", Path.cwd())).resolve()
ATT = ROOT / "source_materials" / "C题" / "附件"
TOL = 1e-6
FLOW_MAX_KWH = 5000.0 / 6.0


def load_source() -> tuple[list[str], np.ndarray, np.ndarray, np.ndarray]:
    wb = load_workbook(ATT / "附件2.xlsx", read_only=True, data_only=True)
    try:
        rows_load = list(wb.worksheets[0].iter_rows(min_row=2, values_only=True))
        rows_pv = list(wb.worksheets[1].iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()
    dates = [value[0].date().isoformat() if isinstance(value[0], datetime) else str(value[0])[:10] for value in rows_load]
    load = np.asarray([[float(x) for x in row[1:]] for row in rows_load]) / 6.0
    pv = np.asarray([[float(x) for x in row[1:]] for row in rows_pv]) / 6.0
    wb = load_workbook(ATT / "附件4.xlsx", read_only=True, data_only=True)
    try:
        price_rows = list(wb.worksheets[0].iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()
    price = np.asarray([[float(x) for x in row[1:]] for row in price_rows])
    return dates, load, pv, price


def weekday_price(price: np.ndarray, dates: list[str], day_idx: int) -> np.ndarray:
    target = datetime.fromisoformat(dates[day_idx]).date()
    candidates = [i for i in range(max(0, day_idx - 28), day_idx) if datetime.fromisoformat(dates[i]).date().weekday() == target.weekday()]
    if not candidates:
        candidates = [day_idx - 1]
    return np.mean(price[candidates], axis=0)


def event_table(detail: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for day, group in detail.groupby("date", sort=False):
        values = group.sort_values("slot")["emergency_purchase_kwh"].to_numpy(float)
        start = None
        event_id = 0
        for idx, value in enumerate(values):
            if value > 1e-7 and start is None:
                start = idx
            if start is not None and (value <= 1e-7 or idx == len(values) - 1):
                end = idx if value <= 1e-7 else idx + 1
                event_id += 1
                rows.append((day, event_id, start + 1, end + 1, float(np.sum(values[start:end]))))
                start = None
    return pd.DataFrame(rows, columns=["date", "event_id", "start_slot", "end_slot_exclusive", "purchase_kwh"])


def check(prefix: str, q43: bool, dates: list[str], load: np.ndarray, pv: np.ndarray, price: np.ndarray) -> dict:
    detail = pd.read_csv(ROOT / "results" / f"{prefix}_detail.csv")
    daily = pd.read_csv(ROOT / "results" / f"{prefix}_daily.csv")
    storage = pd.read_csv(ROOT / "results" / f"{prefix}_storage_4h.csv")
    events = pd.read_csv(ROOT / "results" / f"{prefix}_emergency_events.csv")
    date_to_idx = {d: i for i, d in enumerate(dates)}
    source_diff = price_diff = forecast_price_diff = balance_diff = state_diff = 0.0
    cost_diff = 0.0
    past_lock = 0
    q43_price_cache: dict[tuple[int, int], np.ndarray] = {}
    for row in detail.itertuples(index=False):
        i, s = date_to_idx[str(row.date)], int(row.slot) - 1
        source_diff = max(source_diff, abs(float(row.load_kwh) - load[i, s]), abs(float(row.pv_kwh) - pv[i, s]))
        price_diff = max(price_diff, abs(float(row.actual_price_yuan_per_kwh) - price[i, s]))
        base_price_today = weekday_price(price, dates, i)
        if q43:
            issue_slot = {"00:00": 0, "06:00": 36, "12:00": 72, "18:00": 108}[str(row.executed_plan_issue_time)]
            key = (i, issue_slot)
            if key not in q43_price_cache:
                if issue_slot == 0:
                    q43_price_cache[key] = base_price_today.copy()
                else:
                    scale = float(np.sum(price[i, :issue_slot]) / np.sum(base_price_today[:issue_slot]))
                    residual_rows = []
                    for history_idx in range(max(7, i - 28), i):
                        history_base = weekday_price(price, dates, history_idx)
                        residual_rows.append(price[history_idx] - history_base)
                    residual = np.asarray(residual_rows)
                    lag = residual[:, :-1].reshape(-1)
                    lead = residual[:, 1:].reshape(-1)
                    denominator = float(np.dot(lag, lag))
                    rho = 0.0 if denominator <= 0.0 else float(np.clip(np.dot(lag, lead) / denominator, -0.99, 0.99))
                    last_residual = float(price[i, issue_slot - 1] - base_price_today[issue_slot - 1] * scale)
                    remaining = base_price_today[issue_slot:] * scale + last_residual * rho ** np.arange(1, 145 - issue_slot)
                    q43_price_cache[key] = np.maximum(remaining, 0.0)
            expected_forecast = q43_price_cache[key][s if issue_slot == 0 else s - issue_slot]
        else:
            issue_slot = 0
            expected_forecast = base_price_today[s]
        forecast_price_diff = max(forecast_price_diff, abs(float(row.forecast_price_yuan_per_kwh) - expected_forecast))
        balance = float(row.normal_purchase_kwh if not q43 else row.effective_purchase_kwh) + float(row.pv_kwh) + float(row.discharge_bus_kwh) + float(row.emergency_purchase_kwh) - float(row.load_kwh) - float(row.charge_bus_kwh) - float(row.surplus_discard_kwh)
        state = float(row.soc_close_kwh) - float(row.soc_open_kwh) - 0.9 * float(row.charge_bus_kwh) + float(row.discharge_bus_kwh) / 0.9
        balance_diff = max(balance_diff, abs(balance))
        state_diff = max(state_diff, abs(state))
        if q43:
            calc = price[i, s] * (float(row.original_purchase_kwh) - 0.5 * float(row.downward_adjustment_kwh) + 1.5 * float(row.upward_adjustment_kwh) + 5.0 * float(row.emergency_purchase_kwh))
            past_lock += int(s < issue_slot)
        else:
            calc = price[i, s] * (float(row.normal_purchase_kwh) + 5.0 * float(row.emergency_purchase_kwh))
        cost_diff = max(cost_diff, abs(calc - float(row.total_cost_yuan if q43 else row.normal_cost_yuan + row.emergency_cost_yuan)))

    grouped = detail.groupby("date", sort=False)
    total_recomputed = grouped["total_cost_yuan"].sum() if q43 else grouped["normal_cost_yuan"].sum() + grouped["emergency_cost_yuan"].sum()
    daily_diff = float(np.max(np.abs(total_recomputed.to_numpy() - daily["total_cost_yuan"].to_numpy())))
    expected_events = event_table(detail)
    keys = ["date", "event_id", "start_slot", "end_slot_exclusive"]
    events_exact = len(expected_events) == len(events) and expected_events[keys].astype(str).equals(events[keys].astype(str))
    event_diff = float(np.max(np.abs(expected_events["purchase_kwh"].to_numpy() - events["purchase_kwh"].to_numpy()))) if len(events) else 0.0
    soc_min = float(detail[["soc_open_kwh", "soc_close_kwh"]].min().min())
    soc_max = float(detail[["soc_open_kwh", "soc_close_kwh"]].max().max())
    max_charge = float(detail["charge_bus_kwh"].max())
    max_discharge = float(detail["discharge_bus_kwh"].max())
    crossday = float(np.max(np.abs(daily["soc_open_kwh"].to_numpy()[1:] - daily["soc_close_kwh"].to_numpy()[:-1])))
    info = int(np.sum(pd.to_datetime(detail["train_end_date"]) >= pd.to_datetime(detail["date"])))
    flags = {
        "row_counts": len(detail) == 334 * 144 and len(daily) == 334 and len(storage) == 334 * 6,
        "source": source_diff <= TOL,
        "actual_price": price_diff <= TOL,
        "causal_price_forecast": forecast_price_diff <= TOL,
        "slot_cost": cost_diff <= TOL,
        "balance": balance_diff <= TOL,
        "state": state_diff <= TOL,
        "soc_bounds": soc_min >= 1200 - TOL and soc_max <= 10800 + TOL,
        "storage_power": max_charge <= FLOW_MAX_KWH + TOL and max_discharge <= FLOW_MAX_KWH + TOL,
        "crossday": crossday <= TOL,
        "daily": daily_diff <= TOL,
        "events": events_exact and event_diff <= TOL,
        "information_cutoff": info == 0,
        "past_lock": past_lock == 0,
    }
    flags = {key: bool(value) for key, value in flags.items()}
    return {
        "status": "PASS" if all(flags.values()) else "FAIL",
        "pass_flags": flags,
        "metrics": {"detail_rows": len(detail), "daily_rows": len(daily), "storage_rows": len(storage), "event_rows": len(events), "max_source_difference": source_diff, "max_actual_price_difference": price_diff, "max_causal_price_forecast_difference": forecast_price_diff, "max_slot_cost_difference": cost_diff, "max_balance_residual_kwh": balance_diff, "max_state_residual_kwh": state_diff, "soc_min_kwh": soc_min, "soc_max_kwh": soc_max,
                    "storage_power_limit_kw": 5000.0, "storage_flow_limit_kwh_per_slot": FLOW_MAX_KWH,
                    "max_charge_kwh_per_slot": max_charge, "max_discharge_kwh_per_slot": max_discharge,
                    "max_crossday_soc_gap_kwh": crossday, "max_daily_cost_difference_yuan": daily_diff, "events_exact": events_exact, "max_event_amount_difference_kwh": event_diff, "information_cutoff_violations": info, "past_lock_violations": past_lock},
    }


def main() -> None:
    dates, load, pv, price = load_source()
    q42 = check("q4_2", False, dates, load, pv, price)
    q43 = check("q4_3", True, dates, load, pv, price)
    output = {"schema_version": 1, "question": "Q4", "status": "PASS" if q42["status"] == q43["status"] == "PASS" else "FAIL", "q4_2": q42, "q4_3": q43, "independence": "Reads attachments 2 and 4 plus frozen Q4 CSV outputs directly; does not import common_data, dispatch_core, q2_solve, q3_solve, or q4_solve."}
    path = ROOT / "validation" / "q4_independent_recompute.json"
    path.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": output["status"], "q4_2": q42["metrics"], "q4_3": q43["metrics"]}, ensure_ascii=False))
    if output["status"] != "PASS":
        raise RuntimeError("Q4独立复算失败")


if __name__ == "__main__":
    main()
