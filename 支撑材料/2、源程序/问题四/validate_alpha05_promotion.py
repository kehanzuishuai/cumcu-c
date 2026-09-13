from __future__ import annotations

import json
import os
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(os.environ.get("CUMCM_PROJECT_ROOT", Path.cwd())).resolve()
STAGE = ROOT / "_promotion_staging" / "alpha05_final"
ATT = ROOT / "source_materials" / "C题" / "附件"
TOL = 1e-6
POS = 1e-7
FLOW = 5000.0 / 6.0


def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value)[:10]).date()


def load_sources():
    wb = load_workbook(ATT / "附件1.xlsx", read_only=True, data_only=True)
    try:
        fixed = np.asarray([float(r[1]) for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True)])
    finally:
        wb.close()
    wb = load_workbook(ATT / "附件2.xlsx", read_only=True, data_only=True)
    try:
        mats, dates = [], None
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            current = [parse_date(r[0]) for r in rows]
            if dates is None:
                dates = current
            elif dates != current:
                raise ValueError("附件2日期轴不一致")
            mats.append(np.asarray([[float(v) for v in r[1:]] for r in rows]))
    finally:
        wb.close()
    wb = load_workbook(ATT / "附件4.xlsx", read_only=True, data_only=True)
    try:
        rows = list(wb.worksheets[0].iter_rows(min_row=2, values_only=True))
        realtime = np.asarray([[float(v) for v in r[1:]] for r in rows])
    finally:
        wb.close()
    return dates, mats[0], mats[1], fixed, realtime


def weighted_load(load: np.ndarray, dates: list[date], day_idx: int) -> np.ndarray:
    cand = [i for i in range(max(0, day_idx - 28), day_idx) if dates[i].weekday() == dates[day_idx].weekday()]
    if not cand:
        cand = [day_idx - 1]
    return np.average(load[cand], axis=0, weights=np.arange(1, len(cand) + 1, dtype=float))


def weekday_price(price: np.ndarray, dates: list[date], day_idx: int) -> np.ndarray:
    cand = [i for i in range(max(0, day_idx - 28), day_idx) if dates[i].weekday() == dates[day_idx].weekday()]
    if not cand:
        cand = [day_idx - 1]
    return np.mean(price[cand], axis=0)


def rebuild_events(detail: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for day, group in detail.groupby("date", sort=True):
        values = group.sort_values("slot")["emergency_purchase_kwh"].to_numpy(float)
        start, event_id = None, 0
        for i, v in enumerate(values):
            if v > POS and start is None:
                start = i
            if start is not None and (v <= POS or i == 143):
                end = i if v <= POS else i + 1
                event_id += 1
                rows.append((day, event_id, start + 1, end + 1, float(values[start:end].sum())))
                start = None
    return pd.DataFrame(rows, columns=["date", "event_id", "start_slot", "end_slot_exclusive", "purchase_kwh"])


def check(prefix: str, q43: bool, dates, load, pv, fixed, realtime):
    base = STAGE / "results"
    detail = pd.read_csv(base / f"{prefix}_detail.csv", encoding="utf-8-sig").sort_values(["date", "slot"]).reset_index(drop=True)
    daily = pd.read_csv(base / f"{prefix}_daily.csv", encoding="utf-8-sig").sort_values("date").reset_index(drop=True)
    storage = pd.read_csv(base / f"{prefix}_storage_4h.csv", encoding="utf-8-sig").sort_values(["date", "block"]).reset_index(drop=True)
    events = pd.read_csv(base / f"{prefix}_emergency_events.csv", encoding="utf-8-sig")
    versions = pd.read_csv(base / f"{prefix}_plan_versions.csv", encoding="utf-8-sig")
    summary = json.loads((base / ("q4_summary.json" if q43 else "q3_summary.json")).read_text(encoding="utf-8"))
    dmap = {d.isoformat(): i for i, d in enumerate(dates)}
    source_diff = actual_price_diff = forecast_price_diff = load_feedback_diff = 0.0
    for day, group in detail.groupby("date", sort=True):
        i = dmap[str(day)]
        g = group.sort_values("slot")
        slots = g["slot"].to_numpy(int) - 1
        source_diff = max(source_diff, float(np.max(np.abs(g["actual_load_kw"].to_numpy() - load[i, slots]))), float(np.max(np.abs(g["actual_pv_kw"].to_numpy() - pv[i, slots]))))
        base_load = weighted_load(load, dates, i)
        base_price = weekday_price(realtime, dates, i)
        for row in g.itertuples(index=False):
            s = int(row.slot) - 1
            issue = int(str(row.executed_plan_issue_time)[:2])
            start = issue * 6
            ratio = 1.0
            if start:
                raw = float(load[i, :start].sum() / base_load[:start].sum())
                ratio = 1.0 + .5 * (raw - 1.0)
            load_feedback_diff = max(load_feedback_diff, abs(float(row.forecast_load_kw) - float(base_load[s] * ratio)))
            expected_actual_price = float(realtime[i, s] if q43 else fixed[s])
            actual_price_diff = max(actual_price_diff, abs(float(row.actual_price_yuan_per_kwh) - expected_actual_price))
            if q43:
                forecast_price_diff = max(forecast_price_diff, abs(float(row.forecast_price_yuan_per_kwh) - float(base_price[s])))
    q0 = detail["original_purchase_kwh"].to_numpy(float)
    effective = detail["effective_purchase_kwh"].to_numpy(float)
    upward, downward = np.maximum(effective - q0, 0), np.maximum(q0 - effective, 0)
    price = np.asarray([realtime[dmap[str(r.date)], int(r.slot) - 1] if q43 else fixed[int(r.slot) - 1] for r in detail.itertuples(index=False)])
    base_cost, down_cost, up_cost = price*q0, .5*price*downward, 1.5*price*upward
    normal, emerg_cost = base_cost-down_cost+up_cost, 5*price*detail["emergency_purchase_kwh"].to_numpy(float)
    total = normal + emerg_cost
    cost_diff = float(max(
        np.max(np.abs(upward-detail["upward_adjustment_kwh"].to_numpy(float))),
        np.max(np.abs(downward-detail["downward_adjustment_kwh"].to_numpy(float))),
        np.max(np.abs(base_cost-detail["base_plan_cost_yuan"].to_numpy(float))),
        np.max(np.abs(down_cost-detail["downward_credit_yuan"].to_numpy(float))),
        np.max(np.abs(up_cost-detail["upward_adjustment_cost_yuan"].to_numpy(float))),
        np.max(np.abs(normal-detail["normal_settlement_cost_yuan"].to_numpy(float))),
        np.max(np.abs(emerg_cost-detail["emergency_cost_yuan"].to_numpy(float))),
        np.max(np.abs(total-detail["total_cost_yuan"].to_numpy(float))),
    ))
    balance = effective + detail["pv_kwh"] + detail["discharge_bus_kwh"] + detail["emergency_purchase_kwh"] - detail["load_kwh"] - detail["charge_bus_kwh"] - detail["surplus_discard_kwh"]
    state = detail["soc_close_kwh"] - detail["soc_open_kwh"] - .9*detail["charge_bus_kwh"] + detail["discharge_bus_kwh"]/.9
    grouped = detail.assign(_total=total).groupby("date", sort=True)
    daily_diff = float(np.max(np.abs(grouped["_total"].sum().to_numpy() - daily["total_cost_yuan"].to_numpy())))
    temp = detail.assign(block=(detail["slot"].to_numpy(int)-1)//24+1)
    rebuilt_storage = temp.groupby(["date","block"],sort=True).agg(charge_kwh=("charge_bus_kwh","sum"),discharge_kwh=("discharge_bus_kwh","sum")).reset_index()
    storage_diff = float(max(np.max(np.abs(rebuilt_storage[c].to_numpy()-storage[c].to_numpy())) for c in ("charge_kwh","discharge_kwh")))
    expected_events = rebuild_events(detail)
    keys = ["date","event_id","start_slot","end_slot_exclusive"]
    event_keys = len(events)==len(expected_events) and expected_events[keys].astype(str).equals(events[keys].astype(str))
    event_diff = float(np.max(np.abs(expected_events["purchase_kwh"].to_numpy()-events["purchase_kwh"].to_numpy()))) if len(events) else 0.0
    latest = versions.sort_values(["date","target_slot","issue_slot"]).groupby(["date","target_slot"],as_index=False).tail(1).sort_values(["date","target_slot"]).reset_index(drop=True)
    version_diff = float(np.max(np.abs(latest["new_effective_plan_kwh"].to_numpy()-effective)))
    past_lock = int(np.sum((versions["issue_slot"]>0)&(versions["target_slot"]<=versions["issue_slot"])))
    info = int(np.sum(pd.to_datetime(detail["train_end_date"])>=pd.to_datetime(detail["date"])))
    cross = float(np.max(np.abs(daily["soc_open_kwh"].to_numpy()[1:]-daily["soc_close_kwh"].to_numpy()[:-1])))
    simultaneous = int(np.sum((detail["charge_bus_kwh"]>POS)&(detail["discharge_bus_kwh"]>POS)))
    emerg_charge = int(np.sum((detail["emergency_purchase_kwh"]>POS)&(detail["charge_bus_kwh"]>POS)))
    main = summary["q4_3"]["strategies"]["q4_3_all_main"] if q43 else summary["strategies"]["q3_all_main"]
    summary_diff = max(abs(float(total.sum())-float(main["total_cost_yuan"])),abs(float(detail["emergency_purchase_kwh"].sum())-float(main["emergency_purchase_kwh"])),abs(float(detail.iloc[-1]["soc_close_kwh"])-float(main["soc_final_kwh"])))
    metrics = {"max_source_difference":source_diff,"max_actual_price_difference":actual_price_diff,"max_price_forecast_difference":forecast_price_diff,
        "max_alpha05_load_forecast_difference":load_feedback_diff,"max_cost_difference":cost_diff,"max_balance_residual_kwh":float(np.max(np.abs(balance))),
        "max_state_residual_kwh":float(np.max(np.abs(state))),"max_daily_cost_difference_yuan":daily_diff,"max_storage_difference_kwh":storage_diff,
        "event_keys_exact":bool(event_keys),"max_event_difference_kwh":event_diff,"max_version_difference_kwh":version_diff,"past_lock_violations":past_lock,
        "information_cutoff_violations":info,"max_crossday_soc_gap_kwh":cross,"simultaneous_slots":simultaneous,"emergency_charge_slots":emerg_charge,
        "soc_min_kwh":float(detail[["soc_open_kwh","soc_close_kwh"]].min().min()),"soc_max_kwh":float(detail[["soc_open_kwh","soc_close_kwh"]].max().max()),
        "max_charge_kwh":float(detail["charge_bus_kwh"].max()),"max_discharge_kwh":float(detail["discharge_bus_kwh"].max()),"summary_difference":summary_diff}
    flags = {"rows":len(detail)==334*144 and len(daily)==334 and len(storage)==334*6,"source":source_diff<=1e-9,"actual_price":actual_price_diff<=TOL,
        "price_forecast_no_ar1":(not q43) or forecast_price_diff<=TOL,"alpha05_load_feedback":load_feedback_diff<=TOL,"cost":cost_diff<=TOL,
        "balance":metrics["max_balance_residual_kwh"]<=TOL,"state":metrics["max_state_residual_kwh"]<=TOL,
        "soc_bounds":metrics["soc_min_kwh"]>=1200-TOL and metrics["soc_max_kwh"]<=10800+TOL,
        "storage_power":metrics["max_charge_kwh"]<=FLOW+TOL and metrics["max_discharge_kwh"]<=FLOW+TOL,
        "crossday":cross<=TOL,"mutual_exclusion":simultaneous==0,"emergency_not_charging":emerg_charge==0,
        "daily":daily_diff<=TOL,"storage":storage_diff<=TOL,"events":event_keys and event_diff<=TOL,
        "versions":past_lock==0 and version_diff<=TOL,"information_cutoff":info==0,"summary":summary_diff<=TOL}
    return {"status":"PASS" if all(flags.values()) else "FAIL","pass_flags":{k:bool(v) for k,v in flags.items()},"metrics":metrics}


def main():
    dates, load, pv, fixed, realtime = load_sources()
    q3 = check("q3",False,dates,load,pv,fixed,realtime)
    q43 = check("q4_3",True,dates,load,pv,fixed,realtime)
    native = json.loads((STAGE/"validation/promotion_gate_native.json").read_text(encoding="utf-8"))
    payload = {"schema_version":1,"status":"PASS" if q3["status"]==q43["status"]=="PASS" and native["status"]=="PASS" else "FAIL",
        "q3":q3,"q4_3":q43,"terminal_soc_inventory_comparability":native["terminal_soc_inventory_comparability"],
        "independence":"Reads raw attachments and staged CSV/JSON directly; does not import common_data, dispatch_core, q2_solve, q3_solve, or q4_solve."}
    (STAGE/"validation").mkdir(parents=True,exist_ok=True)
    (STAGE/"validation/promotion_gate_independent.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding="utf-8")
    (STAGE/"validation/q3_independent_recompute.json").write_text(json.dumps({"schema_version":1,"question":"Q3",**q3},ensure_ascii=False,indent=2),encoding="utf-8")
    # Preserve Q4-2 evidence by referencing its unchanged formal validation while independently recomputing only the promoted Q4-3.
    formal_q4 = json.loads((ROOT/"validation/q4_independent_recompute.json").read_text(encoding="utf-8"))
    (STAGE/"validation/q4_independent_recompute.json").write_text(json.dumps({"schema_version":1,"question":"Q4","status":"PASS" if formal_q4["q4_2"]["status"]=="PASS" and q43["status"]=="PASS" else "FAIL","q4_2":formal_q4["q4_2"],"q4_3":q43,"independence":payload["independence"]},ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"status":payload["status"],"q3":q3["metrics"],"q4_3":q43["metrics"]},ensure_ascii=False))
    if payload["status"]!="PASS":
        raise RuntimeError("independent promotion gate failed")


if __name__=="__main__":
    main()
