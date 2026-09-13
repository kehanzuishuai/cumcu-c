from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
ATTACHMENT2 = ROOT / "source_materials" / "C题" / "附件" / "附件2.xlsx"
DETAIL = ROOT / "results" / "q2_component_ablation_detail.csv"
DAILY = ROOT / "results" / "q2_component_ablation_daily.csv"
SUMMARY = ROOT / "results" / "q2_component_ablation_summary.json"
OUTPUT = ROOT / "validation" / "q2_component_ablation_independent.json"
STRATEGIES = ["A0_base_point", "A1_fixed_q80", "A2_dynamic_risk", "A3_dynamic_plus_terminal_value", "O2_online_scheme_reconstruction"]
TOL = 1e-6
POS_TOL = 1e-7
FLOW_LIMIT = 5000.0 / 6.0


def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value)).date()


def load_source() -> tuple[list[date], np.ndarray, np.ndarray]:
    workbook = load_workbook(ATTACHMENT2, read_only=True, data_only=True)
    try:
        matrices = []
        dates = None
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            current_dates = [parse_date(row[0]) for row in rows]
            if dates is None:
                dates = current_dates
            elif dates != current_dates:
                raise ValueError("附件2日期轴不一致")
            matrices.append(np.asarray([[float(value) for value in row[1:]] for row in rows]))
    finally:
        workbook.close()
    assert dates is not None
    return dates, matrices[0], matrices[1]


def max_abs(values: np.ndarray) -> float:
    return float(np.max(np.abs(values))) if values.size else 0.0


def main() -> None:
    source_dates, source_load, source_pv = load_source()
    detail = pd.read_csv(DETAIL, encoding="utf-8-sig", dtype={"date": str, "train_end_date": str, "risk_history_end_date": str})
    daily = pd.read_csv(DAILY, encoding="utf-8-sig", dtype={"date": str, "train_end_date": str, "risk_history_start_date": str, "risk_history_end_date": str})
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    expected_load = source_load[31:].reshape(-1)
    expected_pv = source_pv[31:].reshape(-1)
    strategy_results = {}
    all_pass = True
    for strategy in STRATEGIES:
        rows = detail.loc[detail["strategy"] == strategy].sort_values(["date", "slot"]).reset_index(drop=True)
        days = daily.loc[daily["strategy"] == strategy].sort_values("date").reset_index(drop=True)
        load_diff = max_abs(rows["actual_load_kw"].to_numpy() - expected_load)
        pv_diff = max_abs(rows["actual_pv_kw"].to_numpy() - expected_pv)
        point_link = max_abs(rows["effective_net_kw"].to_numpy() - rows["point_net_kw"].to_numpy() - rows["risk_uplift_kw"].to_numpy())
        load_kwh = rows["actual_load_kw"].to_numpy() / 6.0
        pv_kwh = rows["actual_pv_kw"].to_numpy() / 6.0
        balance = rows["normal_purchase_kwh"].to_numpy() + pv_kwh + rows["discharge_kwh"].to_numpy() + rows["emergency_purchase_kwh"].to_numpy() - load_kwh - rows["charge_kwh"].to_numpy() - rows["surplus_kwh"].to_numpy()
        state = rows["soc_close_kwh"].to_numpy() - rows["soc_open_kwh"].to_numpy() - 0.9 * rows["charge_kwh"].to_numpy() + rows["discharge_kwh"].to_numpy() / 0.9
        normal_cost = rows["price_yuan_per_kwh"].to_numpy() * rows["normal_purchase_kwh"].to_numpy()
        emergency_cost = 5.0 * rows["price_yuan_per_kwh"].to_numpy() * rows["emergency_purchase_kwh"].to_numpy()
        cost_diff = max(max_abs(normal_cost - rows["normal_cost_yuan"].to_numpy()), max_abs(emergency_cost - rows["emergency_cost_yuan"].to_numpy()))
        recomputed = rows.groupby("date", sort=True).agg(
            normal_purchase_kwh=("normal_purchase_kwh", "sum"),
            emergency_purchase_kwh=("emergency_purchase_kwh", "sum"),
            normal_cost_yuan=("normal_cost_yuan", "sum"),
            emergency_cost_yuan=("emergency_cost_yuan", "sum"),
            soc_open_kwh=("soc_open_kwh", "first"),
            soc_close_kwh=("soc_close_kwh", "last"),
        ).reset_index()
        recomputed["total_cost_yuan"] = recomputed["normal_cost_yuan"] + recomputed["emergency_cost_yuan"]
        daily_cols = ["normal_purchase_kwh", "emergency_purchase_kwh", "normal_cost_yuan", "emergency_cost_yuan", "total_cost_yuan", "soc_open_kwh", "soc_close_kwh"]
        daily_diff = max(max_abs(recomputed[column].to_numpy() - days[column].to_numpy()) for column in daily_cols)
        crossday = max_abs(days["soc_open_kwh"].to_numpy()[1:] - days["soc_close_kwh"].to_numpy()[:-1])
        selected = rows["selected_quantile"].dropna().unique()
        quantile_valid = strategy == "A0_base_point" and len(selected) == 0 or strategy != "A0_base_point" and set(np.round(selected.astype(float), 10)).issubset({0.7, 0.8, 0.9})
        fixed_q80 = strategy != "A1_fixed_q80" or set(np.round(selected.astype(float), 10)) == {0.8}
        info_violations = int(np.sum(rows["train_end_date"] >= rows["date"]) + np.sum(rows["risk_history_end_date"] >= rows["date"]))
        simultaneous = int(np.sum((rows["charge_kwh"] > POS_TOL) & (rows["discharge_kwh"] > POS_TOL)))
        emergency_charge = int(np.sum((rows["emergency_purchase_kwh"] > POS_TOL) & (rows["charge_kwh"] > POS_TOL)))
        aggregate = summary["strategies"][strategy]
        aggregate_diff = max(
            abs(float(recomputed[column].sum()) - float(aggregate[column]))
            for column in ("normal_purchase_kwh", "emergency_purchase_kwh", "normal_cost_yuan", "emergency_cost_yuan", "total_cost_yuan")
        )
        aggregate_diff = max(aggregate_diff, abs(float(recomputed.iloc[-1]["soc_close_kwh"]) - float(aggregate["soc_final_kwh"])))
        metrics = {
            "detail_rows": int(len(rows)), "daily_rows": int(len(days)),
            "max_source_load_difference_kw": load_diff, "max_source_pv_difference_kw": pv_diff,
            "max_effective_forecast_link_difference_kw": point_link,
            "max_balance_residual_kwh": max_abs(balance), "max_state_residual_kwh": max_abs(state),
            "max_slot_cost_difference_yuan": cost_diff, "max_daily_difference": daily_diff,
            "max_crossday_soc_gap_kwh": crossday,
            "soc_min_kwh": float(min(rows["soc_open_kwh"].min(), rows["soc_close_kwh"].min())),
            "soc_max_kwh": float(max(rows["soc_open_kwh"].max(), rows["soc_close_kwh"].max())),
            "max_charge_kwh_per_slot": float(rows["charge_kwh"].max()),
            "max_discharge_kwh_per_slot": float(rows["discharge_kwh"].max()),
            "simultaneous_slots": simultaneous, "emergency_charge_overlap_slots": emergency_charge,
            "information_cutoff_violations": info_violations, "max_aggregate_difference": aggregate_diff,
        }
        flags = {
            "row_counts": len(rows) == 334 * 144 and len(days) == 334,
            "source": max(load_diff, pv_diff) <= 1e-9, "forecast_link": point_link <= 1e-9,
            "balance": metrics["max_balance_residual_kwh"] <= TOL, "state": metrics["max_state_residual_kwh"] <= TOL,
            "cost": cost_diff <= TOL and daily_diff <= TOL and aggregate_diff <= TOL,
            "soc": metrics["soc_min_kwh"] >= 1200 - TOL and metrics["soc_max_kwh"] <= 10800 + TOL and crossday <= TOL,
            "power": metrics["max_charge_kwh_per_slot"] <= FLOW_LIMIT + TOL and metrics["max_discharge_kwh_per_slot"] <= FLOW_LIMIT + TOL,
            "mutual_exclusion": simultaneous == 0, "emergency_not_charging": emergency_charge == 0,
            "information": info_violations == 0, "risk_grid": bool(quantile_valid and fixed_q80),
        }
        status = "PASS" if all(flags.values()) else "FAIL"
        all_pass = all_pass and status == "PASS"
        strategy_results[strategy] = {"status": status, "pass_flags": flags, "metrics": metrics}
    frozen_unchanged = bool(summary["frozen_main_unchanged"] and summary["frozen_hashes_before"] == summary["frozen_hashes_after"])
    payload = {
        "schema_version": 1, "question": "Q2 component ablation", "status": "PASS" if all_pass and frozen_unchanged else "FAIL",
        "strategies": strategy_results, "frozen_main_unchanged": frozen_unchanged,
        "independence": "Reads attachment2 and ablation CSV/JSON directly; does not import q2_component_ablation, q2_solve, dispatch_core, or common_data.",
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": payload["status"], "strategies": {k: v["status"] for k, v in strategy_results.items()}, "frozen_main_unchanged": frozen_unchanged}, ensure_ascii=False))
    if payload["status"] != "PASS":
        raise RuntimeError("Q2组件消融独立复算失败")


if __name__ == "__main__":
    main()
