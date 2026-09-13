#!/usr/bin/env python3
"""Read-only deterministic pre-audit for a CUMCM C submission project.

This script intentionally does not decide mathematical correctness. It gathers
evidence for the semantic and visual stages defined by the skill.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET


TEXT_EXTENSIONS = {
    ".tex", ".md", ".txt", ".py", ".m", ".r", ".jl", ".c", ".cpp",
    ".h", ".hpp", ".json", ".yaml", ".yml", ".csv", ".tsv", ".ipynb",
}
PAPER_EXTENSIONS = {".pdf", ".docx", ".tex", ".md"}
BUILD_SUFFIXES = {
    ".aux", ".log", ".out", ".fls", ".fdb_latexmk", ".xdv", ".synctex.gz",
}
CACHE_DIRS = {"__pycache__", ".ipynb_checkpoints", "cache", ".cache"}
OLD_NAME_RE = re.compile(r"(?:old|backup|bak|copy|final[_ -]?\d+|旧|备份|副本)", re.I)
IMAGE_SUFFIXES = {".pdf", ".png", ".jpg", ".jpeg", ".svg", ".eps"}
DRAWING_PLACEHOLDER_RE = re.compile(
    r"(?:待绘制|待补(?:充)?(?:流程图|图|表)|待插入(?:流程图|图|表)|图表待补|"
    r"此处(?:待)?(?:插入|放置|添加)(?:[^\n。；;]{0,10})?(?:流程图|图|表)|"
    r"(?:TODO|TBD)\s*[:：-]?\s*(?:figure|table|图|表)|"
    r"(?:figure|table)\s*placeholder)",
    re.I,
)
SOURCE_ATTRIBUTION_RE = re.compile(
    r"(?:题目|题设|赛题|题中|附件(?:\s*\d+)?)"
    r"(?:中)?(?:已)?(?:给出|给定|提供|规定|要求|可知|表明)"
    r"|(?:根据|依据|由)(?:题目|题设|赛题|题意|附件(?:\s*\d+)?)(?:可知|得|中)?"
    r"|(?:原始|附件|模拟|关联|样本|上述)?数据(?:中)?(?:给出|提供|表明|显示|可知)",
    re.I,
)
SOURCE_SENSITIVE_TERMS = (
    "参数", "阈值", "权重", "分布", "相关系数", "相关矩阵", "关系矩阵",
    "邻接矩阵", "协方差矩阵", "替代边", "互补边", "模拟边", "置信水平",
    "风险系数", "服务水平", "聚类数", "窗口", "情景", "概率", "边数",
)
PAPER_ONLY_DEFAULT_DEFERRALS = (
    ("SCOPE-FULL-001", "代码执行与复现", "PAPER_ONLY 不运行或判定代码；转 FULL 核验。"),
    ("SCOPE-FULL-002", "结果文件与模型—代码—结果一致性", "PAPER_ONLY 不因结果文件未提供而失败；转 FULL 核验。"),
    ("SCOPE-FULL-003", "支撑材料目录/压缩包与工程卫生", "PAPER_ONLY 不审支撑包和构建垃圾；转 FULL 核验。"),
    ("SCOPE-FULL-004", "AI 工具使用详情.pdf", "PAPER_ONLY 只审正文 AI 声明；详情文件存在性转 FULL 核验。"),
    ("SCOPE-FULL-006", "最终 PDF 逐页视觉确认", "PAPER_ONLY 可以仅以 LaTeX/Markdown 源码完成正文审核；最终 PDF 视觉检查转 FULL/VISUAL_ONLY。"),
)

META_PHRASES = [
    "需要说明的是", "值得注意的是", "需要强调的是", "进一步表明", "由此形成",
    "从而构建", "统一框架", "递进式体系", "提供决策支持", "具有重要意义",
    "效果显著", "具有较高参考价值", "提供理论依据", "具有重要参考价值",
]
LOGIC_CONNECTORS = (
    "由于", "因此", "从而", "进而", "由此", "所以", "故而", "虽然", "但是", "然而",
    "同时", "此外", "并且", "不仅", "而且", "若", "则", "当", "如果", "为了", "通过",
    "基于", "考虑到", "在此基础上", "另一方面", "进一步",
)
ABSTRACT_NOUN_TERMS = (
    "机制", "体系", "框架", "能力", "水平", "效应", "关系", "特征", "结构", "模式",
    "过程", "策略", "路径", "逻辑", "维度", "价值", "意义", "支撑", "赋能", "协同",
    "闭环", "优化", "提升", "构建", "实现",
)
DEFAULT_TERMINOLOGY_GROUPS = (
    ("暖启动", ("暖启动", "热启动")),
    ("样本外评估", ("样本外评测", "样本外评价", "样本外评估")),
    ("超额产量", ("超额产量", "超产量")),
    ("关联情景方案", ("关联情景方案", "关联方案")),
)
SHOWCASE_CITATION_RE = re.compile(
    r"(?:(?:相关|具体|上述|该)(?:方法|算法|模型|理论|证明)?\s*)?"
    r"(?:详见|参见|见|可参见|可参考|可见于)\s*(?:相关)?文献(?!综述|列表|回顾|计量)",
)
MODEL_PATTERNS = {
    "LMM/线性混合效应模型": r"线性混合效应|\bLMM\b",
    "Cox": r"\bCox\b|比例风险模型",
    "AFT": r"\bAFT\b|加速失效",
    "XGBoost": r"\bXGBoost\b",
    "LightGBM": r"\bLightGBM\b",
    "Prophet": r"\bProphet\b",
    "CVaR": r"\bCVaR\b|条件风险价值",
    "NSGA-II": r"NSGA[-‐‑–—]?II",
    "报童模型": r"报童模型|newsvendor",
    "ARIMA": r"\b(?:S?ARIMA)\b",
    "TOPSIS": r"\bTOPSIS\b",
}
CITATION_COMMAND_RE = re.compile(
    r"\\(?P<command>[A-Za-z@]*cite[A-Za-z@]*)\*?"
    r"(?:\s*\[[^\]]*\]){0,2}\s*\{(?P<keys>[^{}]+)\}",
    re.I,
)
CITATION_CONTEXT_TERMS = (
    "参考", "文献", "研究", "提出", "指出", "表明", "证明", "根据", "见", "采用",
    "模型", "算法", "方法", "理论", "定理", "数据来源", "标准", "已有工作",
)
PAPER_ONLY_CORE_MODULES = (
    "段落与排版源码",
    "AI/模板化语言",
    "长句与可读性",
    "段落/小节过渡",
    "模型百科式表达",
    "术语一致性",
    "问题分析与最终模型边界",
    "模型选择证据链",
    "跨问接口",
    "模型输出→最终决策映射",
    "图表正文叙事",
    "公式与符号",
    "引用格式",
    "引用跳转",
    "引用语义正确性",
    "参考文献数量与正文对应关系",
)


@dataclass
class Finding:
    priority: str
    check_id: str
    title: str
    location: str
    evidence: str
    rule: str
    recommendation: str
    auto_fixable: bool = False
    certainty: str = "DETERMINISTIC"
    stage_status: str = "CONFIRMED_IN_SCOPE"


def strip_latex_comments(text: str) -> str:
    cleaned: list[str] = []
    for line in text.splitlines():
        cut = None
        for idx, char in enumerate(line):
            if char != "%":
                continue
            backslashes = 0
            pos = idx - 1
            while pos >= 0 and line[pos] == "\\":
                backslashes += 1
                pos -= 1
            if backslashes % 2 == 0:
                cut = idx
                break
        cleaned.append(line if cut is None else line[:cut])
    return "\n".join(cleaned)


def safe_read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def compact_text(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def _mask_non_newlines(match: re.Match[str]) -> str:
    """Hide non-prose LaTeX while preserving line positions for findings."""
    return re.sub(r"[^\n]", " ", match.group(0))


def prose_for_language_audit(text: str) -> str:
    """Return conservative prose for heuristic readability candidate discovery.

    This is intentionally not a LaTeX parser. It removes the most common math,
    float and code regions so formulas and commands do not inflate sentence or
    connector counts. Semantic language findings still require human review.
    """
    prose = strip_latex_comments(text)
    environments = (
        "equation", "equation*", "align", "align*", "aligned", "gather", "gather*",
        "multline", "multline*", "figure", "figure*", "table", "table*", "tikzpicture",
        "lstlisting", "verbatim",
    )
    for environment in environments:
        pattern = rf"\\begin\{{{re.escape(environment)}\}}.*?\\end\{{{re.escape(environment)}\}}"
        prose = re.sub(pattern, _mask_non_newlines, prose, flags=re.S)
    for pattern in (
        r"\$\$.*?\$\$", r"\\\[.*?\\\]", r"\\\(.*?\\\)",
        r"(?<!\$)\$(?!\$).*?(?<!\$)\$(?!\$)",
    ):
        prose = re.sub(pattern, _mask_non_newlines, prose, flags=re.S)
    for _ in range(3):
        prose = re.sub(
            r"\\(?:textbf|textit|emph|underline|textrm|textsf|texttt)\s*\{([^{}]*)\}",
            lambda match: match.group(1),
            prose,
        )
    prose = CITATION_COMMAND_RE.sub(_mask_non_newlines, prose)
    prose = re.sub(
        r"\\(?:ref|eqref|label|includegraphics|url|href)\s*(?:\[[^]]*\])?\s*\{[^{}]*\}",
        _mask_non_newlines,
        prose,
    )
    prose = re.sub(r"\\[A-Za-z@]+\*?(?:\[[^]]*\])?", " ", prose)
    return prose.replace("{", " ").replace("}", " ")


def find_citation_commands(text: str) -> list[dict[str, Any]]:
    """Return real LaTeX citation commands and their keys.

    Numeric text such as ``[4,5]`` is deliberately not treated as a citation.
    """
    cleaned = strip_latex_comments(text)
    records: list[dict[str, Any]] = []
    for match in CITATION_COMMAND_RE.finditer(cleaned):
        command = match.group("command")
        if command.lower() == "nocite":
            continue
        keys = [
            key.strip() for key in match.group("keys").split(",")
            if key.strip() and "#" not in key and key.strip() != "*"
        ]
        if not keys:
            continue
        records.append(
            {
                "command": command,
                "keys": keys,
                "line": cleaned.count("\n", 0, match.start()) + 1,
                "start": match.start(),
                "end": match.end(),
            }
        )
    return records


def bib_entry_keys(text: str) -> list[str]:
    keys: list[str] = []
    for match in re.finditer(r"@(?P<kind>[A-Za-z]+)\s*\{\s*(?P<key>[^,\s=]+)", text):
        if match.group("kind").lower() in {"string", "comment", "preamble"}:
            continue
        keys.append(match.group("key"))
    return keys


def declared_bibliography_files(
    source_units: list[tuple[Path, str]], project_root: Path
) -> list[Path]:
    """Resolve bibliography files explicitly loaded by the final TeX source."""
    resolved: list[Path] = []
    project_root = project_root.resolve()
    for source_path, text in source_units:
        declarations: list[str] = []
        declarations.extend(
            name.strip()
            for group in re.findall(r"\\bibliography\s*\{([^{}]+)\}", text)
            for name in group.split(",")
            if name.strip()
        )
        declarations.extend(
            match.group(1).strip()
            for match in re.finditer(r"\\addbibresource(?:\s*\[[^\]]*\])?\s*\{([^{}]+)\}", text)
            if match.group(1).strip()
        )
        for name in declarations:
            candidate = source_path.parent / name
            if not candidate.suffix:
                candidate = candidate.with_suffix(".bib")
            if not candidate.exists():
                fallback = project_root / (name if Path(name).suffix else f"{name}.bib")
                candidate = fallback
            candidate = candidate.resolve()
            try:
                candidate.relative_to(project_root)
            except ValueError:
                continue
            if candidate.is_file() and candidate not in resolved:
                resolved.append(candidate)
    return resolved


def prose_for_reference_audit(text: str) -> str:
    """Mask non-prose and real citation commands while preserving line numbers."""
    prose = strip_latex_comments(text)
    environments = (
        "equation", "equation*", "align", "align*", "aligned", "gather", "gather*",
        "multline", "multline*", "figure", "figure*", "table", "table*", "tikzpicture",
        "lstlisting", "verbatim", "thebibliography",
    )
    for environment in environments:
        pattern = rf"\\begin\{{{re.escape(environment)}\}}.*?\\end\{{{re.escape(environment)}\}}"
        prose = re.sub(pattern, _mask_non_newlines, prose, flags=re.S)
    for pattern in (
        r"\$\$.*?\$\$", r"\\\[.*?\\\]", r"\\\(.*?\\\)",
        r"(?<!\$)\$(?!\$).*?(?<!\$)\$(?!\$)",
    ):
        prose = re.sub(pattern, _mask_non_newlines, prose, flags=re.S)
    prose = CITATION_COMMAND_RE.sub(_mask_non_newlines, prose)
    return prose


def manual_numeric_reference_tokens(text: str) -> list[dict[str, Any]]:
    """Locate numeric prose tokens that may be hand-typed citations.

    The result is evidence for review, not an automatic failure. In particular,
    intervals such as ``[4,5]`` remain non-citation tokens unless their context
    makes a citation claim or the final PDF links them to the bibliography.
    """
    prose = prose_for_reference_audit(text)
    pattern = re.compile(
        r"(?P<token>\[\s*\d{1,4}(?:\s*[-–—,，]\s*\d{1,4})*\s*\]"
        r"|[（(]\s*\d{1,3}(?:\s*[-–—,，]\s*\d{1,3})*\s*[）)])"
    )
    records: list[dict[str, Any]] = []
    for match in pattern.finditer(prose):
        before = prose[max(0, match.start() - 45):match.start()]
        after = prose[match.end():min(len(prose), match.end() + 25)]
        compact_context = re.sub(r"\s+", "", before + match.group("token") + after)
        before_clause = re.split(r"[。！？!?；;\n]", before)[-1]
        after_clause = re.split(r"[。！？!?；;\n]", after)[0]
        citation_like = any(
            term in before_clause[-30:] or term in after_clause[:12] for term in CITATION_CONTEXT_TERMS
        )
        records.append(
            {
                "token": re.sub(r"\s+", "", match.group("token")),
                "line": prose.count("\n", 0, match.start()) + 1,
                "context": compact_context[:150],
                "citation_like": citation_like,
            }
        )
    return records


def latex_indentation_evidence(text: str) -> dict[str, list[dict[str, Any]]]:
    """Collect source-level paragraph indentation evidence without visual claims."""
    cleaned = strip_latex_comments(text)
    settings: list[dict[str, Any]] = []
    setting_patterns = (
        r"\\setlength\s*\{\s*\\parindent\s*\}\s*\{\s*([^{}]+?)\s*\}",
        r"\\parindent\s*=\s*([^\s%]+)",
    )
    for pattern in setting_patterns:
        for match in re.finditer(pattern, cleaned):
            settings.append(
                {
                    "value": match.group(1).strip(),
                    "line": cleaned.count("\n", 0, match.start()) + 1,
                }
            )

    noindent: list[dict[str, Any]] = []
    for match in re.finditer(r"\\noindent\b", cleaned):
        line_start = cleaned.rfind("\n", 0, match.start()) + 1
        line_end = cleaned.find("\n", match.end())
        line_end = len(cleaned) if line_end == -1 else line_end
        context = cleaned[line_start:line_end].strip()
        if re.search(r"keywords?|关键词|关键字|title|author|date", context, re.I):
            continue
        noindent.append(
            {
                "line": cleaned.count("\n", 0, match.start()) + 1,
                "context": context[:160],
            }
        )

    manual_spacing: list[dict[str, Any]] = []
    manual_pattern = re.compile(
        r"(?m)^[ \t]*(?P<token>　{1,3}|~{2,}|(?:\\quad\s*){2,}"
        r"|\\hspace\*?\s*\{\s*(?:1\.5|2(?:\.0)?)\s*(?:em|zw|\\ccwd)\s*\})"
    )
    for match in manual_pattern.finditer(cleaned):
        manual_spacing.append(
            {
                "token": match.group("token"),
                "line": cleaned.count("\n", 0, match.start()) + 1,
            }
        )
    return {"settings": settings, "noindent": noindent, "manual_spacing": manual_spacing}


def collect_latex_sources(main_source: Path, project_root: Path) -> list[tuple[Path, str]]:
    """Read the main TeX file and its local input/include/subfile closure."""
    collected: list[tuple[Path, str]] = []
    seen: set[Path] = set()

    def visit(path: Path) -> None:
        resolved = path.resolve()
        if resolved in seen or not resolved.is_file() or resolved.suffix.lower() != ".tex":
            return
        try:
            resolved.relative_to(project_root.resolve())
        except ValueError:
            return
        seen.add(resolved)
        raw = safe_read_text(resolved)
        cleaned = strip_latex_comments(raw)
        collected.append((resolved, cleaned))
        for match in re.finditer(r"\\(?:input|include|subfile)\s*\{([^{}]+)\}", cleaned):
            name = match.group(1).strip()
            candidate = resolved.parent / name
            if not candidate.suffix:
                candidate = candidate.with_suffix(".tex")
            visit(candidate)

    visit(main_source)
    return collected


def inspect_pdf_reference_links(path: Path) -> dict[str, Any] | None:
    """Inspect internal PDF links whose destination is in the bibliography pages."""
    try:
        import fitz
    except ImportError:
        return None
    try:
        document = fitz.open(path)
    except Exception:
        return None
    try:
        reference_start: int | None = None
        for page_index in range(len(document)):
            page_text = document[page_index].get_text("text")
            if re.search(r"(?m)^\s*(?:参\s*考\s*文\s*献|References?)\s*$", page_text, re.I):
                reference_start = page_index
                break
        records: list[dict[str, Any]] = []
        if reference_start is not None:
            for page_index in range(reference_start):
                page = document[page_index]
                for link in page.get_links():
                    target_page = link.get("page", -1)
                    if not isinstance(target_page, int) or target_page < reference_start:
                        continue
                    rect = fitz.Rect(link.get("from"))
                    expanded = fitz.Rect(rect.x0 - 1, rect.y0 - 1, rect.x1 + 1, rect.y1 + 1)
                    display = re.sub(r"\s+", "", page.get_textbox(expanded))
                    records.append(
                        {
                            "source_page": page_index + 1,
                            "target_page": target_page + 1,
                            "display": display,
                        }
                    )
        return {
            "reference_start_page": None if reference_start is None else reference_start + 1,
            "links": records,
        }
    finally:
        document.close()


def readability_candidates(text: str) -> dict[str, list[dict[str, Any]]]:
    """Locate readability candidates without deciding that the prose fails."""
    prose = prose_for_language_audit(text)
    candidates: dict[str, list[dict[str, Any]]] = {
        "long_sentences": [],
        "dense_connectors": [],
        "abstract_stacks": [],
    }
    for match in re.finditer(r"[^。！？!?；;\n]+[。！？!?；;]?", prose):
        sentence = re.sub(r"\s+", " ", match.group(0)).strip()
        if not sentence:
            continue
        content_length = len(re.findall(r"[\u4e00-\u9fffA-Za-z0-9]", sentence))
        connector_hits = [term for term in LOGIC_CONNECTORS for _ in re.finditer(re.escape(term), sentence)]
        abstract_hits = [term for term in ABSTRACT_NOUN_TERMS for _ in re.finditer(re.escape(term), sentence)]
        record = {
            "line": prose.count("\n", 0, match.start()) + 1,
            "length": content_length,
            "connector_count": len(connector_hits),
            "abstract_noun_count": len(abstract_hits),
            "excerpt": sentence[:180],
        }
        if content_length >= 115 or (content_length >= 85 and len(connector_hits) >= 4):
            candidates["long_sentences"].append(record)
        if content_length >= 55 and len(connector_hits) >= 5:
            candidates["dense_connectors"].append(record)
        if content_length >= 55 and len(abstract_hits) >= 8:
            candidates["abstract_stacks"].append(record)
    return candidates


def normalized_terminology_groups(config: Any = None) -> list[tuple[str, tuple[str, ...]]]:
    """Merge conservative defaults with project-specific terminology groups."""
    groups: dict[str, list[str]] = {
        canonical: list(variants) for canonical, variants in DEFAULT_TERMINOLOGY_GROUPS
    }
    items: list[dict[str, Any]] = []
    if isinstance(config, dict):
        items = [
            {"canonical": canonical, "variants": variants}
            for canonical, variants in config.items()
        ]
    elif isinstance(config, list):
        items = [item for item in config if isinstance(item, dict)]

    for item in items:
        canonical = str(item.get("canonical", "")).strip()
        variants_value = item.get("variants", [])
        if not canonical or not isinstance(variants_value, list):
            continue
        variants = [value.strip() for value in variants_value if isinstance(value, str) and value.strip()]
        if canonical not in variants:
            variants.insert(0, canonical)
        groups[canonical] = list(dict.fromkeys(variants))

    result: list[tuple[str, tuple[str, ...]]] = []
    seen_variant_sets: set[frozenset[str]] = set()
    for canonical, variants in groups.items():
        unique_variants = tuple(dict.fromkeys(variant for variant in variants if variant))
        variant_set = frozenset(unique_variants)
        if len(unique_variants) < 2 or variant_set in seen_variant_sets:
            continue
        seen_variant_sets.add(variant_set)
        result.append((canonical, unique_variants))
    return result


def terminology_consistency_candidates(
    text: str, config: Any = None
) -> list[dict[str, Any]]:
    """Locate co-occurring terminology variants without deciding synonymy."""
    prose = prose_for_language_audit(text)
    records: list[dict[str, Any]] = []
    for canonical, variants in normalized_terminology_groups(config):
        occurrences: dict[str, list[int]] = {}
        for variant in variants:
            lines = [
                prose.count("\n", 0, match.start()) + 1
                for match in re.finditer(re.escape(variant), prose)
            ]
            if lines:
                occurrences[variant] = lines
        if len(occurrences) >= 2:
            records.append(
                {
                    "canonical": canonical,
                    "variants": list(occurrences),
                    "lines": occurrences,
                }
            )
    return records


def showcase_citation_candidates(text: str) -> list[dict[str, Any]]:
    """Locate prose that displays a reference instead of attaching it to a claim."""
    prose = prose_for_language_audit(text)
    records: list[dict[str, Any]] = []
    for match in SHOWCASE_CITATION_RE.finditer(prose):
        line_start = prose.rfind("\n", 0, match.start()) + 1
        line_end = prose.find("\n", match.end())
        line_end = len(prose) if line_end == -1 else line_end
        excerpt = re.sub(r"\s+", " ", prose[line_start:line_end]).strip()
        records.append(
            {
                "line": prose.count("\n", 0, match.start()) + 1,
                "phrase": re.sub(r"\s+", "", match.group(0)),
                "excerpt": excerpt[:180],
            }
        )
    return records


def drawing_placeholder_candidates(text: str) -> list[dict[str, Any]]:
    """Locate explicit unfinished figure/table markers while ignoring comments."""
    cleaned = strip_latex_comments(text)
    records: list[dict[str, Any]] = []
    for match in DRAWING_PLACEHOLDER_RE.finditer(cleaned):
        line_start = cleaned.rfind("\n", 0, match.start()) + 1
        line_end = cleaned.find("\n", match.end())
        line_end = len(cleaned) if line_end == -1 else line_end
        records.append(
            {
                "line": cleaned.count("\n", 0, match.start()) + 1,
                "marker": re.sub(r"\s+", " ", match.group(0)).strip(),
                "excerpt": re.sub(r"\s+", " ", cleaned[line_start:line_end]).strip()[:180],
            }
        )
    return records


def source_attribution_candidates(text: str) -> list[dict[str, Any]]:
    """Locate source claims about model parameters or constructed relations.

    These are review candidates only. Whether a claim falsely attributes a
    team-defined value to the problem requires comparison with the problem,
    attachments, or a frozen interpretation.
    """
    cleaned = strip_latex_comments(text)
    records: list[dict[str, Any]] = []
    seen_sentences: set[tuple[int, int]] = set()
    for match in SOURCE_ATTRIBUTION_RE.finditer(cleaned):
        start_candidates = [cleaned.rfind(mark, 0, match.start()) for mark in "。！？；;\n"]
        sentence_start = max(start_candidates) + 1
        end_candidates = [
            position for mark in "。！？；;\n"
            if (position := cleaned.find(mark, match.end())) != -1
        ]
        sentence_end = min(end_candidates) + 1 if end_candidates else len(cleaned)
        span = (sentence_start, sentence_end)
        if span in seen_sentences:
            continue
        seen_sentences.add(span)
        excerpt = re.sub(r"\s+", " ", cleaned[sentence_start:sentence_end]).strip()
        math_tokens: list[str] = []
        for groups in re.findall(r"\$([^$]{1,40})\$|\\\(([^)]{1,40})\\\)", excerpt):
            math_tokens.extend(re.sub(r"\s+", "", token) for token in groups if token)
        math_tokens.extend(re.sub(r"\s+", "", token) for token in re.findall(r"`([^`]{1,40})`", excerpt))
        math_tokens.extend(re.findall(r"\b[A-Za-z][A-Za-z0-9]*_[A-Za-z0-9]+\b", excerpt))
        number_tokens = re.findall(r"\d+(?:\.\d+)?\s*(?:%|条|个|类|组|阶|年|天|周|月)?", excerpt)
        term_tokens = [term for term in SOURCE_SENSITIVE_TERMS if term in excerpt]
        if not math_tokens and not term_tokens:
            continue
        objects = list(dict.fromkeys(math_tokens + number_tokens + term_tokens))[:8]
        records.append(
            {
                "line": cleaned.count("\n", 0, sentence_start) + 1,
                "claim": re.sub(r"\s+", " ", match.group(0)).strip(),
                "objects": objects,
                "excerpt": excerpt[:220],
            }
        )
    return records


def text_is_extractable(text: str | None, minimum: int = 30) -> bool:
    if text is None:
        return False
    meaningful = re.sub(r"[\s\f\W_]+", "", text, flags=re.UNICODE)
    return len(meaningful) >= minimum


def find_abstract_evidence(text: str) -> tuple[bool, str]:
    patterns = [
        (r"\\begin\{(?:abstract|cnabstract|zhabstract|cabstract)\*?\}", "摘要环境"),
        (r"\\(?:zhabstract|cnabstract|cabstract|abstracttext|makeabstract)\s*\{", "摘要命令"),
        (r"\\(?:section|chapter|heading|biaoti)\*?\s*\{\s*摘\s*要\s*\}", "摘要标题命令"),
        (r"(?:^|\n)\s*(?:\\textbf\s*\{\s*)?摘\s*要(?:\s*\})?\s*(?:[:：]|\n)", "摘要文本标题"),
    ]
    for pattern, label in patterns:
        if re.search(pattern, text, re.I | re.M):
            return True, label
    return False, ""


def find_keyword_payloads(text: str) -> list[str]:
    payloads: list[str] = []
    command_patterns = [
        r"\\(?:keywords?|keywordscn|ckeywords|zhkeywords|keywordcn|KeyWords)\s*\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}",
        r"\\(?:keywordline|makekeywords)\s*\{([^{}]+)\}",
    ]
    for pattern in command_patterns:
        payloads.extend(match.group(1) for match in re.finditer(pattern, text, re.I | re.S))
    label_pattern = re.compile(
        r"(?:关键词|关键字|Key\s*words?)\s*[:：]\s*([^\n\\]{2,180})",
        re.I,
    )
    payloads.extend(match.group(1).strip() for match in label_pattern.finditer(text))
    return payloads


def count_keywords(payload: str) -> int | None:
    cleaned = re.sub(r"\\(?:quad|qquad|enspace|hspace\*?\{[^}]*\})", "；", payload)
    cleaned = re.sub(r"\\[a-zA-Z]+\s*", "", cleaned)
    parts = [part.strip(" \t{}。.") for part in re.split(r"[，,；;、|]+", cleaned)]
    parts = [part for part in parts if part]
    if len(parts) == 1 and " " in parts[0]:
        ascii_parts = [part for part in re.split(r"\s{2,}|\s+", parts[0]) if part]
        if len(ascii_parts) > 1:
            parts = ascii_parts
    return len(parts) if parts else None


def ooxml_metadata(path: Path) -> dict[str, str]:
    metadata: dict[str, str] = {}
    try:
        with zipfile.ZipFile(path) as archive:
            for member in ("docProps/core.xml", "docProps/app.xml", "docProps/custom.xml"):
                if member not in archive.namelist():
                    continue
                try:
                    root = ET.fromstring(archive.read(member))
                except ET.ParseError:
                    continue
                for node in root.iter():
                    value = (node.text or "").strip()
                    if value:
                        key = node.tag.rsplit("}", 1)[-1]
                        metadata[f"{member}:{key}"] = value
    except (OSError, zipfile.BadZipFile):
        pass
    return metadata


def pdf_metadata(path: Path) -> dict[str, str]:
    if not shutil.which("pdfinfo"):
        return {}
    result = run_command(["pdfinfo", str(path)], timeout=60)
    if result.returncode != 0:
        return {}
    metadata: dict[str, str] = {}
    for line in result.stdout.splitlines():
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        if key.strip() in {"Title", "Subject", "Keywords", "Author", "Creator", "Producer"} and value.strip():
            metadata[f"pdf:{key.strip()}"] = value.strip()
    return metadata


def image_metadata(path: Path) -> dict[str, str]:
    try:
        from PIL import ExifTags, Image
    except ImportError:
        return {}
    metadata: dict[str, str] = {}
    try:
        with Image.open(path) as image:
            for key, value in image.info.items():
                if isinstance(value, (str, int, float)) and str(value).strip():
                    metadata[f"image:{key}"] = str(value)
            exif = image.getexif()
            for key, value in exif.items():
                name = ExifTags.TAGS.get(key, str(key))
                if isinstance(value, bytes):
                    value = value.decode("utf-8", errors="ignore")
                if isinstance(value, (str, int, float)) and str(value).strip():
                    metadata[f"exif:{name}"] = str(value)
    except (OSError, ValueError):
        pass
    return metadata


def file_metadata(path: Path) -> dict[str, str]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return pdf_metadata(path)
    if suffix in {".docx", ".xlsx", ".xlsm", ".pptx"}:
        return ooxml_metadata(path)
    if suffix in {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp"}:
        return image_metadata(path)
    return {}


def run_command(args: list[str], cwd: Path | None = None, timeout: int = 60) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        args,
        cwd=str(cwd) if cwd else None,
        text=True,
        encoding="utf-8",
        errors="replace",
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        timeout=timeout,
        shell=False,
    )


def pdf_text(path: Path, first_page_only: bool = False) -> str | None:
    if not shutil.which("pdftotext"):
        return None
    cmd = ["pdftotext"]
    if first_page_only:
        cmd += ["-f", "1", "-l", "1"]
    cmd += ["-layout", str(path), "-"]
    result = run_command(cmd, timeout=120)
    return result.stdout if result.returncode == 0 else None


def pdf_page_count(path: Path) -> int | None:
    if not shutil.which("pdfinfo"):
        return None
    result = run_command(["pdfinfo", str(path)])
    match = re.search(r"^Pages:\s+(\d+)", result.stdout, re.M)
    return int(match.group(1)) if match else None


def pdf_font_names(path: Path) -> list[str] | None:
    if not shutil.which("pdffonts"):
        return None
    result = run_command(["pdffonts", str(path)], timeout=120)
    if result.returncode != 0:
        return None
    lines = result.stdout.splitlines()[2:]
    names = []
    for line in lines:
        if line.strip():
            names.append(line.split()[0])
    return sorted(set(names))


def docx_text(path: Path) -> str | None:
    try:
        with zipfile.ZipFile(path) as archive:
            data = archive.read("word/document.xml")
    except (OSError, KeyError, zipfile.BadZipFile):
        return None
    root = ET.fromstring(data)
    return "\n".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))


def xlsx_text(path: Path) -> str:
    values: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name == "xl/sharedStrings.xml" or name.startswith("xl/worksheets/"):
                    try:
                        root = ET.fromstring(archive.read(name))
                    except ET.ParseError:
                        continue
                    for node in root.iter():
                        if node.tag.endswith("}t") and node.text:
                            values.append(node.text)
    except (OSError, zipfile.BadZipFile):
        return ""
    return "\n".join(values)


def column_number(cell_ref: str) -> int:
    letters = re.match(r"[A-Za-z]+", cell_ref)
    if not letters:
        return 0
    value = 0
    for char in letters.group(0).upper():
        value = value * 26 + ord(char) - 64
    return value


def xlsx_summary(path: Path) -> dict[str, dict[str, int]]:
    result: dict[str, dict[str, int]] = {}
    try:
        with zipfile.ZipFile(path) as archive:
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            relationship_map = {
                node.attrib.get("Id", ""): node.attrib.get("Target", "")
                for node in rels.iter()
                if node.tag.endswith("}Relationship")
            }
            relationship_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
            for sheet in workbook.iter():
                if not sheet.tag.endswith("}sheet"):
                    continue
                sheet_name = sheet.attrib.get("name", "")
                target = relationship_map.get(sheet.attrib.get(relationship_ns, ""), "")
                if not target:
                    continue
                member = target.lstrip("/")
                if not member.startswith("xl/"):
                    member = "xl/" + member
                member = str(Path(member))
                xml = ET.fromstring(archive.read(member))
                row_count = 0
                max_column = 0
                for row in xml.iter():
                    if row.tag.endswith("}row"):
                        row_count += 1
                    elif row.tag.endswith("}c"):
                        max_column = max(max_column, column_number(row.attrib.get("r", "")))
                result[sheet_name] = {"rows": row_count, "columns": max_column}
    except (OSError, KeyError, ET.ParseError, zipfile.BadZipFile):
        return {}
    return result


def csv_summary(path: Path) -> tuple[int, int] | None:
    try:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
            rows = list(csv.reader(handle, delimiter=delimiter))
    except OSError:
        return None
    return len(rows), max((len(row) for row in rows), default=0)


def parse_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if not isinstance(value, str):
        return None
    cleaned = value.strip().replace(",", "").replace("，", "")
    percent = cleaned.endswith("%")
    cleaned = cleaned.rstrip("%").strip()
    match = re.fullmatch(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?", cleaned)
    if not match:
        return None
    number = float(cleaned)
    return number / 100.0 if percent else number


def values_close(left: Any, right: Any, atol: float, rtol: float) -> bool:
    left_number = parse_number(left)
    right_number = parse_number(right)
    if left_number is not None and right_number is not None:
        return math.isclose(left_number, right_number, abs_tol=atol, rel_tol=rtol)
    return str(left).strip() == str(right).strip()


def read_table(path: Path, sheet: str | None = None) -> list[list[Any]] | None:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        try:
            with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
                return [list(row) for row in csv.reader(handle, delimiter=delimiter)]
        except OSError:
            return None
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return None
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            worksheet = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            workbook.close()
            return rows
        except (OSError, KeyError, ValueError, zipfile.BadZipFile):
            return None
    return None


def compare_tables(
    generated: Path,
    frozen: Path,
    *,
    sheet: str | None,
    atol: float,
    rtol: float,
    key_columns: list[str] | None,
    ignore_columns: list[str] | None,
) -> dict[str, Any] | None:
    left = read_table(generated, sheet)
    right = read_table(frozen, sheet)
    if left is None or right is None:
        return None
    result: dict[str, Any] = {
        "equal": True,
        "generated_shape": [len(left), max((len(row) for row in left), default=0)],
        "frozen_shape": [len(right), max((len(row) for row in right), default=0)],
        "mismatches": [],
    }
    if key_columns and left and right:
        left_header = [str(value).strip() for value in left[0]]
        right_header = [str(value).strip() for value in right[0]]
        if left_header != right_header:
            result["equal"] = False
            result["mismatches"].append("表头不一致")
            return result
        try:
            key_indexes = [left_header.index(name) for name in key_columns]
        except ValueError:
            result["equal"] = False
            result["mismatches"].append("key_columns 在表头中不存在")
            return result
        ignored = {left_header.index(name) for name in (ignore_columns or []) if name in left_header}
        left_map = {tuple(row[index] if index < len(row) else None for index in key_indexes): row for row in left[1:]}
        right_map = {tuple(row[index] if index < len(row) else None for index in key_indexes): row for row in right[1:]}
        if set(left_map) != set(right_map):
            result["equal"] = False
            missing = sorted(set(right_map) - set(left_map), key=str)[:5]
            extra = sorted(set(left_map) - set(right_map), key=str)[:5]
            result["mismatches"].append(f"键集合不一致，缺少={missing}，新增={extra}")
        row_pairs = [(key, left_map[key], right_map[key]) for key in left_map.keys() & right_map.keys()]
    else:
        ignored = set()
        if len(left) != len(right):
            result["equal"] = False
            result["mismatches"].append(f"行数不一致：{len(left)} != {len(right)}")
        row_pairs = [(index + 1, a, b) for index, (a, b) in enumerate(zip(left, right))]
    for row_key, left_row, right_row in row_pairs:
        if len(left_row) != len(right_row):
            result["equal"] = False
            result["mismatches"].append(f"行 {row_key} 列数不一致：{len(left_row)} != {len(right_row)}")
            if len(result["mismatches"]) >= 10:
                break
        for column, (left_value, right_value) in enumerate(zip(left_row, right_row)):
            if column in ignored:
                continue
            if not values_close(left_value, right_value, atol, rtol):
                result["equal"] = False
                result["mismatches"].append(
                    f"行 {row_key} 列 {column + 1}: {left_value!r} != {right_value!r}"
                )
                if len(result["mismatches"]) >= 10:
                    break
        if len(result["mismatches"]) >= 10:
            break
    return result


def json_path_value(payload: Any, path: str) -> Any:
    value = payload
    for part in path.split(".") if path else []:
        if isinstance(value, list):
            value = value[int(part)]
        elif isinstance(value, dict):
            value = value[part]
        else:
            raise KeyError(part)
    return value


def structured_value(root: Path, spec: dict[str, Any]) -> tuple[Any, str] | None:
    raw_path = spec.get("path")
    if not isinstance(raw_path, str) or not raw_path:
        return None
    path = Path(raw_path)
    path = path if path.is_absolute() else root / path
    if not path.exists():
        return None
    suffix = path.suffix.lower()
    try:
        if suffix == ".json":
            value = json_path_value(json.loads(safe_read_text(path)), str(spec.get("json_path", "")))
            return value, str(path)
        if suffix in {".csv", ".tsv"}:
            delimiter = "\t" if suffix == ".tsv" else ","
            with path.open("r", encoding="utf-8-sig", errors="ignore", newline="") as handle:
                rows = list(csv.DictReader(handle, delimiter=delimiter))
            row_index = int(spec.get("row", 0))
            column = str(spec["column"])
            return rows[row_index][column], f"{path}:row={row_index},column={column}"
        if suffix in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
            except ImportError:
                return None
            workbook = load_workbook(path, read_only=True, data_only=True)
            worksheet = workbook[str(spec.get("sheet"))] if spec.get("sheet") else workbook[workbook.sheetnames[0]]
            cell = str(spec.get("cell", "A1"))
            value = worksheet[cell].value
            workbook.close()
            return value, f"{path}:{worksheet.title}!{cell}"
    except (OSError, KeyError, IndexError, ValueError, json.JSONDecodeError):
        return None
    return None


def is_hidden(path: Path, root: Path) -> bool:
    try:
        relative = path.relative_to(root)
    except ValueError:
        return False
    return any(part.startswith(".") for part in relative.parts)


def list_project_files(root: Path) -> list[Path]:
    return sorted(
        path for path in root.rglob("*")
        if path.is_file() and not is_hidden(path, root)
    )


def relative(path: Path | None, root: Path) -> str:
    if path is None:
        return "—"
    try:
        return str(path.relative_to(root))
    except ValueError:
        return str(path)


class SubmissionAudit:
    def __init__(self, args: argparse.Namespace) -> None:
        self.args = args
        self.root = Path(args.project).resolve()
        if args.manifest:
            manifest_candidate = Path(args.manifest)
            self.manifest_path = (
                manifest_candidate.resolve()
                if manifest_candidate.is_absolute()
                else (self.root / manifest_candidate).resolve()
            )
        else:
            self.manifest_path = self.root / "audit_manifest.json"
        self.manifest: dict[str, Any] = {}
        self.findings: list[Finding] = []
        self.passes: list[dict[str, str]] = []
        self.unverified: list[dict[str, str]] = []
        self.deferred: list[dict[str, str]] = []
        self.not_ready = False
        self.scope_incomplete = False
        self.files: list[Path] = []
        self.paper: Path | None = None
        self.source: Path | None = None
        self.paper_text = ""
        self.source_text = ""
        self.source_units: list[tuple[Path, str]] = []
        self.paper_text_extractable = False
        self.artifact_snapshots: dict[int, bytes] = {}
        self.entrypoint_succeeded = False
        self.visual_output: dict[str, Any] = {}
        self.source_attributions: list[dict[str, Any]] = []
        if args.mode == "paper-only":
            for check_id, item, reason in PAPER_ONLY_DEFAULT_DEFERRALS:
                self.defer(check_id, "DEFERRED_TO_FULL", item, reason, "FULL")

    def add(
        self,
        priority: str,
        check_id: str,
        title: str,
        location: str,
        evidence: str,
        rule: str,
        recommendation: str,
        auto_fixable: bool = False,
        certainty: str = "DETERMINISTIC",
        stage_status: str = "CONFIRMED_IN_SCOPE",
    ) -> None:
        self.findings.append(
            Finding(
                priority, check_id, title, location, evidence, rule, recommendation,
                auto_fixable, certainty, stage_status,
            )
        )

    def warn(
        self,
        priority: str,
        check_id: str,
        title: str,
        location: str,
        evidence: str,
        rule: str,
        recommendation: str,
    ) -> None:
        self.add(priority, check_id, title, location, evidence, rule, recommendation, False, "HEURISTIC")

    def passed(self, check_id: str, message: str) -> None:
        self.passes.append({"check_id": check_id, "message": message})

    def mark_unverified(self, check_id: str, message: str) -> None:
        self.unverified.append({"check_id": check_id, "message": message})

    def defer(
        self,
        check_id: str,
        status: str,
        item: str,
        reason: str,
        next_mode: str,
    ) -> None:
        record = {
            "check_id": check_id,
            "status": status,
            "item": item,
            "reason": reason,
            "next_mode": next_mode,
        }
        for index, existing in enumerate(self.deferred):
            if existing["check_id"] == check_id:
                self.deferred[index] = record
                return
        self.deferred.append(record)

    def resolve_manifest_path(self, value: Any) -> Path | None:
        if not isinstance(value, str) or not value.strip():
            return None
        candidate = Path(value)
        return candidate.resolve() if candidate.is_absolute() else (self.root / candidate).resolve()

    def load_manifest(self) -> None:
        if self.manifest_path.exists():
            try:
                self.manifest = json.loads(safe_read_text(self.manifest_path))
                self.passed("SCOPE-001", f"已读取审查清单 {relative(self.manifest_path, self.root)}")
            except (OSError, json.JSONDecodeError) as exc:
                if self.args.mode == "paper-only":
                    self.defer(
                        "SCOPE-FULL-005", "DEFERRED_TO_FULL", "audit_manifest.json 解析",
                        f"PAPER_ONLY 不依赖交付清单；当前解析错误：{exc}", "FULL",
                    )
                else:
                    self.add("P1", "SCOPE-002", "审查清单无法解析", relative(self.manifest_path, self.root), str(exc),
                             "FULL 审查需要可靠的冻结和交付清单。", "修复 audit_manifest.json 后重新运行。")
                    self.scope_incomplete = True
        else:
            if self.args.mode == "paper-only":
                self.defer(
                    "SCOPE-FULL-005", "OUT_OF_SCOPE", "audit_manifest.json",
                    "PAPER_ONLY 不要求冻结与交付清单；FULL 再核验。", "FULL",
                )
            else:
                self.mark_unverified("SCOPE-003", "未提供 audit_manifest.json；将使用文件发现和冻结文档作有限判断。")

    def check_readiness(self) -> None:
        if self.args.mode not in {"full", "recheck"}:
            return
        try:
            self.manifest_path.relative_to(self.root)
        except ValueError:
            self.add(
                "P0", "READY-007", "冻结清单不在 paper_final 根目录内",
                str(self.manifest_path), "manifest path escapes project root",
                "FULL/RECHECK 必须只读当前冻结包内部的 audit_manifest.json。",
                "把审查根目录指向 paper_final，并使用其中的 audit_manifest.json。",
            )
            self.not_ready = True
            return
        if self.manifest:
            blockers = self.manifest.get("blockers", [])
            if self.manifest.get("frozen") is not True or blockers:
                evidence = f"frozen={self.manifest.get('frozen')!r}, blockers={blockers!r}"
                self.add("P0", "READY-001", "未通过 Final Audit 冻结门禁", relative(self.manifest_path, self.root), evidence,
                         "核心模型、结果、代码和图表冻结后才能执行最终审查。", "返回写作/建模阶段，清零 BLOCKER 并冻结最终版本。")
                self.not_ready = True
            else:
                self.passed("READY-002", "清单确认项目已冻结且无 BLOCKER")
                self.audit_freeze_provenance()
        else:
            self.add(
                "P0", "READY-003", "FULL/RECHECK 缺少最终冻结清单", str(self.root),
                "未读取 audit_manifest.json",
                "FULL/RECHECK 只审核最终成稿冻结阶段生成的版本。",
                "先生成新的 paper_final 冻结包，再从该目录执行审查。",
            )
            self.not_ready = True

    def audit_freeze_provenance(self) -> None:
        """Verify the final-freeze interface without changing audit semantics."""
        issues: list[str] = []
        if self.manifest.get("freeze_schema_version") != 1:
            issues.append("freeze_schema_version must be 1")
        standard = self.manifest.get("standard")
        if not isinstance(standard, dict) or not standard.get("version") or not standard.get("sha256"):
            issues.append("standard version/hash is missing")

        frozen_files = self.manifest.get("frozen_files")
        frozen_index: dict[str, str] = {}
        if not isinstance(frozen_files, list) or not frozen_files:
            issues.append("frozen_files is missing or empty")
        else:
            for index, item in enumerate(frozen_files):
                if not isinstance(item, dict):
                    issues.append(f"frozen_files[{index}] is not an object")
                    continue
                raw_path = item.get("path")
                expected = item.get("sha256")
                if not isinstance(raw_path, str) or not raw_path.strip() or Path(raw_path).is_absolute():
                    issues.append(f"frozen_files[{index}] path is not project-relative")
                    continue
                path = (self.root / raw_path).resolve()
                try:
                    relative_path = path.relative_to(self.root).as_posix()
                except ValueError:
                    issues.append(f"frozen_files[{index}] path escapes the frozen root")
                    continue
                if relative_path in frozen_index:
                    issues.append(f"duplicate frozen_files path: {relative_path}")
                    continue
                if not path.is_file():
                    issues.append(f"frozen file is missing: {relative_path}")
                    continue
                actual = sha256_file(path)
                if not isinstance(expected, str) or expected.lower() != actual:
                    issues.append(f"frozen file SHA-256 mismatch: {relative_path}")
                    continue
                frozen_index[relative_path] = actual

        upstream = self.manifest.get("upstream_manifests")
        upstream_payloads: dict[str, dict[str, Any]] = {}
        if not isinstance(upstream, dict):
            issues.append("upstream_manifests is missing")
        else:
            for key, expected_status in (("paper_ready", "PAPER_READY"), ("figures", "FIGURES_READY")):
                item = upstream.get(key)
                if not isinstance(item, dict):
                    issues.append(f"upstream_manifests.{key} is missing")
                    continue
                raw_path = item.get("path")
                expected_hash = item.get("sha256")
                if not isinstance(raw_path, str) or Path(raw_path).is_absolute():
                    issues.append(f"upstream_manifests.{key}.path is invalid")
                    continue
                path = (self.root / raw_path).resolve()
                try:
                    relative_path = path.relative_to(self.root).as_posix()
                except ValueError:
                    issues.append(f"upstream_manifests.{key}.path escapes the frozen root")
                    continue
                if not path.is_file():
                    issues.append(f"upstream manifest is missing: {relative_path}")
                    continue
                actual = sha256_file(path)
                if not isinstance(expected_hash, str) or expected_hash.lower() != actual:
                    issues.append(f"upstream manifest SHA-256 mismatch: {relative_path}")
                    continue
                if frozen_index.get(relative_path) != actual:
                    issues.append(f"upstream manifest is not covered by frozen_files: {relative_path}")
                try:
                    payload = json.loads(safe_read_text(path))
                except (OSError, json.JSONDecodeError) as exc:
                    issues.append(f"upstream manifest cannot be parsed: {relative_path}: {exc}")
                    continue
                if not isinstance(payload, dict) or payload.get("status") != expected_status:
                    issues.append(f"upstream manifest status is not {expected_status}: {relative_path}")
                    continue
                expected_schema = 1 if key == "paper_ready" else 2
                if payload.get("schema_version") != expected_schema:
                    issues.append(
                        f"upstream manifest schema_version is not {expected_schema}: {relative_path}"
                    )
                    continue
                if payload.get("standard") != standard:
                    issues.append(f"upstream manifest standard is inconsistent: {relative_path}")
                upstream_payloads[key] = payload

        paper_ready = upstream_payloads.get("paper_ready")
        figures = upstream_payloads.get("figures")
        if paper_ready and figures:
            required_roles = {
                "interpretation_final",
                "model_plan_final",
                "ambiguity_decisions",
                "granularity_contract",
                "hard_constraints",
                "interpretation_decisions",
                "model_assumptions",
            }
            paper_files = paper_ready.get("files", [])
            actual_roles = {
                item.get("role") for item in paper_files if isinstance(item, dict)
            }
            if actual_roles != required_roles:
                issues.append("paper_ready_manifest does not contain the two finals and five contracts")

            assets = figures.get("assets", [])
            if not isinstance(assets, list) or not assets:
                issues.append("figure_manifest assets is empty")
                assets = []
            for index, asset in enumerate(assets):
                if not isinstance(asset, dict) or asset.get("status") != "FINAL":
                    issues.append(f"figure asset is not FINAL: index={index}")
                    continue
                asset_type = asset.get("asset_type")
                required_qa = (
                    ("density", "evidence_richness", "portfolio", "a4_visual")
                    if asset_type == "figure"
                    else ("logic_match", "editability", "a4_visual")
                )
                qa = asset.get("qa")
                if asset_type not in {"figure", "flowchart"} or not isinstance(qa, dict):
                    issues.append(f"figure asset type/QA is invalid: index={index}")
                elif any(qa.get(key) != "PASS" for key in required_qa):
                    issues.append(f"figure asset has unresolved QA: index={index}")
                if asset_type == "flowchart" and asset.get("human_approved") is not True:
                    issues.append(f"flowchart lacks human approval: index={index}")
                brief_contract = asset.get("brief_contract")
                if (
                    not isinstance(brief_contract, dict)
                    or brief_contract.get("schema_version") != 2
                    or not re.fullmatch(r"[0-9a-f]{64}", str(brief_contract.get("semantic_sha256", "")))
                    or not isinstance(brief_contract.get("body_context"), dict)
                ):
                    issues.append(f"figure asset lacks strict Brief/body snapshot: index={index}")
                if asset_type == "figure" and not isinstance(asset.get("script"), dict):
                    issues.append(f"figure asset lacks drawing script: index={index}")
                if asset_type == "flowchart":
                    source = asset.get("source")
                    if (
                        not isinstance(source, dict)
                        or not isinstance(source.get("path"), str)
                        or Path(source["path"]).suffix.lower() != ".pptx"
                    ):
                        issues.append(f"flowchart lacks editable PPTX source: index={index}")

            paper_ready_item = upstream.get("paper_ready", {})
            paper_ready_path = self.resolve_manifest_path(paper_ready_item.get("path"))
            if paper_ready_path and figures.get("paper_ready_manifest_sha256") != sha256_file(paper_ready_path):
                issues.append("figure_manifest points to a stale paper_ready_manifest")

            referenced_files: list[tuple[str, Any]] = []
            for item in paper_ready.get("files", []):
                if isinstance(item, dict) and isinstance(item.get("path"), str):
                    referenced_files.append((item["path"], item.get("sha256")))
            for asset in assets:
                if not isinstance(asset, dict):
                    continue
                for key in ("brief", "source", "script"):
                    item = asset.get(key)
                    if isinstance(item, dict) and isinstance(item.get("path"), str):
                        referenced_files.append((item["path"], item.get("sha256")))
                for key in ("data_sources", "outputs"):
                    for item in asset.get(key, []):
                        if isinstance(item, dict) and isinstance(item.get("path"), str):
                            referenced_files.append((item["path"], item.get("sha256")))
            for raw_path, expected_hash in referenced_files:
                path = (self.root / raw_path).resolve()
                try:
                    relative_path = path.relative_to(self.root).as_posix()
                except ValueError:
                    issues.append(f"upstream dependency escapes the frozen root: {raw_path}")
                    continue
                if relative_path not in frozen_index:
                    issues.append(f"upstream dependency is not covered by frozen_files: {relative_path}")
                elif not isinstance(expected_hash, str) or expected_hash.lower() != frozen_index[relative_path]:
                    issues.append(f"upstream dependency SHA-256 is inconsistent: {relative_path}")

        if issues:
            evidence = "; ".join(issues[:12])
            if len(issues) > 12:
                evidence += f"; ... total={len(issues)}"
            self.add(
                "P0", "READY-005", "最终冻结来源或文件 hash 无法核验",
                relative(self.manifest_path, self.root), evidence,
                "FULL/RECHECK 必须只读来源清楚且 hash 与冻结清单一致的 paper_final 包。",
                "返回冻结阶段重建新版本；不要在现有冻结目录内直接替换文件。",
            )
            self.not_ready = True
        else:
            self.passed("READY-006", f"冻结来源与 {len(frozen_index)} 个文件的 SHA-256 均一致")

    def discover_files(self) -> None:
        if not self.root.is_dir():
            raise SystemExit(f"project root is not a directory: {self.root}")
        self.files = list_project_files(self.root)
        paper_value = self.args.paper or self.manifest.get("paper")
        self.paper = self.resolve_manifest_path(paper_value)
        ambiguous_papers: list[Path] = []
        if self.paper is None:
            candidates = [
                path for path in self.files
                if path.suffix.lower() in {".pdf", ".docx"}
                and "AI 工具使用详情" not in path.name
                and "support" not in {part.lower() for part in path.parts}
                and "支撑材料" not in path.parts
            ]
            if len(candidates) == 1:
                self.paper = candidates[0]
            elif candidates:
                preferred = [path for path in candidates if re.search(r"final|paper|论文", path.stem, re.I)]
                if len(preferred) == 1:
                    self.paper = preferred[0]
                else:
                    ambiguous_papers = candidates
        source_value = self.manifest.get("source")
        self.source = self.resolve_manifest_path(source_value)
        if self.source is None:
            source_candidates = [path for path in self.files if path.suffix.lower() in {".tex", ".md"}]
            preferred = [path for path in source_candidates if re.search(r"main|paper|example|论文", path.stem, re.I)]
            if len(preferred) == 1:
                self.source = preferred[0]
            elif len(source_candidates) == 1:
                self.source = source_candidates[0]

        if ambiguous_papers:
            evidence = ", ".join(relative(path, self.root) for path in ambiguous_papers[:10])
            if self.args.mode == "paper-only" and self.source and self.source.exists():
                self.defer(
                    "SCOPE-FULL-006", "DEFERRED_TO_FULL", "最终 PDF 逐页视觉确认",
                    f"发现多个 PDF/Word 候选（{evidence}）；PAPER_ONLY 以已锁定源码审正文，FULL 再锁定唯一成品。",
                    "FULL",
                )
            else:
                self.add("P1", "FILE-001", "最终论文文件不唯一", str(self.root), evidence,
                         "Final Audit 必须锁定唯一最终论文。", "在清单中填写 paper，删除或移出旧版本。")
                self.scope_incomplete = True

        if self.paper and self.paper.exists():
            self.passed("FILE-002", f"锁定最终论文：{relative(self.paper, self.root)}")
        elif self.args.mode == "paper-only" and self.source and self.source.exists():
            self.passed("FILE-002", f"PAPER_ONLY 以源文件为正文审查主体：{relative(self.source, self.root)}")
        else:
            self.add("P0" if self.args.mode in {"full", "recheck"} else "P1", "FILE-003", "未找到最终 PDF/Word",
                     str(self.root), f"paper={paper_value!r}", "最终交卷必须锁定单一电子论文。",
                     "提供 --paper 或在 audit_manifest.json 中填写 paper。")
            self.scope_incomplete = True
        if self.source and self.source.exists():
            self.passed("FILE-004", f"锁定论文源文件：{relative(self.source, self.root)}")
        else:
            self.mark_unverified("FILE-005", "未找到唯一 LaTeX/Markdown 源文件；部分结构与引用检查将受限。")

        if self.args.mode in {"full", "recheck"}:
            problem_files = self.manifest.get("problem_files", [])
            verified_problem_files = [self.resolve_manifest_path(item) for item in problem_files]
            if not verified_problem_files or any(path is None or not path.exists() for path in verified_problem_files):
                self.mark_unverified("SCOPE-004", "未从清单确认原题/附件；题意覆盖与反向证据链必须由语义审查补证。")
                self.scope_incomplete = True

    def load_document_text(self) -> None:
        if self.paper and self.paper.exists():
            suffix = self.paper.suffix.lower()
            if suffix == ".pdf":
                self.paper_text = pdf_text(self.paper) or ""
            elif suffix == ".docx":
                self.paper_text = docx_text(self.paper) or ""
            self.paper_text_extractable = text_is_extractable(self.paper_text)
        if self.source and self.source.exists():
            raw = safe_read_text(self.source)
            self.source_text = strip_latex_comments(raw) if self.source.suffix.lower() == ".tex" else raw
            if self.source.suffix.lower() == ".tex":
                self.source_units = collect_latex_sources(self.source, self.root)

    def latex_bundle_text(self) -> str:
        if self.source_units:
            return "\n\n".join(text for _, text in self.source_units)
        return self.source_text

    def audit_paper_file(self) -> None:
        if not self.paper or not self.paper.exists():
            return
        suffix = self.paper.suffix.lower()
        size = self.paper.stat().st_size
        if size > 20 * 1024 * 1024:
            self.add("P0", "SUBMIT-001", "电子论文超过 20 MB", relative(self.paper, self.root), f"{size / 1024 / 1024:.2f} MB",
                     "2026 S0：电子论文不超过 20 MB。", "压缩图片或清理嵌入资源后重新生成 PDF。")
        else:
            self.passed("SUBMIT-002", f"电子论文大小 {size / 1024 / 1024:.2f} MB，未超过 20 MB")
        if suffix not in {".pdf", ".docx"}:
            self.add("P0", "SUBMIT-003", "电子论文格式不合规", relative(self.paper, self.root), suffix,
                     "2026 S0：电子论文为单个 PDF 或 Word。", "导出为单个 PDF/Word。")
            return
        if suffix == ".docx":
            if not self.paper_text:
                self.add("P1", "DOCX-001", "Word 论文无法读取", relative(self.paper, self.root), "document.xml 读取失败",
                         "最终论文必须可正常打开和审查。", "修复或重新导出 Word，并优先提供 PDF。")
            self.mark_unverified("DOCX-002", "Word 分页、摘要占版和字体需使用文档渲染后人工检查。")
            return

        pages = pdf_page_count(self.paper)
        if pages is not None:
            self.passed("PDF-001", f"PDF 总页数：{pages}")
        else:
            self.mark_unverified("PDF-002", "无法调用 pdfinfo 获取页数。")
        first = pdf_text(self.paper, first_page_only=True)
        if not text_is_extractable(first):
            self.mark_unverified(
                "PDF-003",
                "PDF 第 1 页文本不可提取或内容过少；不能据此判摘要缺失，必须查看第 1 页渲染图。",
            )
        else:
            first_compact = compact_text(first)
            has_abstract = bool(re.search(r"摘要", first_compact))
            has_keywords = bool(re.search(r"关键词|关键字|Keywords?", first_compact, re.I))
            if has_abstract and has_keywords:
                self.passed("PDF-004", "第 1 页检测到摘要与关键词")
            else:
                self.mark_unverified(
                    "PDF-005",
                    "第 1 页可提取文本中未同时识别摘要与关键词；可能是字体编码或自定义排版，需视觉确认。",
                )
            forbidden_front = [term for term in ["承诺书", "编号专用页"] if term in first]
            if forbidden_front:
                self.add("P0", "PDF-006", "电子论文包含承诺书或编号专用页", "PDF 第 1 页", ", ".join(forbidden_front),
                         "2026 S0：电子论文不放承诺书和编号专用页。", "从电子版删除相应页面。", True)
            if "问题重述" in first:
                self.warn("P2", "PDF-007", "问题重述疑似进入摘要页", "PDF 第 1 页", "检测到“问题重述”",
                          "S1：标题、摘要和关键词独占第 1 页。", "查看第 1 页渲染图；确认后再调整摘要与分页。")
        if self.paper_text:
            if re.search(r"(?m)^\s*目\s*录\s*$", self.paper_text):
                self.add("P0", "PDF-008", "正式论文存在目录", relative(self.paper, self.root), "提取文本中发现独立“目录”标题",
                         "2026 S0：正文不设目录。", "删除目录并重新编译。", True)
            appendix_page = None
            if pages and shutil.which("pdftotext"):
                for page in range(2, pages + 1):
                    result = run_command(["pdftotext", "-f", str(page), "-l", str(page), "-layout", str(self.paper), "-"], timeout=30)
                    if re.search(r"(?m)^\s*附\s*录\s*$", result.stdout):
                        appendix_page = page
                        break
            if appendix_page is not None:
                body_pages = max(0, appendix_page - 2)
                if body_pages > 30:
                    self.add("P0", "PDF-009", "正文超过 30 页", f"PDF 第 2—{appendix_page - 1} 页", f"估计正文 {body_pages} 页",
                             "2026 S0：正文不超过 30 页。", "删减低价值内容或调整排版，不压缩到不可读。")
                else:
                    self.passed("PDF-010", f"估计正文 {body_pages} 页，未超过 30 页")
            elif pages and pages > 31:
                self.mark_unverified("PDF-011", f"PDF 共 {pages} 页但未定位附录起始页，需人工确认正文是否超过 30 页。")
        fonts = pdf_font_names(self.paper)
        if fonts is None:
            self.mark_unverified("FONT-001", "无法调用 pdffonts 获取字体清单。")
        else:
            joined = " ".join(fonts).lower()
            self.passed("FONT-002", "PDF 字体清单：" + ", ".join(fonts[:20]))
            if not re.search(r"times|newroman|termes", joined):
                self.mark_unverified(
                    "FONT-003",
                    "PDF 字体别名中未识别 Times New Roman；子集字体/别名可能导致误判，需视觉或样式确认。",
                )
            if not re.search(r"simsun|song|fandol|stsong|宋", joined):
                self.mark_unverified(
                    "FONT-004",
                    "PDF 字体别名中未识别宋体系列；子集字体/CJK 别名可能导致误判，需视觉或样式确认。",
                )

    def render_visual_assets(self) -> None:
        if self.args.mode == "paper-only" and not self.args.render_pdf:
            return
        if not (self.args.render_pdf or self.args.mode == "visual-only"):
            self.mark_unverified(
                "VISUAL-001",
                "尚未生成最终 PDF 全页渲染和 contact sheet；FULL/RECHECK 应加 --render-pdf。",
            )
            return
        if self.paper is None or not self.paper.exists() or self.paper.suffix.lower() != ".pdf":
            self.mark_unverified("VISUAL-002", "没有最终 PDF，无法执行逐页视觉审查准备。")
            return
        output_value = self.args.visual_dir or ".final-audit-visual"
        output_dir = Path(output_value)
        output_dir = output_dir.resolve() if output_dir.is_absolute() else (self.root / output_dir).resolve()
        helper = Path(__file__).with_name("render_pdf_for_audit.py")
        command = [sys.executable, str(helper), str(self.paper), "--output-dir", str(output_dir)]
        try:
            result = run_command(command, cwd=self.root, timeout=300)
        except (OSError, subprocess.TimeoutExpired) as exc:
            self.mark_unverified("VISUAL-003", f"PDF 渲染未完成：{exc}")
            return
        if result.returncode != 0:
            self.mark_unverified("VISUAL-004", f"PDF 渲染失败，需改用其他渲染工具：{result.stdout[-1500:]}")
            return
        try:
            self.visual_output = json.loads(result.stdout)
        except json.JSONDecodeError:
            self.mark_unverified("VISUAL-005", "PDF 渲染完成但无法解析视觉清单输出。")
            return
        contact_sheet = str(self.visual_output.get("contact_sheet", ""))
        self.passed(
            "VISUAL-006",
            f"已渲染最终 PDF 共 {self.visual_output.get('page_count')} 页；contact sheet: {contact_sheet}",
        )
        candidate_pages = self.visual_output.get("candidate_pages", [])
        if candidate_pages:
            self.warn(
                "P3",
                "VISUAL-007",
                "PDF 渲染检测到空白候选页",
                contact_sheet,
                f"候选页：{candidate_pages}",
                "图像指标只能定位候选页，不能自动判断强制分页、标题孤行或排版失败。",
                "逐页查看候选页及其相邻页，确认后再记录视觉问题。",
            )

    def audit_source_structure(self) -> None:
        is_tex = bool(self.source and self.source.suffix.lower() == ".tex")
        text = self.latex_bundle_text() if is_tex else (self.source_text or self.paper_text)
        if not text:
            self.mark_unverified("STRUCT-001", "没有可提取文本，无法执行结构、语言和引用检查。")
            return
        self.audit_source_attributions(text)
        if is_tex:
            if "\\tableofcontents" in text:
                self.add("P0", "STRUCT-002", "LaTeX 启用了目录", relative(self.source, self.root), "发现 \\tableofcontents",
                         "2026 S0：正文不设目录。", "删除或注释目录命令。", True)
            abstract_found, abstract_evidence = find_abstract_evidence(text)
            if abstract_found:
                self.passed("ABSTRACT-001", f"LaTeX 检测到摘要：{abstract_evidence}")
            elif self.paper_text_extractable and find_abstract_evidence(self.paper_text)[0]:
                self.passed("ABSTRACT-001", "最终 PDF 文本检测到摘要标题")
            else:
                self.mark_unverified(
                    "ABSTRACT-002",
                    "未识别标准/常见自定义摘要写法；不能据此判摘要缺失，需检查最终 PDF 第 1 页。",
                )
            keyword_payloads = find_keyword_payloads(text)
            if not keyword_payloads and self.paper_text_extractable:
                keyword_payloads = find_keyword_payloads(self.paper_text)
            keyword_counts = [(payload, count_keywords(payload)) for payload in keyword_payloads]
            exact_five = [payload for payload, count in keyword_counts if count == 5]
            definite_non_five = [(payload, count) for payload, count in keyword_counts if count is not None and count != 5]
            if exact_five:
                self.passed("ABSTRACT-003", "在关键词命令或标签文本中识别到 5 个关键词")
            elif len(definite_non_five) == 1:
                raw, count = definite_non_five[0]
                self.add("P1", "ABSTRACT-004", "关键词数量不是 5 个", relative(self.source, self.root), f"检测到 {count} 个：{raw}",
                         "S1：摘要关键词固定 5 个。", "核对最终 PDF 后调整为 5 个关键词。")
            else:
                self.mark_unverified(
                    "ABSTRACT-005",
                    "未可靠解析关键词数量；支持自定义命令和标签文本，仍无法确定时由最终 PDF 视觉审查确认。",
                )
            self.audit_latex_figures(text)
            self.audit_latex_indentation()
            self.audit_latex_references(text)
            self.audit_dynamic_headings(text)
            self.audit_team_structure_rules(text)
            self.audit_duplicate_labels(text)
            self.audit_latex_compilation()
        else:
            reference_lines = re.findall(r"(?m)^\s*\[(\d+)\]\s+", text)
            if reference_lines and len(reference_lines) < 10:
                self.add("P1", "REF-001", "参考文献少于 10 篇", relative(self.source or self.paper, self.root), f"检测到 {len(reference_lines)} 篇",
                         "S1：参考文献不少于 10 篇且正文实际引用。", "补充真正使用的领域、模型和官方来源文献。")
        self.audit_ai_statement(text)
        self.audit_language(text)

    def audit_latex_indentation(self) -> None:
        units = self.source_units or ([(self.source, self.source_text)] if self.source else [])
        if not units:
            self.mark_unverified("INDENT-000", "未提供 LaTeX 源码，无法检查段首缩进设置与异常覆盖。")
            return
        all_settings: list[tuple[Path, dict[str, Any]]] = []
        noindent_hits: list[tuple[Path, dict[str, Any]]] = []
        manual_hits: list[tuple[Path, dict[str, Any]]] = []
        for path, unit_text in units:
            evidence = latex_indentation_evidence(unit_text)
            all_settings.extend((path, item) for item in evidence["settings"])
            noindent_hits.extend((path, item) for item in evidence["noindent"])
            manual_hits.extend((path, item) for item in evidence["manual_spacing"])

        raw_values = {
            re.sub(r"\s+", "", str(item["value"])).lower() for _, item in all_settings
        }

        def canonical_parindent(value: str) -> str:
            if re.fullmatch(r"2(?:\.0)?(?:em|zw|\\ccwd)", value):
                return "TWO_CJK_CHARS"
            if re.fullmatch(r"0(?:\.0)?(?:pt|em|zw)?", value):
                return "ZERO"
            return value

        normalized_values = {canonical_parindent(value) for value in raw_values}
        recognized_two_character = "TWO_CJK_CHARS" in normalized_values
        zero_values = {value for value in normalized_values if value == "ZERO"}
        locations = [
            f"{relative(path, self.root)}:{item['line']}={item['value']}" for path, item in all_settings
        ]
        if zero_values:
            self.warn(
                "P2", "INDENT-001", "LaTeX 源码显式取消正文首行缩进",
                ", ".join(locations[:12]), ", ".join(sorted(zero_values)),
                "S1：正文统一首行缩进 2 个汉字；局部取消必须有明确版式理由。",
                "核对作用域；正文段落恢复统一 parindent，摘要、标题等例外由最终 PDF 确认。",
            )
        if len(normalized_values) > 1:
            self.warn(
                "P3", "INDENT-002", "LaTeX 多处 parindent 设置不一致候选",
                ", ".join(locations[:12]), ", ".join(sorted(normalized_values)),
                "不同章节不得无依据切换正文缩进规则。",
                "确认各设置作用域，保留一个正文全局规则；合法局部例外需在视觉审查中确认。",
            )
        if recognized_two_character and not zero_values and len(normalized_values) == 1:
            self.passed("INDENT-003", f"源码识别到统一的 2 字符级 parindent 设置：{', '.join(locations[:6])}")
        elif not all_settings:
            self.mark_unverified(
                "INDENT-004",
                "源码未显式识别 parindent 设置；可能由文档类/宏包提供，PAPER_ONLY 不能据此判失败，需检查类文件并由最终 PDF 确认。",
            )
        elif not recognized_two_character and not zero_values:
            self.mark_unverified(
                "INDENT-005",
                f"识别到 parindent={', '.join(sorted(raw_values))}，但无法确认等价于 2 个汉字；需核对类文件和最终 PDF。",
            )

        if noindent_hits:
            evidence = " | ".join(
                f"{relative(path, self.root)}:{item['line']} {item['context']}" for path, item in noindent_hits[:12]
            )
            self.warn(
                "P3", "INDENT-006", f"发现异常 \\noindent 候选 {len(noindent_hits)} 处",
                relative(self.source, self.root), evidence,
                "正文段落不应无依据绕过统一首行缩进；摘要、关键词或特殊块可作为例外。",
                "逐处确认作用对象；只移除正文中的异常 noindent，合法特殊块保留并交视觉审查确认。",
            )
        if manual_hits:
            evidence = " | ".join(
                f"{relative(path, self.root)}:{item['line']} {item['token']}" for path, item in manual_hits[:12]
            )
            self.warn(
                "P3", "INDENT-007", f"发现人工空格模拟缩进候选 {len(manual_hits)} 处",
                relative(self.source, self.root), evidence,
                "正文缩进应由统一段落样式控制，不用全角空格、重复 quad 或 hspace 模拟。",
                "确认后删除人工空格，统一由 parindent/文档类控制。",
            )

    def audit_drawing_placeholders(self, text: str) -> None:
        candidates = drawing_placeholder_candidates(text)
        if not candidates:
            return
        evidence = " | ".join(
            f"行 {item['line']} {item['marker']}：{item['excerpt']}"
            for item in candidates[:12]
        )
        location = relative(self.source or self.paper, self.root)
        if self.args.mode == "paper-only":
            self.defer(
                "DRAWING-001",
                "DEFERRED_TO_DRAWING",
                f"发现 {len(candidates)} 处图表占位",
                f"{location}；{evidence}",
                "DRAWING/FULL",
            )
            return
        if self.args.mode in {"visual-only", "full", "recheck"}:
            mode_label = self.args.mode.upper().replace("-", "_")
            self.add(
                "P1", "DRAWING-002", f"最终 PDF 仍有图表占位 {len(candidates)} 处",
                location, evidence,
                "VISUAL_ONLY/FULL 审查的是绘图完成后的最终 PDF，不得保留“待绘制/待补图/此处插入”等标记。",
                "完成图表、补齐图前图后叙事并重新生成最终 PDF。",
                stage_status=f"{mode_label}_BLOCKER",
            )

    def audit_source_attributions(self, text: str) -> None:
        candidates = source_attribution_candidates(text)
        self.source_attributions = candidates
        if not candidates:
            return
        evidence = " | ".join(
            f"行 {item['line']} {item['claim']}：{item['excerpt']}"
            for item in candidates[:10]
        )
        self.warn(
            "P2", "SOURCE-001", f"发现题面/附件来源性声明候选 {len(candidates)} 处",
            relative(self.source or self.paper, self.root), evidence,
            "涉及参数、分布、相关系数、关系矩阵、模拟边、阈值或权重的‘题目给出/根据题设’声明必须反查原题或附件。自设、模拟或估计量不得冒充题设事实。",
            "逐项填写‘题面事实与模型设定来源矩阵’；只有对照原题/附件后确认冒充来源，才升为 P1。",
        )

    def audit_latex_figures(self, text: str) -> None:
        source_dir = self.source.parent if self.source else self.root
        graphics_roots = [source_dir, self.root, self.root / "figures"]
        for block in re.findall(r"\\graphicspath\s*\{((?:\s*\{[^{}]+\}\s*)+)\}", text, re.S):
            for item in re.findall(r"\{([^{}]+)\}", block):
                graphics_roots.append((source_dir / item).resolve())
        figure_blocks = re.findall(r"\\begin\{figure\*?\}(.*?)\\end\{figure\*?\}", text, re.S)
        table_blocks = re.findall(r"\\begin\{table\*?\}(.*?)\\end\{table\*?\}", text, re.S)
        self.passed("FIG-001", f"LaTeX 检测到 {len(figure_blocks)} 个 figure、{len(table_blocks)} 个 table")
        for index, block in enumerate(figure_blocks, 1):
            if not re.search(r"\\caption(?:\[[^]]*\])?\{", block):
                self.add("P2", "FIG-002", "图片缺少 caption", f"figure #{index}", "未找到 \\caption",
                         "S1：图题完整并置于图下。", "补充准确图题。", True)
            if not re.search(r"\\label\{[^}]+\}", block):
                self.add("P2", "FIG-003", "图片缺少 label", f"figure #{index}", "未找到 \\label",
                         "图表编号和正文引用必须一致。", "补充唯一 label 并使用交叉引用。", True)
            for name in re.findall(r"\\includegraphics(?:\[[^]]*\])?\{([^}]+)\}", block):
                stem = Path(name).stem
                if re.fullmatch(r"\d+|图\d+", stem):
                    self.warn("P3", "FIG-004", "图片文件名缺少语义", name, stem,
                              "文件名应可追溯且便于版本核验。", "改用如 q2_risk_frontier.pdf 的语义名称。")
                candidate = Path(name)
                extensions = [""] if candidate.suffix else [".pdf", ".png", ".jpg", ".jpeg", ".svg"]
                exists = any((base / (name + ext)).exists() for base in graphics_roots for ext in extensions)
                if not exists:
                    self.mark_unverified(
                        "FIG-005",
                        f"静态路径扫描未找到图片 {name}；可能由宏/构建步骤生成，以 LaTeX 编译日志为准。",
                    )
        for index, block in enumerate(table_blocks, 1):
            if not all(token in block for token in ["\\toprule", "\\midrule", "\\bottomrule"]):
                self.warn("P3", "TABLE-001", "表格未检测到 booktabs 三线命令", f"table #{index}", "缺少 top/mid/bottomrule",
                          "S1：正文表格优先三线表，但允许其他等价实现。", "查看最终 PDF；若不是等价三线表再调整。")
        naked = re.finditer(
            r"\\(?:sub)*section\*?\{([^}]+)\}\s*\\begin\{(figure\*?|table\*?)\}", text, re.S
        )
        for match in naked:
            heading = match.group(1).strip()
            if heading in {"符号说明", "支撑材料文件清单"}:
                continue
            line = text.count("\n", 0, match.start()) + 1
            self.warn("P2", "FIG-006", "重要图表疑似裸接小标题", f"{relative(self.source, self.root)}:{line}", match.group(0)[:180],
                      "S1：图前说明用途、问题和统计口径。", "人工核对上下文；确认后补自然引入和图后量化解释。")

    def audit_latex_references(self, text: str) -> None:
        bibitem_keys = re.findall(r"\\bibitem(?:\[[^]]*\])?\{([^}]+)\}", text)
        citation_commands = find_citation_commands(text)
        cited: set[str] = {key for command in citation_commands for key in command["keys"]}
        nocite_groups = re.findall(r"\\nocite\*?\s*\{([^{}]+)\}", strip_latex_comments(text), re.I)
        bib_keys: list[str] = []
        available_bibs = [path for path in self.files if path.suffix.lower() == ".bib"]
        loaded_bibs = declared_bibliography_files(self.source_units, self.root) if self.source_units else []
        bibliography_scope_resolved = True
        bibliography_scope_verified = bool(loaded_bibs or bibitem_keys)
        if loaded_bibs:
            selected_bibs = loaded_bibs
        elif len(available_bibs) == 1:
            selected_bibs = available_bibs
            self.mark_unverified(
                "REF-010",
                f"未识别 bibliography/addbibresource 声明；暂按唯一 .bib 文件 {relative(available_bibs[0], self.root)} 核对，需确认最终构建实际载入该文件。",
            )
        elif bibitem_keys or not available_bibs:
            selected_bibs = []
        else:
            selected_bibs = []
            bibliography_scope_resolved = False
            self.mark_unverified(
                "REF-011",
                "存在多个 .bib 文件但未识别最终加载声明；无法可靠判断 citation key、实际文末数量和孤儿文献。",
            )
        for path in selected_bibs:
            bib_keys.extend(bib_entry_keys(safe_read_text(path)))
        all_keys = set(bibitem_keys) | set(bib_keys)
        nocited: set[str] = set()
        for group in nocite_groups:
            if group.strip() == "*":
                nocited.update(all_keys)
            else:
                nocited.update(key.strip() for key in group.split(",") if key.strip())
        if bibitem_keys:
            rendered_keys = set(bibitem_keys)
        else:
            rendered_keys = (cited | nocited) & all_keys
        ref_count = len(rendered_keys)
        if bibliography_scope_resolved:
            if ref_count < 10:
                self.add("P1", "REF-002", "参考文献少于 10 篇", relative(self.source, self.root), f"检测到 {ref_count} 个正文实际引用/文末列出的唯一文献键",
                         "S1：参考文献不少于 10 篇且正文对应引用。", "补充真正使用的领域、方法、标准和数据来源文献。")
            else:
                self.passed("REF-003", f"检测到 {ref_count} 个正文实际引用/文末列出的唯一文献键（≥10）")
            missing = sorted(cited - all_keys)
            if missing:
                self.add("P1", "REF-004", "正文引用键在文末不存在", relative(self.source, self.root), ", ".join(missing[:20]),
                         "正文引用与文末文献必须一一对应。", "补充文献条目或修正引用键。", True)
            orphan = sorted(set(bibitem_keys) - cited if bibitem_keys else nocited - cited)
            if orphan:
                self.add("P2", "REF-005", "存在文末孤儿文献", relative(self.source, self.root), ", ".join(orphan[:20]),
                         "文末只列正文实际引用的资料。", "在真实引用位置标注或删除无关条目。")
        else:
            missing = []
            orphan = []
        manual_tokens = manual_numeric_reference_tokens(text)
        citation_like = [item for item in manual_tokens if item["citation_like"]]
        if citation_like:
            evidence = " | ".join(
                f"行 {item['line']} {item['token']}：{item['context']}" for item in citation_like[:15]
            )
            self.warn(
                "P1", "REF-007", f"发现手工输入文献编号候选 {len(citation_like)} 处",
                relative(self.source, self.root), evidence,
                "正文文献编号必须由 citation command 生成；不得手工输入（2）或 [2] 冒充引用。",
                "逐处核对上下文；确认是文献引用时改用 cite/parencite/supercite 等统一命令并重新编译。",
            )
        if bibliography_scope_verified and citation_commands and ref_count >= 10 and not missing and not orphan and not citation_like:
            commands = sorted({item["command"] for item in citation_commands})
            self.passed(
                "REF-008",
                f"检测到 {len(citation_commands)} 处 citation command（{', '.join(commands)}），引用键与文末条目对应且未发现手工编号候选",
            )
        elif not citation_commands and ref_count:
            self.mark_unverified(
                "REF-009",
                "文末存在参考文献，但未识别 cite/parencite/supercite 等 citation command；需确认是否使用了自定义引用宏或手工编号。",
            )
        for label, pattern in MODEL_PATTERNS.items():
            match = re.search(pattern, text, re.I)
            if not match:
                continue
            start = max(text.rfind("\n\n", 0, match.start()), 0)
            end_pos = text.find("\n\n", match.end())
            end = len(text) if end_pos == -1 else end_pos
            paragraph = text[start:end]
            if not find_citation_commands(paragraph):
                line = text.count("\n", 0, match.start()) + 1
                self.mark_unverified(
                    "REF-006",
                    f"{label} 首次命中所在段未识别 citation command（{relative(self.source, self.root)}:{line}）；手工 [n]/（n）不算有效引用，需确认是否为正式引入。",
                )

    def audit_pdf_reference_links(self) -> None:
        if self.paper is None or not self.paper.exists() or self.paper.suffix.lower() != ".pdf":
            self.mark_unverified("REF-LINK-001", "没有最终 PDF，无法检查引用编号到参考文献的实际跳转。")
            return
        inspection = inspect_pdf_reference_links(self.paper)
        if inspection is None:
            self.mark_unverified("REF-LINK-002", "当前环境无法读取 PDF 内部链接注释；引用跳转必须人工点击核验。")
            return
        reference_start = inspection.get("reference_start_page")
        if reference_start is None:
            self.mark_unverified("REF-LINK-003", "PDF 中未可靠定位参考文献起始页，无法自动判定哪些内部链接指向 bibliography。")
            return
        links = [
            link for link in inspection.get("links", [])
            if re.search(r"\d", str(link.get("display", "")))
        ]
        if not links:
            source_text = self.latex_bundle_text() or self.source_text
            commands = find_citation_commands(source_text) if source_text else []
            if commands:
                self.add(
                    "P1", "REF-LINK-004", "最终 PDF 缺少可用的文献编号跳转",
                    relative(self.paper, self.root),
                    f"源码检测到 {len(commands)} 处 citation command，参考文献起始页为第 {reference_start} 页，但没有数字内部链接指向 bibliography。",
                    "citation command、文末条目与最终 PDF 跳转应保持一致。",
                    "检查 hyperref/模板链接设置，重新编译并逐项点击正文文献编号。",
                )
            else:
                self.mark_unverified(
                    "REF-LINK-004",
                    f"已定位参考文献起始页为第 {reference_start} 页，但未读取到数字内部链接，且没有源码 citation command 证据；需人工确认。",
                )
            return

        source_text = self.latex_bundle_text() or self.source_text
        literal_tokens = manual_numeric_reference_tokens(source_text) if source_text else []
        literal_by_token: dict[str, list[dict[str, Any]]] = {}
        literal_by_signature: dict[tuple[str, ...], list[dict[str, Any]]] = {}
        for item in literal_tokens:
            literal_by_token.setdefault(item["token"], []).append(item)
            signature = tuple(re.findall(r"\d+", item["token"]))
            if signature:
                literal_by_signature.setdefault(signature, []).append(item)
        suspicious: list[str] = []
        style_mismatches: list[str] = []
        bibliography_uses_brackets = bool(
            re.search(r"(?m)^\s*\[\s*\d+\s*\]", self.paper_text or "")
        )
        for link in links:
            display = re.sub(r"\s+", "", str(link.get("display", "")))
            signature = tuple(re.findall(r"\d+", display))
            matching_literals = literal_by_token.get(display, []) or literal_by_signature.get(signature, [])
            interval_like = bool(
                re.fullmatch(r"\[(?:0\s*[,，]\s*1|(?:19|20)\d{2}\s*[,，–—-]\s*(?:19|20)\d{2})\]", display)
            )
            if matching_literals or interval_like:
                source_context = "；".join(
                    f"源码行 {item['line']} {item['context']}" for item in matching_literals[:3]
                ) or "显示文本本身像普通区间/年份范围"
                suspicious.append(
                    f"PDF 第 {link['source_page']} 页 {display} → 第 {link['target_page']} 页；{source_context}"
                )
            if bibliography_uses_brackets and re.fullmatch(
                r"[（(]\s*\d+(?:\s*[,，]\s*\d+)*\s*[)）]", display
            ):
                style_mismatches.append(
                    f"PDF 第 {link['source_page']} 页 {display} → 第 {link['target_page']} 页"
                )
        if style_mismatches:
            self.add(
                "P1", "REF-LINK-007", "正文引用与文末编号样式不一致",
                relative(self.paper, self.root), " | ".join(style_mismatches[:12]),
                "正文与文末编号应由统一 citation/bibliography 样式生成；文末为 [n] 时不得在正文手写（n）。",
                "删除手工括号编号，统一 citation command 与 bibliography style 后重新编译并复核跳转。",
            )
        if suspicious:
            self.warn(
                "P1", "REF-LINK-005", f"发现非 citation command 数字文本误链参考文献候选 {len(suspicious)} 处",
                relative(self.paper, self.root), " | ".join(suspicious[:12]),
                "只有 citation command 生成的编号可以链接到 bibliography；集合、区间、数组和年份范围不得误链。",
                "对照源码与 PDF 逐处点击；确认误链后修正宏/链接范围并重新编译。",
            )
        else:
            self.passed(
                "REF-LINK-006",
                f"读取到 {len(links)} 个正文→参考文献内部链接，未发现与源码手工数字 token 直接匹配的误链接候选；仍需按唯一 citation key 人工抽查目的条目。",
            )

    def audit_dynamic_headings(self, text: str) -> None:
        headings = re.findall(r"\\(?:subsection|subsubsection)\*?\{([^}]+)\}", text)
        generic = [
            "问题思路与模型依据", "模型建立与求解", "结果与对比分析", "模型检验与灵敏度分析",
            "问题思路", "结果分析", "模型检验",
        ]
        repeated = {title: headings.count(title) for title in generic if headings.count(title) >= 2}
        if repeated:
            self.warn("P2", "STRUCT-003", "多问小标题疑似机械重复", relative(self.source, self.root), json.dumps(repeated, ensure_ascii=False),
                      "S1：固定功能闭环，但三级标题按真实数学内容动态命名。", "人工确认是否确为模板复用；按真实数学内容定点重命名。")

    def audit_team_structure_rules(self, text: str) -> None:
        bold_count = len(re.findall(r"\\textbf\s*\{", text))
        if 10 <= bold_count <= 20:
            self.passed("STRUCT-004", f"检测到 {bold_count} 个 \\textbf 信息锚点（目标 10—20）")
        else:
            self.mark_unverified(
                "STRUCT-005",
                f"仅按 \\textbf 检测到 {bold_count} 个加粗；自定义命令/样式可能未计入，需视觉确认信息锚点约 10—20 个。",
            )

        assumption = re.search(
            r"\\section\*?\{模型假设\}(.*?)(?=\\section\*?\{|\\begin\{thebibliography\}|\Z)",
            text,
            re.S,
        )
        if not assumption:
            self.mark_unverified(
                "STRUCT-006",
                "未按内置 section 语法定位模型假设；自定义标题命令可能未识别，需语义审查确认。",
            )
        else:
            body = assumption.group(1).strip()
            items = len(re.findall(r"\\item\b", body))
            if len(re.sub(r"\\\w+(?:\[[^]]*\])?\{?", "", body).strip()) < 12:
                self.add("P1", "STRUCT-007", "模型假设章节内容为空或过少", relative(self.source, self.root), body[:160],
                         "S1：假设应覆盖必要现实简化和边界。", "补充 4—6 条真正被模型使用的宽口径假设。")
            elif items and not 4 <= items <= 6:
                self.warn("P2", "STRUCT-008", "模型假设条数偏离 4—6", relative(self.source, self.root), f"检测到 {items} 个 item",
                          "S1：模型假设通常 4—6 条，不机械堆砌。", "人工核对内容后合并重复假设或补充遗漏边界。")

        if not re.search(r"\\section\*?\{[^}]*数据预处理[^}]*\}", text):
            self.mark_unverified("STRUCT-009", "未检测到独立数据预处理章节；若为数据型 C 题，需按 S1 补充 1—3 页数据审计与处理影响。")

        if not re.search(r"\\section\*?\{[^}]*模型评价[^}]*改进[^}]*结论[^}]*\}", text):
            self.mark_unverified(
                "STRUCT-010",
                "未按内置 section 语法识别合并末章；需在最终目录结构中人工确认“模型评价、改进推广与结论”。",
            )

        restatement = re.search(
            r"\\section\*?\{问题重述\}(.*?)(?=\\section\*?\{|\\begin\{thebibliography\}|\Z)",
            text,
            re.S,
        )
        if restatement:
            headings = re.findall(r"\\subsection\*?\{([^}]+)\}", restatement.group(1))
            has_background = any("背景" in heading for heading in headings)
            has_tasks = any(any(term in heading for term in ["任务", "关键条件", "需要解决的问题"]) for heading in headings)
            if not (has_background and has_tasks):
                self.warn("P2", "STRUCT-011", "问题重述两层结构疑似不完整", relative(self.source, self.root), ", ".join(headings),
                          "S1：问题重述采用问题背景 + 问题任务与关键条件。", "人工核对自定义标题和正文内容后再调整。")

    def audit_duplicate_labels(self, text: str) -> None:
        labels = re.findall(r"\\label\s*\{([^}]+)\}", text)
        duplicates = sorted({label for label in labels if labels.count(label) > 1})
        if duplicates:
            self.add(
                "P1",
                "LATEX-001",
                "LaTeX 存在重复 label",
                relative(self.source, self.root),
                ", ".join(duplicates[:30]),
                "图、表、公式和章节交叉引用必须唯一且可追踪。",
                "为重复 label 改用唯一名称，重新编译并复核全部引用。",
                True,
            )
        else:
            self.passed("LATEX-002", f"检测到 {len(labels)} 个 label，未发现重复")

    def audit_latex_compilation(self) -> None:
        if not self.args.compile_latex:
            self.mark_unverified("LATEX-003", "未执行临时 LaTeX 编译；FULL/RECHECK 应加 --compile-latex。")
            return
        latex_config = self.manifest.get("latex", {})
        if not isinstance(latex_config, dict):
            latex_config = {}
        source = self.resolve_manifest_path(latex_config.get("source")) or self.source
        if source is None or not source.exists() or source.suffix.lower() != ".tex":
            self.mark_unverified("LATEX-004", "未锁定可编译的最终 .tex 主文件。")
            return
        engine = str(latex_config.get("engine", "xelatex")).lower()
        if engine not in {"xelatex", "pdflatex", "lualatex"}:
            self.mark_unverified("LATEX-005", f"未支持的 LaTeX engine={engine!r}；请改为 xelatex/pdflatex/lualatex。")
            return
        extra_args = latex_config.get("args", [])
        if not isinstance(extra_args, list) or not all(isinstance(item, str) for item in extra_args):
            self.mark_unverified("LATEX-006", "latex.args 不是安全字符串数组，未执行编译。")
            return
        with tempfile.TemporaryDirectory(prefix="cumcm-latex-audit-") as temp_name:
            build_dir = Path(temp_name)
            if shutil.which("latexmk"):
                mode = {"xelatex": "-xelatex", "pdflatex": "-pdf", "lualatex": "-lualatex"}[engine]
                command = [
                    "latexmk", mode, "-interaction=nonstopmode", "-file-line-error",
                    f"-outdir={build_dir}", *extra_args, source.name,
                ]
            elif shutil.which(engine):
                command = [
                    engine, "-interaction=nonstopmode", "-file-line-error",
                    f"-output-directory={build_dir}", *extra_args, source.name,
                ]
            else:
                self.mark_unverified("LATEX-007", f"环境中没有 latexmk 或 {engine}，无法执行编译质量检查。")
                return
            timeout = int(latex_config.get("timeout_seconds", 180))
            try:
                result = run_command(command, cwd=source.parent, timeout=timeout)
            except (OSError, subprocess.TimeoutExpired) as exc:
                self.add("P1", "LATEX-008", "LaTeX 编译未完成", relative(source, self.root), str(exc),
                         "最终 LaTeX 工程应可从统一入口正常编译。", "修复依赖、路径或超时后重新编译。")
                return
            log_path = build_dir / f"{source.stem}.log"
            log_text = safe_read_text(log_path) if log_path.exists() else ""
            combined = result.stdout + "\n" + log_text
            missing_files = sorted(set(
                re.findall(r"(?:LaTeX Error|Package [^\n]+ Error): File [`']([^`']+)[`'] not found", combined, re.I)
            ))
            image_files_only = bool(missing_files) and all(
                Path(name).suffix.lower() in IMAGE_SUFFIXES for name in missing_files
            )
            explicit_drawing_placeholders = drawing_placeholder_candidates(self.latex_bundle_text())
            deferred_drawing_failure = (
                self.args.mode == "paper-only"
                and image_files_only
                and bool(explicit_drawing_placeholders)
            )
            if result.returncode != 0 or not (build_dir / f"{source.stem}.pdf").exists():
                tail = result.stdout[-2500:].strip()
                if deferred_drawing_failure:
                    self.defer(
                        "LATEX-009", "DEFERRED_TO_DRAWING", "占位图导致 LaTeX 暂不能完整编译",
                        f"缺失图片：{', '.join(missing_files[:20])}", "DRAWING/FULL",
                    )
                else:
                    self.add("P1", "LATEX-009", "最终 LaTeX 无法正常编译", relative(source, self.root),
                             f"exit={result.returncode}\n{tail}", "最终工程应可生成可审查 PDF。",
                             "按首个真实错误修复，清理后临时全量编译，再复核日志。")
            else:
                self.passed("LATEX-010", f"临时编译成功：{engine} / {source.name}")

            if missing_files:
                if deferred_drawing_failure:
                    self.defer(
                        "LATEX-011", "DEFERRED_TO_DRAWING", "编译日志确认占位图片尚未生成",
                        ", ".join(missing_files[:20]), "DRAWING/FULL",
                    )
                else:
                    self.add("P1", "LATEX-011", "编译日志确认图片/输入文件缺失", relative(source, self.root),
                             ", ".join(missing_files[:20]), "最终工程中的全部图片和输入文件必须可读取。",
                             "补齐文件或修正相对路径后重新编译。")

            undefined_citations = sorted(set(re.findall(r"Citation [`']([^`']+)[`'] .* undefined", combined, re.I)))
            if undefined_citations or re.search(r"undefined citations", combined, re.I):
                evidence = ", ".join(undefined_citations[:30]) or "日志报告 undefined citations"
                self.add("P1", "LATEX-012", "LaTeX 存在 undefined citation", relative(source, self.root), evidence,
                         "正文引用必须解析到真实文献条目。", "补齐 bib/bbl 条目并重新编译至引用稳定。")

            undefined_refs = sorted(set(re.findall(r"Reference [`']([^`']+)[`'] .* undefined", combined, re.I)))
            if undefined_refs or re.search(r"undefined references", combined, re.I):
                evidence = ", ".join(undefined_refs[:30]) or "日志报告 undefined references"
                self.add("P1", "LATEX-013", "LaTeX 存在 undefined reference", relative(source, self.root), evidence,
                         "章节、图、表和公式编号必须完整解析。", "补齐 label/ref 并重新编译至引用稳定。")

            overfull_values = [float(value) for value in re.findall(r"Overfull \\hbox \(([-+]?\d+(?:\.\d+)?)pt too wide\)", combined)]
            if overfull_values:
                self.warn("P2", "LATEX-014", "LaTeX 存在 Overfull hbox", relative(source, self.root),
                          f"共 {len(overfull_values)} 处，最大 {max(overfull_values):.2f} pt",
                          "越界内容可能在最终 PDF 中溢出或碰撞。", "逐页查看对应位置；仅修复最终 PDF 中真实可见的溢出。")

    def audit_ai_statement(self, text: str) -> None:
        statement_pos = text.find("AI 工具使用声明")
        reference_candidates = [pos for pos in [text.find("参考文献"), text.find("\\begin{thebibliography}")] if pos >= 0]
        reference_pos = min(reference_candidates) if reference_candidates else -1
        if statement_pos < 0:
            self.add("P0", "AI-001", "缺少 AI 工具使用声明", relative(self.source or self.paper, self.root), "未找到一级标题“AI 工具使用声明”",
                     "2026 S0：参考文献之前必须设置 AI 工具使用声明。", "按实际使用情况加入规定文本。")
        elif reference_pos >= 0 and statement_pos > reference_pos:
            self.add("P0", "AI-002", "AI 工具使用声明位置错误", relative(self.source or self.paper, self.root), "声明位于参考文献之后",
                     "2026 S0：AI 工具使用声明位于参考文献之前。", "移动声明到参考文献之前。", True)
        else:
            self.passed("AI-003", "检测到 AI 工具使用声明且位置在参考文献之前")
        ai_used = self.manifest.get("ai_used")
        details = [path for path in self.files if path.name == "AI 工具使用详情.pdf"]
        if ai_used is True:
            if "使用了 AI 工具" not in text:
                self.add("P0", "AI-004", "AI 使用声明与清单不一致", relative(self.source or self.paper, self.root), "清单 ai_used=true，但正文未检测到使用 AI 的规定表述",
                         "AI 使用必须如实声明。", "核对实际使用并采用规定文本。")
            if not details:
                if self.args.mode == "paper-only":
                    self.defer(
                        "SCOPE-FULL-004", "DEFERRED_TO_FULL", "AI 工具使用详情.pdf",
                        "清单 ai_used=true；PAPER_ONLY 只核对正文声明，详情文件存在性属于 FULL。",
                        "FULL",
                    )
                else:
                    self.add("P0", "AI-005", "缺少 AI 工具使用详情.pdf", str(self.root), "清单 ai_used=true",
                             "2026 S0：使用 AI 时支撑材料包含 AI 工具使用详情.pdf。", "生成详情 PDF 并纳入支撑材料。")
            else:
                self.passed("AI-006", f"发现 {relative(details[0], self.root)}")
        elif ai_used is False:
            if "未使用任何 AI 工具" not in text:
                self.add("P0", "AI-007", "未使用 AI 的声明文本未确认", relative(self.source or self.paper, self.root), "清单 ai_used=false",
                         "2026 S0：未使用时采用规定声明。", "加入准确的未使用 AI 声明。")
        else:
            self.mark_unverified("AI-008", "清单未声明 ai_used；无法核对正文声明与真实使用记录的一致性。")

    def audit_language(self, text: str) -> None:
        for phrase in META_PHRASES:
            positions = [match.start() for match in re.finditer(re.escape(phrase), text)]
            if len(positions) >= 3:
                lines = [str(text.count("\n", 0, pos) + 1) for pos in positions[:8]]
                self.warn("P3", "LANG-001", f"模板化措辞“{phrase}”出现 {len(positions)} 次",
                          f"行 {', '.join(lines)}", phrase,
                          "S1：能直接陈述事实时不增加元话语或无证据评价。", "逐处结合上下文确认，只保留确有强调作用的表达。")
        chains = list(re.finditer(r"[^\n，。；]{1,18}[—→][^\n，。；]{1,18}[—→][^\n，。；]{1,18}", text))
        if len(chains) >= 3:
            examples = " | ".join(match.group(0).strip() for match in chains[:3])
            self.warn("P3", "LANG-002", "A—B—C 概念链偏多", relative(self.source or self.paper, self.root), examples,
                      "S1：普通处理不要反复包装成框架/体系。", "人工确认用途，只在摘要总述或总流程图保留真正必要的概念链。")

        candidate_groups = readability_candidates(text)
        specifications = [
            (
                "long_sentences", "LANG-003", "发现长句候选",
                "句长只是定位线索；只有主谓距离过远、职责过多或难以一次读清时才需要拆分。",
                "人工检查是否同时承担背景、条件、方法、结果和评价；必要时按事实—推断—动作定点拆句。",
            ),
            (
                "dense_connectors", "LANG-004", "发现多层逻辑连接候选",
                "连接词数量不能直接证明语言有问题，应核对因果、转折、递进和条件层级是否清楚。",
                "人工还原逻辑层级；若一次读不清，再拆成条件/事实、推断和本文动作。",
            ),
            (
                "abstract_stacks", "LANG-005", "发现抽象名词堆叠候选",
                "技术名词本身不是 AI 痕迹；只有抽象包装替代对象、动作、变量或证据时才构成问题。",
                "人工确认后将抽象名词还原为对象—动作—数据/变量—结果，不做全篇批量替换。",
            ),
        ]
        for key, check_id, title, rule, recommendation in specifications:
            items = candidate_groups[key]
            if not items:
                continue
            evidence = " | ".join(
                f"行 {item['line']}（{item['length']} 字/连接词 {item['connector_count']}/抽象词 {item['abstract_noun_count']}）：{item['excerpt']}"
                for item in items[:6]
            )
            self.warn(
                "P3", check_id, f"{title} {len(items)} 处",
                relative(self.source or self.paper, self.root), evidence, rule, recommendation,
            )

        showcase_items = showcase_citation_candidates(text)
        if showcase_items:
            evidence = " | ".join(
                f"行 {item['line']} {item['phrase']}：{item['excerpt']}"
                for item in showcase_items[:8]
            )
            self.warn(
                "P2", "LANG-006", f"发现展示式引用表达候选 {len(showcase_items)} 处",
                relative(self.source or self.paper, self.root), evidence,
                "引用应就近附着于其支持的方法名、性质或句子主张；“详见文献”本身不能说明支持对象。",
                "人工核对 citation 的真实支持对象；优先改成“方法/主张 + citation + 本题作用”，不得直接删除可靠引用。",
            )

        terminology_items = terminology_consistency_candidates(
            text, self.manifest.get("terminology_groups")
        )
        if terminology_items:
            evidence = " | ".join(
                f"{item['canonical']}："
                + "、".join(
                    f"{variant}(行 {','.join(str(line) for line in item['lines'][variant][:6])})"
                    for variant in item["variants"]
                )
                for item in terminology_items[:8]
            )
            self.warn(
                "P2", "TERM-001", f"发现术语变体共现候选 {len(terminology_items)} 组",
                relative(self.source or self.paper, self.root), evidence,
                "同一技术概念、动作、模型和方案应使用稳定名称；不同概念则必须明确定义差异。",
                "逐组确认这些写法在本文是否同义；同义时统一规范名称，不同义时补充区分定义，禁止自动批量替换。",
            )

    def audit_required_files(self) -> None:
        required = self.manifest.get("required_files", [])
        if self.args.mode in {"full", "recheck"} and not required:
            self.mark_unverified("RESULT-001", "清单未列 required_files；强制结果模板需由题面覆盖审查补充。")
        for item in required:
            if not isinstance(item, dict):
                continue
            path = self.resolve_manifest_path(item.get("path"))
            if path is None or not path.exists():
                self.add("P0", "RESULT-002", "强制结果文件缺失", str(item.get("path")), "文件不存在",
                         "题目指定结果文件属于硬交付约束。", "按精确文件名生成最终结果文件。")
                continue
            self.passed("RESULT-003", f"强制结果文件存在：{relative(path, self.root)}")
            expected_hash = item.get("sha256")
            if isinstance(expected_hash, str) and expected_hash:
                actual_hash = sha256_file(path)
                if actual_hash.lower() != expected_hash.lower():
                    self.add("P0", "RESULT-009", "结果文件 hash 与冻结清单不一致", relative(path, self.root),
                             f"actual={actual_hash}, expected={expected_hash}",
                             "最终结果文件必须与冻结版本一致。", "确认最终版本并更新文件或经人工批准后更新冻结 hash。")
                else:
                    self.passed("RESULT-010", f"结果文件 SHA-256 一致：{relative(path, self.root)}")
            if path.suffix.lower() in {".csv", ".tsv"}:
                summary = csv_summary(path)
                if summary:
                    rows, columns = summary
                    self.check_expected_dimension(item, path, rows, columns)
            elif path.suffix.lower() == ".xlsx":
                summary = xlsx_summary(path)
                if not summary:
                    self.add("P1", "RESULT-004", "XLSX 无法读取", relative(path, self.root), "工作簿结构解析失败",
                             "结果文件必须可正常打开。", "重新保存工作簿并检查损坏。")
                    continue
                expected_sheets = item.get("sheets", [])
                missing_sheets = [name for name in expected_sheets if name not in summary]
                if missing_sheets:
                    self.add("P0", "RESULT-005", "XLSX 缺少指定工作表", relative(path, self.root), ", ".join(missing_sheets),
                             "题目指定工作表名和结构属于硬约束。", "按模板恢复精确工作表名称。")
                expected_rows = item.get("rows", {})
                if isinstance(expected_rows, dict):
                    for sheet, count in expected_rows.items():
                        if sheet in summary and summary[sheet]["rows"] != count:
                            self.add("P0", "RESULT-006", "XLSX 行数不符合清单", f"{relative(path, self.root)}:{sheet}",
                                     f"实际 {summary[sheet]['rows']}，期望 {count}", "结果模板行列属于硬约束。", "核对缺行、重复行和表头口径。")
                expected_columns = item.get("columns", {})
                if isinstance(expected_columns, dict):
                    for sheet, count in expected_columns.items():
                        if sheet in summary and summary[sheet]["columns"] != count:
                            self.add("P0", "RESULT-011", "XLSX 列数不符合清单", f"{relative(path, self.root)}:{sheet}",
                                     f"实际 {summary[sheet]['columns']}，期望 {count}",
                                     "结果模板字段顺序和列数属于硬约束。", "核对缺列、隐藏字段和模板结构。")

    def check_expected_dimension(self, item: dict[str, Any], path: Path, rows: int, columns: int) -> None:
        expected_rows = item.get("rows")
        expected_columns = item.get("columns")
        if isinstance(expected_rows, int) and rows != expected_rows:
            self.add("P0", "RESULT-007", "结果文件行数不符合清单", relative(path, self.root), f"实际 {rows}，期望 {expected_rows}",
                     "结果模板行列属于硬约束。", "核对表头、缺行和重复行。")
        if isinstance(expected_columns, int) and columns != expected_columns:
            self.add("P0", "RESULT-008", "结果文件列数不符合清单", relative(path, self.root), f"实际 {columns}，期望 {expected_columns}",
                     "结果模板字段顺序和列数属于硬约束。", "按模板修正字段。")

    def reproducibility_config(self) -> dict[str, Any]:
        config = self.manifest.get("reproducibility", {})
        return config if isinstance(config, dict) else {}

    def artifact_checks(self) -> list[dict[str, Any]]:
        config = self.reproducibility_config()
        checks = config.get("artifact_checks", self.manifest.get("artifact_checks", []))
        return [item for item in checks if isinstance(item, dict)] if isinstance(checks, list) else []

    def paper_claims(self) -> list[dict[str, Any]]:
        config = self.reproducibility_config()
        claims = config.get("paper_claims", self.manifest.get("paper_claims", []))
        return [item for item in claims if isinstance(item, dict)] if isinstance(claims, list) else []

    def snapshot_overwritten_artifacts(self) -> None:
        for index, item in enumerate(self.artifact_checks()):
            generated = self.resolve_manifest_path(item.get("generated"))
            frozen = self.resolve_manifest_path(item.get("frozen"))
            if generated and frozen and generated == frozen and frozen.exists():
                try:
                    self.artifact_snapshots[index] = frozen.read_bytes()
                except OSError:
                    self.mark_unverified("REPRO-006", f"无法快照将被重写的冻结结果：{relative(frozen, self.root)}")

    def audit_artifact_consistency(self, fresh_run: bool) -> None:
        checks = self.artifact_checks()
        if not checks:
            self.mark_unverified(
                "REPRO-007",
                "manifest 未配置 reproducibility.artifact_checks；运行成功仍不能证明重算结果与冻结结果一致。",
            )
            return
        for index, item in enumerate(checks):
            check_name = str(item.get("id", f"artifact-{index + 1}"))
            generated = self.resolve_manifest_path(item.get("generated"))
            frozen = self.resolve_manifest_path(item.get("frozen"))
            if generated is None or frozen is None:
                self.mark_unverified("REPRO-008", f"{check_name}: generated/frozen 路径未完整配置。")
                continue
            if generated == frozen and not fresh_run and index not in self.artifact_snapshots:
                self.mark_unverified(
                    "REPRO-019",
                    f"{check_name}: generated 与 frozen 为同一路径，未执行入口时不能自比较证明复现一致。",
                )
                continue
            if not generated.exists():
                if fresh_run:
                    self.add("P0", "REPRO-009", "重算关键结果文件缺失", relative(generated, self.root), check_name,
                             "最终入口应生成 manifest 声明的全部关键结果。", "修正输出路径或入口，并重新运行一致性审计。")
                else:
                    self.mark_unverified(
                        "REPRO-009",
                        f"{check_name}: 未执行入口且当前没有 generated 文件，无法判断重算一致性。",
                    )
                continue
            frozen_for_compare = frozen
            temp_path: Path | None = None
            if index in self.artifact_snapshots:
                handle = tempfile.NamedTemporaryFile(suffix=frozen.suffix, delete=False)
                handle.write(self.artifact_snapshots[index])
                handle.close()
                temp_path = Path(handle.name)
                frozen_for_compare = temp_path
            elif not frozen.exists():
                self.add("P0", "REPRO-010", "冻结基准结果文件缺失", relative(frozen, self.root), check_name,
                         "复现审计必须有不可歧义的冻结基准。", "补充冻结结果或 SHA-256 基准后重新审计。")
                continue
            method = str(item.get("method", "hash")).lower()
            try:
                if method == "hash":
                    generated_hash = sha256_file(generated)
                    frozen_hash = sha256_file(frozen_for_compare)
                    if generated_hash != frozen_hash:
                        self.add("P0", "REPRO-011", "重算文件与冻结文件 hash 不一致", check_name,
                                 f"generated={generated_hash}, frozen={frozen_hash}",
                                 "模型—代码—结果文件必须来自同一冻结版本。", "定位差异来源；不得只因程序 exit=0 判复现通过。")
                    else:
                        self.passed("REPRO-012", f"{check_name}: 重算文件与冻结文件 SHA-256 一致")
                elif method == "table":
                    raw_keys = item.get("key_columns", [])
                    raw_ignored = item.get("ignore_columns", [])
                    key_columns = [str(value) for value in raw_keys] if isinstance(raw_keys, list) else None
                    ignore_columns = [str(value) for value in raw_ignored] if isinstance(raw_ignored, list) else None
                    comparison = compare_tables(
                        generated,
                        frozen_for_compare,
                        sheet=str(item.get("sheet")) if item.get("sheet") else None,
                        atol=float(item.get("atol", 1e-8)),
                        rtol=float(item.get("rtol", 1e-6)),
                        key_columns=key_columns or None,
                        ignore_columns=ignore_columns or None,
                    )
                    if comparison is None:
                        self.mark_unverified("REPRO-013", f"{check_name}: 无法解析表格进行容差比较。")
                    elif not comparison["equal"]:
                        self.add("P0", "REPRO-014", "重算表格与冻结结果超出容差", check_name,
                                 json.dumps(comparison, ensure_ascii=False),
                                 "关键结果的行列、键和数值必须在声明容差内一致。", "按首个差异追踪模型参数、随机种子、代码版本和导出逻辑。")
                    else:
                        self.passed("REPRO-015", f"{check_name}: 表格结构与数值在容差内一致")
                else:
                    self.mark_unverified("REPRO-016", f"{check_name}: 不支持 method={method!r}，仅支持 hash/table。")
            finally:
                if temp_path:
                    temp_path.unlink(missing_ok=True)

    def audit_paper_claim_consistency(self) -> None:
        claims = self.paper_claims()
        if not claims:
            self.mark_unverified(
                "CHAIN-001",
                "manifest 未配置 paper_claims；正文数字与结果文件的一致性需在反向证据链中人工核对。",
            )
            return
        text = self.paper_text if self.paper_text_extractable else self.source_text
        if not text_is_extractable(text):
            self.mark_unverified("CHAIN-002", "论文文本不可提取；paper_claims 需视觉/人工核对。")
            return
        for index, claim in enumerate(claims):
            claim_id = str(claim.get("id", f"claim-{index + 1}"))
            pattern = claim.get("paper_pattern")
            if not isinstance(pattern, str) or not pattern:
                self.mark_unverified("CHAIN-003", f"{claim_id}: 缺少 paper_pattern。")
                continue
            try:
                matches = list(re.finditer(pattern, text, re.I | re.S))
            except re.error as exc:
                self.mark_unverified("CHAIN-004", f"{claim_id}: paper_pattern 无效：{exc}")
                continue
            group = int(claim.get("group", 1))
            values: list[float] = []
            for match in matches:
                try:
                    number = parse_number(match.group(group))
                except IndexError:
                    number = None
                if number is not None:
                    values.append(number * float(claim.get("paper_scale", 1.0)))
            unique_values = sorted({round(value, 15) for value in values})
            if len(unique_values) != 1:
                self.mark_unverified(
                    "CHAIN-005",
                    f"{claim_id}: 正文正则未唯一定位一个数值（命中 {len(values)} 个，唯一值 {unique_values[:8]}）。",
                )
                continue
            paper_value = unique_values[0]
            result_spec = claim.get("result")
            expected: Any = claim.get("expected")
            location = "manifest expected"
            if isinstance(result_spec, dict):
                extracted = structured_value(self.root, result_spec)
                if extracted is None:
                    self.mark_unverified("CHAIN-006", f"{claim_id}: 无法从结果文件读取对照值。")
                    continue
                expected, location = extracted
            result_value = parse_number(expected)
            if result_value is None:
                self.mark_unverified("CHAIN-007", f"{claim_id}: 对照值不是可比较数值：{expected!r}")
                continue
            result_value *= float(claim.get("result_scale", 1.0))
            atol = float(claim.get("atol", 1e-8))
            rtol = float(claim.get("rtol", 1e-6))
            if not math.isclose(paper_value, result_value, abs_tol=atol, rel_tol=rtol):
                self.add("P0", "CHAIN-008", "正文数字与结果文件不一致", claim_id,
                         f"paper={paper_value}, result={result_value}, source={location}, atol={atol}, rtol={rtol}",
                         "模型—代码—结果文件—图表—正文数字必须来自同一冻结结果。",
                         "沿结果生成路径定位差异，统一结果文件、图表和正文后重新审计。")
            else:
                self.passed("CHAIN-009", f"{claim_id}: 正文数字与结果值在容差内一致")

    def audit_anonymity_and_hygiene(self) -> None:
        forbidden_terms = [str(term) for term in self.manifest.get("forbidden_terms", []) if str(term).strip()]
        identity_terms = self.manifest.get("identity_terms", {})
        if isinstance(identity_terms, dict):
            for values in identity_terms.values():
                if isinstance(values, list):
                    forbidden_terms.extend(str(value) for value in values if str(value).strip())
                elif isinstance(values, str) and values.strip():
                    forbidden_terms.append(values)
        elif isinstance(identity_terms, list):
            forbidden_terms.extend(str(value) for value in identity_terms if str(value).strip())
        forbidden_terms.extend(self.args.forbidden_term or [])
        forbidden_terms = list(dict.fromkeys(forbidden_terms))

        personal_path_patterns = [
            ("Windows 用户路径", re.compile(r"[A-Za-z]:\\Users\\([^\\\s\"']+)[^\s\"']*", re.I)),
            ("macOS 用户路径", re.compile(r"/Users/([^/\s\"']+)[^\s\"']*")),
            ("Linux 用户路径", re.compile(r"/home/([^/\s\"']+)[^\s\"']*")),
        ]
        generic_users = {"user", "users", "username", "public", "shared", "runner", "ubuntu"}
        email_pattern = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I)
        github_pattern = re.compile(r"https?://github\.com/[A-Za-z0-9_.-]+(?:/[A-Za-z0-9_.-]+)?", re.I)
        environment_path_pattern = re.compile(r"/(?:root|workspace|tmp)/[^\s\"']+")
        term_hits: dict[str, list[str]] = {term: [] for term in forbidden_terms}
        personal_path_hits: list[tuple[str, str, str]] = []
        external_link_hits: list[tuple[str, str]] = []
        environment_path_hits: list[tuple[str, str]] = []
        metadata_candidates: list[tuple[str, str, str]] = []

        scan_files = self.files
        if self.args.mode == "paper-only":
            explicit_scope = {
                path.resolve()
                for path in [self.paper, self.source, *(path for path, _ in self.source_units)]
                if path is not None and path.exists()
            }
            scan_files = (
                [path for path in self.files if path.resolve() in explicit_scope]
                if explicit_scope
                else [path for path in self.files if path.suffix.lower() in PAPER_EXTENSIONS]
            )
        elif self.args.mode == "visual-only":
            scan_files = [self.paper] if self.paper and self.paper.exists() else []

        for path in scan_files:
            if any(part in CACHE_DIRS for part in path.parts):
                continue
            if path.resolve() == self.manifest_path:
                continue
            text = ""
            if path.suffix.lower() in TEXT_EXTENSIONS:
                text = safe_read_text(path)
            elif path.suffix.lower() == ".xlsx":
                text = xlsx_text(path)
            elif path.suffix.lower() == ".docx":
                text = docx_text(path) or ""
            elif path.suffix.lower() == ".pdf" and path.stat().st_size <= 20 * 1024 * 1024:
                text = pdf_text(path) or ""
            metadata = file_metadata(path)
            metadata_text = "\n".join(f"{key}: {value}" for key, value in metadata.items())
            combined = path.name + "\n" + text + "\n" + metadata_text
            for term in forbidden_terms:
                if term and term in combined:
                    term_hits[term].append(relative(path, self.root))
            for label, pattern in personal_path_patterns:
                match = pattern.search(combined)
                if match and match.group(1).lower() not in generic_users:
                    personal_path_hits.append((relative(path, self.root), label, match.group(0)))
            external_link_hits.extend((relative(path, self.root), value) for value in email_pattern.findall(combined))
            external_link_hits.extend((relative(path, self.root), match.group(0)) for match in github_pattern.finditer(combined))
            environment_path_hits.extend(
                (relative(path, self.root), match.group(0)) for match in environment_path_pattern.finditer(combined)
            )
            for key, value in metadata.items():
                key_lower = key.lower()
                if any(token in key_lower for token in ["creator", "author", "lastmodifiedby", "artist", "copyright", "owner"]):
                    if not re.search(r"latex|microsoft|libreoffice|wps|adobe|matplotlib|python|openpyxl|anonymous", value, re.I):
                        metadata_candidates.append((relative(path, self.root), key, value))

        for term, locations in term_hits.items():
            if not locations:
                continue
            self.add("P0", "ANON-001", "明确命中项目身份禁词", ", ".join(sorted(set(locations))[:12]), term,
                     "2026 S0：论文、代码、文件名、元数据和支撑材料不得泄露队伍身份。",
                     "删除或匿名替换该身份信息，清理文件元数据后重新全项目扫描。")
        for location, label, value in personal_path_hits[:20]:
            self.add("P1", "ANON-002", f"发现含用户名的{label}", location, value,
                     "个人路径会泄露用户名并破坏复现性。", "改为相对路径，并清理正文、代码、日志和元数据中的原路径。")
        if environment_path_hits:
            examples = " | ".join(f"{location}: {value}" for location, value in environment_path_hits[:8])
            self.warn("P2", "ANON-003", "发现构建环境绝对路径", str(self.root), examples,
                      "绝对路径不一定泄露队伍身份，但会降低支撑材料可复现性。", "人工确认；能改为相对路径时定点修正。")
        if external_link_hits:
            examples = " | ".join(f"{location}: {value}" for location, value in external_link_hits[:12])
            self.mark_unverified(
                "ANON-004",
                f"检测到邮箱或 GitHub 链接，仅作人工确认，不直接视为身份泄露：{examples}",
            )
        if metadata_candidates:
            examples = " | ".join(
                f"{location} [{key}]={value}" for location, key, value in metadata_candidates[:12]
            )
            self.mark_unverified(
                "ANON-005",
                f"PDF/OOXML/图片元数据存在作者类字段；未命中项目身份词，需人工确认是否应清理：{examples}",
            )
        if self.args.mode in {"full", "recheck"}:
            for path in self.files:
                suffix = path.name.lower()
                if any(suffix.endswith(item) for item in BUILD_SUFFIXES) or any(part in CACHE_DIRS for part in path.parts):
                    self.add("P2", "HYGIENE-001", "支撑工程包含构建垃圾或缓存", relative(path, self.root), path.name,
                             "S1：支撑材料只保留最终可复现所需文件。", "从最终支撑包排除构建垃圾和缓存。", True)
                elif OLD_NAME_RE.search(path.name):
                    self.add("P2", "HYGIENE-002", "发现疑似旧版或备份文件", relative(path, self.root), path.name,
                             "最终支撑材料不得混入旧结果和过期代码。", "确认版本后移出最终提交包。")

    def audit_support(self) -> None:
        declared_support = self.manifest.get("has_support_materials")
        if declared_support not in {True, False, None}:
            self.mark_unverified("SUPPORT-000", "has_support_materials 必须为 true/false；当前值无法解释。")
            declared_support = None
        support_dir = self.resolve_manifest_path(self.manifest.get("support_dir"))
        if support_dir is None:
            candidates = [self.root / "support", self.root / "支撑材料"]
            support_dir = next((path for path in candidates if path.is_dir()), None)
        support_archive = self.resolve_manifest_path(self.manifest.get("support_archive"))
        if support_archive is None:
            archives = [path for path in self.files if path.suffix.lower() in {".zip", ".rar"} and re.search(r"support|支撑", path.stem, re.I)]
            support_archive = archives[0] if len(archives) == 1 else None

        has_dir = bool(support_dir and support_dir.is_dir())
        has_archive = bool(support_archive and support_archive.exists())
        declaration_text = self.source_text or self.paper_text
        no_support_statement = "本论文没有支撑材料" in compact_text(declaration_text)
        no_program_statement = "本论文没有用到程序" in compact_text(declaration_text)

        if declared_support is False:
            if self.manifest.get("ai_used") is True:
                self.add("P0", "SUPPORT-009", "支撑材料声明与 AI 使用要求冲突", relative(self.manifest_path, self.root),
                         "has_support_materials=false 且 ai_used=true", "使用 AI 时支撑材料必须包含 AI 工具使用详情.pdf。",
                         "将 has_support_materials 改为 true，并提交 AI 工具使用详情及必要材料。")
            if has_dir or has_archive:
                self.add("P1", "SUPPORT-010", "支撑材料声明与实际文件冲突", str(self.root),
                         f"has_support_materials=false, dir={has_dir}, archive={has_archive}",
                         "清单、附录和最终提交包必须一致。", "确认是否有支撑材料，并统一 manifest、附录和交卷文件。")
            elif no_support_statement:
                self.passed("SUPPORT-011", "清单声明无支撑材料，附录检测到“本论文没有支撑材料”")
            elif text_is_extractable(declaration_text):
                self.add("P1", "SUPPORT-012", "无支撑材料但附录未确认规定说明", relative(self.source or self.paper, self.root),
                         "has_support_materials=false，未检测到“本论文没有支撑材料”",
                         "S0：确实没有支撑材料时，附录应明确说明。", "在附录加入准确说明并重新生成最终 PDF。")
            else:
                self.mark_unverified("SUPPORT-013", "清单声明无支撑材料，但论文文本不可提取；需视觉确认附录规定说明。")
            if self.manifest.get("uses_programs") is False and not no_program_statement:
                self.mark_unverified("SUPPORT-014", "uses_programs=false；需确认附录是否写明“本论文没有用到程序”。")
            return

        if declared_support is None and not has_dir and not has_archive:
            if no_support_statement:
                self.passed("SUPPORT-015", "附录检测到“本论文没有支撑材料”；未要求 support 文件夹")
            else:
                self.mark_unverified(
                    "SUPPORT-016",
                    "manifest 未声明 has_support_materials，且未发现支撑目录/压缩包；不能直接判缺失，请明确 true/false。",
                )
            return

        if declared_support is None:
            self.mark_unverified(
                "SUPPORT-017",
                "已自动发现支撑材料，但 manifest 未明确 has_support_materials=true；请统一声明、附录和最终提交包。",
            )

        if support_dir and support_dir.is_dir():
            self.passed("SUPPORT-001", f"发现支撑材料目录：{relative(support_dir, self.root)}")
            support_files = [path for path in support_dir.rglob("*") if path.is_file()]
            if not any(re.match(r"README|支撑材料说明", path.name, re.I) for path in support_files):
                self.add("P2", "SUPPORT-008", "支撑材料缺少运行/文件说明", relative(support_dir, self.root),
                         "未发现 README 或支撑材料说明文件", "S1：支撑材料应提供文件清单、运行入口和复现说明。",
                         "补充简洁 README，说明目录、依赖、入口、输出和 AI 详情位置。")
        elif declared_support is True and not has_archive and self.args.mode in {"full", "recheck"}:
            self.add("P1", "SUPPORT-002", "清单声明有支撑材料但未找到材料", str(self.root),
                     "has_support_materials=true，目录和压缩包均不存在",
                     "S0：有程序、外部数据或补充结果时必须提交必要支撑材料。", "提供支撑材料并与附录文件清单对应。")
            self.scope_incomplete = True
        if support_archive and support_archive.exists():
            size = support_archive.stat().st_size
            if support_archive.suffix.lower() not in {".zip", ".rar"}:
                self.add("P0", "SUPPORT-003", "支撑材料压缩格式不合规", relative(support_archive, self.root), support_archive.suffix,
                         "2026 S0：支撑材料为 RAR 或 ZIP。", "使用 WinRAR 生成 RAR/ZIP。")
            if size > 20 * 1024 * 1024:
                self.add("P0", "SUPPORT-004", "支撑材料超过 20 MB", relative(support_archive, self.root), f"{size / 1024 / 1024:.2f} MB",
                         "2026 S0：支撑材料不超过 20 MB。", "删除无关/重复文件或压缩必要资源。")
            if support_archive.suffix.lower() == ".zip":
                try:
                    with zipfile.ZipFile(support_archive) as archive:
                        bad = archive.testzip()
                    if bad:
                        raise zipfile.BadZipFile(f"bad member: {bad}")
                    self.passed("SUPPORT-005", "ZIP 支撑材料可正常读取")
                except zipfile.BadZipFile as exc:
                    self.add("P0", "SUPPORT-006", "ZIP 支撑材料损坏", relative(support_archive, self.root), str(exc),
                             "最终支撑材料必须可正常解压。", "重新压缩并完整测试。")
        elif declared_support is True and self.args.mode in {"full", "recheck"}:
            self.mark_unverified("SUPPORT-007", "未锁定最终 RAR/ZIP；无法核对格式、大小和可解压性。")

    def run_entrypoint(self) -> None:
        if not self.args.run_entrypoint:
            self.mark_unverified("REPRO-001", "未执行最终代码入口；代码可复现性仍需明确授权后核验。")
            return
        config = self.reproducibility_config()
        entrypoint = config.get("entrypoint", self.manifest.get("entrypoint"))
        if not isinstance(entrypoint, list) or not entrypoint or not all(isinstance(item, str) for item in entrypoint):
            self.add("P1", "REPRO-002", "清单缺少安全入口命令数组", relative(self.manifest_path, self.root), repr(entrypoint),
                     "代码执行必须使用明确、无 shell 的最终入口。", "将 entrypoint 写成如 [\"python3\", \"run_all.py\"]。")
            return
        workdir = self.resolve_manifest_path(config.get("working_directory", ".")) or self.root
        try:
            workdir.relative_to(self.root)
        except ValueError:
            self.add("P1", "REPRO-017", "代码工作目录越出项目根目录", str(workdir),
                     "working_directory 必须位于项目内", "复现命令不应依赖项目外个人目录。", "改为项目内相对工作目录。")
            return
        if not workdir.is_dir():
            self.add("P1", "REPRO-018", "代码工作目录不存在", relative(workdir, self.root), "目录不存在",
                     "统一入口必须从明确工作目录运行。", "修正 reproducibility.working_directory。")
            return
        timeout = int(config.get("timeout_seconds", self.manifest.get("timeout_seconds", 300)))
        self.snapshot_overwritten_artifacts()
        try:
            result = run_command(entrypoint, cwd=workdir, timeout=timeout)
        except (OSError, subprocess.TimeoutExpired) as exc:
            self.add("P0", "REPRO-003", "最终代码入口无法完成", str(self.root), str(exc),
                     "核心结果必须可由最终代码重跑。", "修复依赖、路径或运行时间问题后重跑。")
            return
        if result.returncode != 0:
            self.add("P0", "REPRO-004", "最终代码入口运行失败", str(self.root),
                     f"exit={result.returncode}\n{result.stdout[-1500:]}", "核心结果必须可由最终代码重跑。",
                     "根据日志修复错误，并确认使用最终数据和配置。")
        else:
            self.passed("REPRO-005", f"最终代码入口运行成功：{' '.join(entrypoint)}")
            self.entrypoint_succeeded = True

    def determine_script_status(self) -> str:
        if self.not_ready:
            return "NOT_READY_FOR_FINAL_AUDIT"
        if self.scope_incomplete:
            return "INCOMPLETE_AUDIT_SCOPE"
        confirmed_priorities = {
            finding.priority for finding in self.findings if finding.certainty != "HEURISTIC"
        }
        priorities = {finding.priority for finding in self.findings}
        if self.args.mode == "paper-only":
            if confirmed_priorities.intersection({"P0", "P1"}):
                return "PAPER_AUDIT_COMPLETE_WITH_BLOCKERS"
            drawing_deferred = any(
                item["status"] == "DEFERRED_TO_DRAWING" for item in self.deferred
            )
            coverage_pending = any(
                item["status"] != "PASS" for item in self.build_paper_only_coverage()
            )
            if priorities or self.unverified or drawing_deferred or coverage_pending:
                return "PAPER_AUDIT_COMPLETE_WITH_WARNINGS"
            return "PAPER_AUDIT_PASS"
        if "P0" in confirmed_priorities:
            return "SCRIPT_BLOCKERS_FOUND"
        if "P1" in confirmed_priorities:
            return "SCRIPT_PRIORITY_FIXES_FOUND"
        if priorities:
            return "SCRIPT_WARNINGS_FOUND"
        if self.unverified:
            return "SCRIPT_UNVERIFIED_ITEMS"
        return "SCRIPT_CHECKS_PASSED"

    def build_paper_only_coverage(self) -> list[dict[str, str]]:
        """Build a non-skippable automatic coverage scaffold for PAPER_ONLY.

        Semantic rows deliberately remain UNVERIFIED until a reviewer reads the
        paper. This prevents an empty script finding list from being mistaken
        for a completed audit.
        """
        findings_by_id = {item.check_id: item for item in self.findings}
        deferred_by_id = {item["check_id"]: item for item in self.deferred}
        pass_ids = {item["check_id"] for item in self.passes}

        def automatic_row(
            module: str,
            finding_ids: set[str],
            pass_candidates: set[str] | None = None,
            pending: str = "需要人工语义审查；脚本不得代判 PASS。",
            deferred_ids: set[str] | None = None,
        ) -> dict[str, str]:
            matched = [findings_by_id[item] for item in finding_ids if item in findings_by_id]
            confirmed_high = [
                item for item in matched if item.priority in {"P0", "P1"} and item.certainty != "HEURISTIC"
            ]
            if confirmed_high:
                evidence = "；".join(f"{item.check_id} {item.title}" for item in confirmed_high[:4])
                return {"module": module, "status": "FAIL", "evidence": evidence}
            if matched:
                evidence = "；".join(f"{item.check_id} {item.title}" for item in matched[:4])
                return {"module": module, "status": "WARN", "evidence": evidence}
            deferred_matches = [
                deferred_by_id[item] for item in (deferred_ids or set()) if item in deferred_by_id
            ]
            if deferred_matches:
                evidence = "；".join(
                    f"{item['status']} {item['item']}" for item in deferred_matches[:4]
                )
                return {"module": module, "status": "WARN", "evidence": evidence}
            if pass_candidates and pass_ids.intersection(pass_candidates):
                evidence = "；".join(sorted(pass_ids.intersection(pass_candidates)))
                return {"module": module, "status": "PASS", "evidence": evidence}
            return {"module": module, "status": "UNVERIFIED", "evidence": pending}

        rows = [
            automatic_row(
                "段落与排版源码",
                {"INDENT-001", "INDENT-002", "INDENT-006", "INDENT-007"},
                {"INDENT-003"},
                "缺少可确认的统一 parindent 证据，或仍需检查类文件/最终 PDF。",
            ),
            automatic_row(
                "AI/模板化语言", {"LANG-001", "LANG-002", "LANG-006"}, pending="需通读正文判断模板化表达，词频不能代判。"
            ),
            automatic_row(
                "长句与可读性", {"LANG-003", "LANG-004", "LANG-005"}, pending="需通读代表段落；无自动候选不等于全文 PASS。"
            ),
            automatic_row("段落/小节过渡", set(), pending="需逐个关键边界核对前段结论与后段任务。"),
            automatic_row("模型百科式表达", set(), pending="需逐个主模型核对本题特征—模型结构—适配理由。"),
            automatic_row(
                "术语一致性", {"TERM-001"}, pending="需建立规范术语表并全文核对；脚本词典未命中不能代判 PASS。"
            ),
            automatic_row("问题分析与最终模型边界", set(), pending="需反向定位首次确定性选模表述。"),
            automatic_row(
                "模型选择证据链", {"SOURCE-001", "SOURCE-002"},
                pending="需核对基线、候选、统一比较、最终选择证据，并反查题面事实与模型自设参数的来源。",
            ),
            automatic_row("跨问接口", set(), pending="仅凭 PAPER_ONLY 可审文字接口；代码实现仍可能 UNVERIFIED。"),
            automatic_row("模型输出→最终决策映射", set(), pending="需逐问追踪输出、阈值/规则与最终答案。"),
            automatic_row(
                "图表正文叙事", {"FIG-002", "FIG-003", "FIG-006"},
                deferred_ids={"DRAWING-001", "LATEX-009", "LATEX-011"},
                pending="需逐张检查图前口径与图后量化解释。",
            ),
            automatic_row("公式与符号", set(), pending="需人工核对符号定义、单位、公式来源与前后一致。"),
            automatic_row(
                "引用格式",
                {"REF-004", "REF-007", "REF-009", "LATEX-012"},
                pending="即使 key 对应正常，仍需核对最终显示样式、首次引用位置与稳定编译，脚本不独立判 PASS。",
            ),
            automatic_row(
                "引用跳转",
                {"REF-LINK-004", "REF-LINK-005", "REF-LINK-007"},
                pending="需按唯一 citation key 点击核对 PDF 目的条目；自动注释扫描不能独立判 PASS。",
            ),
            automatic_row(
                "引用语义正确性", {"LANG-006"}, pending="必须填写引用完整性矩阵；无原文证据时保持 UNVERIFIED。"
            ),
            automatic_row(
                "参考文献数量与正文对应关系",
                {"REF-002", "REF-004", "REF-005"},
                {"REF-008"},
                "需确认文献数量、正文 citation keys 和孤儿条目。",
            ),
        ]
        assert tuple(row["module"] for row in rows) == PAPER_ONLY_CORE_MODULES
        return rows

    def build_report(self) -> tuple[dict[str, Any], str]:
        counts = {priority: 0 for priority in ["P0", "P1", "P2", "P3"]}
        for finding in self.findings:
            counts[finding.priority] = counts.get(finding.priority, 0) + 1
        status = self.determine_script_status()
        audit_coverage = self.build_paper_only_coverage() if self.args.mode == "paper-only" else []
        payload = {
            "script_status": status,
            "mode": self.args.mode.upper(),
            "project": str(self.root),
            "paper": relative(self.paper, self.root),
            "source": relative(self.source, self.root),
            "visual_output": self.visual_output,
            "counts": counts,
            "pass_count": len(self.passes),
            "unverified_count": len(self.unverified),
            "deferred": self.deferred,
            "findings": [asdict(item) for item in self.findings],
            "passes": self.passes,
            "unverified": self.unverified,
            "source_attribution_candidates": self.source_attributions,
            "audit_coverage": audit_coverage,
            "notice": "脚本状态不是最终可提交结论；仍需语义审查与最终 PDF 逐页视觉审查。",
        }
        lines = [
            "# SCRIPT AUDIT REPORT",
            "",
            f"- Script Status: `{status}`",
            f"- Mode: `{self.args.mode.upper()}`",
            f"- Paper: `{payload['paper']}`",
            f"- Source: `{payload['source']}`",
            f"- P0/P1/P2/P3: {counts['P0']}/{counts['P1']}/{counts['P2']}/{counts['P3']}",
            f"- PASS: {len(self.passes)}",
            f"- UNVERIFIED: {len(self.unverified)}",
            f"- DEFERRED / OUT_OF_SCOPE: {len(self.deferred)}",
            "",
            "> 脚本状态不是最终可提交结论；仍需语义审查与最终 PDF 逐页视觉审查。",
            "> P2/P3 为 WARN；HEURISTIC 与 UNVERIFIED 不得自动升级为 FAIL。",
            "",
        ]
        if audit_coverage:
            lines += [
                "## PAPER_ONLY 强制执行顺序",
                "",
                "1. 题目覆盖 + 模型/结果闭环",
                "2. 跨问接口 + 模型输出→最终决策",
                "3. 模型选择 + 题面事实/结构参数来源反查 + 验证",
                "4. 连接过渡 + AI/百科式语言 + 术语一致性",
                "5. 引用格式/跳转/语义正确性",
                "6. 压缩冗余 + 结构",
                "7. 段首缩进等源码排版",
                "",
                "> 自动证据收集可先完成，但不得按脚本命中顺序替代上述人工/语义审查顺序。",
                "",
                "## Audit Coverage（PAPER_ONLY 自动预审）",
                "",
                "> 该表强制保留所有核心模块。语义模块在人工通读前必须为 UNVERIFIED，不能因脚本未命中而写 PASS。",
                "",
                "| 核心模块 | 状态 | 自动证据/待核验事项 |",
                "|---|---|---|",
            ]
            lines.extend(
                f"| {item['module']} | {item['status']} | {item['evidence']} |" for item in audit_coverage
            )
            lines += [
                "",
                "## 术语一致性矩阵（PAPER_ONLY 待语义审查填写）",
                "",
                "> 自动脚本只列已知变体共现候选；必须人工确认是否同义。未命中不代表全文术语一致。",
                "",
                "| 技术概念 | 出现变体 | 位置 | 是否同义 | 统一名称/区分定义 | 状态 |",
                "|---|---|---|---|---|---|",
                "| 待填写 | 待填写 | 待填写 | UNVERIFIED | 待填写 | UNVERIFIED |",
                "",
                "## 题面事实与模型设定来源矩阵（并入模型选择证据链）",
                "",
                "> 只列涉及模型参数、分布、相关结构、关系矩阵、模拟边、阈值或权重的来源性声明。自动候选不能代替原题/附件反查。",
                "",
                "| 参数/关系 | 正文声称来源 | 实际来源/证据 | 状态 |",
                "|---|---|---|---|",
            ]
            if self.source_attributions:
                lines.extend(
                    f"| {', '.join(item['objects'])} | {item['claim']}（行 {item['line']}） | UNVERIFIED：待对照原题/附件/冻结赛题解读 | UNVERIFIED |"
                    for item in self.source_attributions
                )
            else:
                lines.append(
                    "| 待人工通读补全 | 脚本未定位候选，不代表已通过 | 需反查原题/附件 | UNVERIFIED |"
                )
            lines += [
                "",
                "## 引用完整性矩阵（PAPER_ONLY 待语义审查填写）",
                "",
                "> 必须覆盖全部正文引用 occurrence；只有题名/主题、没有文献全文时，原文支持和整行状态不得为 PASS。",
                "",
                "| 正文位置 | 引用编号/key | 当前句子声称什么 | 主题匹配 | 原文支持 | 引用位置 | PDF 跳转 | 状态 |",
                "|---|---|---|---|---|---|---|---|",
                "| 待填写 | 待填写 | 待填写 | UNVERIFIED | UNVERIFIED | UNVERIFIED | UNVERIFIED | UNVERIFIED |",
                "",
            ]
        if self.deferred:
            lines += [
                "## Deferred to FULL / Drawing / Out of Scope",
                "",
                "| 检查项 | 阶段状态 | 原因 | 下一模式 |",
                "|---|---|---|---|",
            ]
            lines.extend(
                f"| {item['item']} | {item['status']} | {item['reason']} | {item['next_mode']} |"
                for item in self.deferred
            )
            lines.append("")
        for priority in ["P0", "P1", "P2", "P3"]:
            items = [finding for finding in self.findings if finding.priority == priority]
            label = f"{priority} (WARN)" if priority in {"P2", "P3"} else priority
            lines += [f"## {label}", ""]
            if not items:
                lines += ["- 无", ""]
                continue
            for item in items:
                lines += [
                    f"### [{item.priority}][{item.check_id}] {item.title}", "",
                    f"- 位置：{item.location}",
                    f"- 证据：{item.evidence}",
                    f"- 规则：{item.rule}",
                    f"- 建议：{item.recommendation}",
                    f"- 判断类型：{item.certainty}",
                    f"- 阶段状态：{item.stage_status}",
                    f"- 确定性自动修复：{'是' if item.auto_fixable else '否'}",
                    "",
                ]
        lines += ["## UNVERIFIED", ""]
        if self.unverified:
            lines.extend(f"- [{item['check_id']}] {item['message']}" for item in self.unverified)
        else:
            lines.append("- 无")
        lines += ["", "## PASS 摘要", ""]
        lines.extend(f"- [{item['check_id']}] {item['message']}" for item in self.passes)
        return payload, "\n".join(lines) + "\n"

    def execute(self) -> int:
        self.load_manifest()
        self.check_readiness()
        if self.not_ready:
            payload, report = self.build_report()
            if self.args.json_out:
                Path(self.args.json_out).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
            if self.args.md_out:
                Path(self.args.md_out).write_text(report, encoding="utf-8")
            print(report, end="")
            return 2
        self.discover_files()
        self.load_document_text()
        if self.args.mode == "paper-only":
            drawing_text = self.latex_bundle_text() or self.source_text or self.paper_text
        elif self.args.mode == "visual-only":
            drawing_text = self.paper_text
        else:
            drawing_text = self.paper_text or self.latex_bundle_text() or self.source_text
        if drawing_text:
            self.audit_drawing_placeholders(drawing_text)
        self.audit_paper_file()
        self.render_visual_assets()
        if self.args.mode != "visual-only":
            self.audit_source_structure()
            self.audit_pdf_reference_links()
        self.audit_anonymity_and_hygiene()
        if self.args.mode in {"full", "recheck"}:
            self.audit_required_files()
            self.audit_support()
            self.run_entrypoint()
            self.audit_artifact_consistency(self.entrypoint_succeeded)
            self.audit_paper_claim_consistency()
        payload, report = self.build_report()
        if self.args.json_out:
            Path(self.args.json_out).write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        if self.args.md_out:
            Path(self.args.md_out).write_text(report, encoding="utf-8")
        print(report, end="")
        if self.not_ready or self.scope_incomplete:
            return 2
        if any(
            finding.priority in {"P0", "P1"} and finding.certainty != "HEURISTIC"
            for finding in self.findings
        ):
            return 1
        return 0


def parse_args(argv: Iterable[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Deterministic pre-audit for a frozen CUMCM C submission.")
    parser.add_argument("project", help="Project root directory")
    parser.add_argument("--mode", choices=["full", "paper-only", "visual-only", "recheck"], default="full")
    parser.add_argument("--paper", help="Final PDF/Word path, relative to project root unless absolute")
    parser.add_argument("--manifest", help="Path to audit_manifest.json")
    parser.add_argument("--run-entrypoint", action="store_true", help="Run manifest entrypoint; requires explicit user authorization")
    parser.add_argument("--compile-latex", action="store_true", help="Compile final LaTeX into a temporary build directory")
    parser.add_argument("--render-pdf", action="store_true", help="Render all final PDF pages and contact sheets")
    parser.add_argument("--visual-dir", help="Visual audit output directory, relative to project root unless absolute")
    parser.add_argument("--forbidden-term", action="append", default=[], help="Project-specific anonymity term; may repeat")
    parser.add_argument("--json-out", help="Write machine-readable findings JSON")
    parser.add_argument("--md-out", help="Write Markdown script report")
    return parser.parse_args(argv)


def main() -> int:
    args = parse_args()
    return SubmissionAudit(args).execute()


if __name__ == "__main__":
    raise SystemExit(main())
