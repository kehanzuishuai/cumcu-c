from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
EXPERIMENT_ROOT = Path(__file__).resolve().parent
TRAJECTORY = EXPERIMENT_ROOT / "results" / "q1_lp_milp_trajectory.csv"
SUMMARY = EXPERIMENT_ROOT / "results" / "q1_lp_milp_crosscheck_summary.json"
OUTPUT = EXPERIMENT_ROOT / "validation" / "q1_lp_milp_independent_recompute.json"
ATTACHMENT = PROJECT_ROOT / "source_materials" / "C题" / "附件" / "附件1.xlsx"

ETA_C = 0.9
ETA_D = 0.9
SOC_INITIAL = 6000.0
SOC_TERMINAL = 6000.0
SOC_MIN = 1200.0
SOC_MAX = 10800.0
FLOW_MAX = 5000.0 / 6.0
TOL = 1e-6
OBJECTIVE_TOL = 1e-5


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_source() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    workbook = load_workbook(ATTACHMENT, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    if len(rows) != 144:
        raise ValueError(f"附件1应为144行，实际为{len(rows)}")
    price = np.asarray([float(row[1]) for row in rows])
    load_kwh = np.asarray([float(row[2]) / 6.0 for row in rows])
    pv_kwh = np.asarray([float(row[3]) / 6.0 for row in rows])
    return price, load_kwh, pv_kwh


def load_trajectory() -> list[dict[str, str]]:
    with TRAJECTORY.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    if len(rows) != 144:
        raise ValueError(f"实验轨迹应为144行，实际为{len(rows)}")
    return rows


def array(rows: list[dict[str, str]], name: str) -> np.ndarray:
    return np.asarray([float(row[name]) for row in rows])


def validate(prefix: str, rows: list[dict[str, str]], price: np.ndarray, load_kwh: np.ndarray, pv_kwh: np.ndarray):
    q = array(rows, f"{prefix}_purchase_kwh")
    charge = array(rows, f"{prefix}_charge_kwh")
    discharge = array(rows, f"{prefix}_discharge_kwh")
    spill = array(rows, f"{prefix}_spill_kwh")
    soc_close = array(rows, f"{prefix}_soc_close_kwh")
    soc_open = np.concatenate(([SOC_INITIAL], soc_close[:-1]))
    balance = q + pv_kwh + discharge - load_kwh - charge - spill
    state = soc_close - soc_open - ETA_C * charge + discharge / ETA_D
    simultaneous = np.minimum(charge, discharge)
    metrics = {
        "purchase_cost_yuan": float(np.dot(price, q)),
        "purchase_kwh": float(np.sum(q)),
        "charge_kwh": float(np.sum(charge)),
        "discharge_kwh": float(np.sum(discharge)),
        "throughput_kwh": float(np.sum(charge + discharge)),
        "balance_max_abs_kwh": float(np.max(np.abs(balance))),
        "state_max_abs_kwh": float(np.max(np.abs(state))),
        "soc_min_kwh": float(min(SOC_INITIAL, np.min(soc_close))),
        "soc_max_kwh": float(max(SOC_INITIAL, np.max(soc_close))),
        "soc_terminal_abs_kwh": float(abs(soc_close[-1] - SOC_TERMINAL)),
        "charge_max_kwh_per_slot": float(np.max(charge)),
        "discharge_max_kwh_per_slot": float(np.max(discharge)),
        "simultaneous_positive_slot_count": int(np.sum(simultaneous > 1e-7)),
        "simultaneous_max_kwh": float(np.max(simultaneous)),
        "nonnegative_max_violation_kwh": float(max(0.0, -np.min(np.concatenate([q, charge, discharge, spill])))),
    }
    metrics["maximum_constraint_residual_kwh"] = max(
        metrics["balance_max_abs_kwh"], metrics["state_max_abs_kwh"],
        metrics["soc_terminal_abs_kwh"], metrics["nonnegative_max_violation_kwh"],
        max(0.0, SOC_MIN - metrics["soc_min_kwh"]), max(0.0, metrics["soc_max_kwh"] - SOC_MAX),
        max(0.0, metrics["charge_max_kwh_per_slot"] - FLOW_MAX),
        max(0.0, metrics["discharge_max_kwh_per_slot"] - FLOW_MAX),
    )
    pass_flags = {
        "balance": metrics["balance_max_abs_kwh"] <= TOL,
        "state": metrics["state_max_abs_kwh"] <= TOL,
        "soc": metrics["soc_min_kwh"] >= SOC_MIN - TOL and metrics["soc_max_kwh"] <= SOC_MAX + TOL and metrics["soc_terminal_abs_kwh"] <= TOL,
        "power": metrics["charge_max_kwh_per_slot"] <= FLOW_MAX + TOL and metrics["discharge_max_kwh_per_slot"] <= FLOW_MAX + TOL,
        "nonnegative": metrics["nonnegative_max_violation_kwh"] <= TOL,
        "mutual_exclusion": metrics["simultaneous_positive_slot_count"] == 0,
    }
    return {"metrics": metrics, "pass_flags": pass_flags, "all_pass": all(pass_flags.values())}


def main() -> None:
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    rows = load_trajectory()
    price, load_kwh, pv_kwh = load_source()
    lp = validate("lp", rows, price, load_kwh, pv_kwh)
    mixed = validate("milp", rows, price, load_kwh, pv_kwh)
    differences = {
        key: float(mixed["metrics"][key] - lp["metrics"][key])
        for key in ("purchase_cost_yuan", "purchase_kwh", "charge_kwh", "discharge_kwh", "throughput_kwh")
    }
    summary_lp = summary["comparison"]["lp"]
    summary_milp = summary["comparison"]["milp"]
    reconciliation = {
        "lp_cost_vs_summary_abs_yuan": abs(lp["metrics"]["purchase_cost_yuan"] - summary_lp["secondary_solution_cost_yuan"]),
        "milp_cost_vs_summary_abs_yuan": abs(mixed["metrics"]["purchase_cost_yuan"] - summary_milp["secondary_solution_cost_yuan"]),
        "lp_purchase_vs_summary_abs_kwh": abs(lp["metrics"]["purchase_kwh"] - summary_lp["daily_purchase_kwh"]),
        "milp_purchase_vs_summary_abs_kwh": abs(mixed["metrics"]["purchase_kwh"] - summary_milp["daily_purchase_kwh"]),
        "lp_throughput_vs_summary_abs_kwh": abs(lp["metrics"]["throughput_kwh"] - summary_lp["throughput_kwh"]),
        "milp_throughput_vs_summary_abs_kwh": abs(mixed["metrics"]["throughput_kwh"] - summary_milp["throughput_kwh"]),
    }
    objective_equal = abs(differences["purchase_cost_yuan"]) <= OBJECTIVE_TOL
    throughput_equal = abs(differences["throughput_kwh"]) <= TOL
    protected_expected = json.loads(
        (EXPERIMENT_ROOT / "validation" / "q1_lp_milp_crosscheck_validation.json").read_text(encoding="utf-8")
    )["protected_sha256_after"]
    protected_actual = {name: sha256(PROJECT_ROOT / name) for name in protected_expected}
    protected_unchanged = protected_expected == protected_actual
    all_pass = bool(
        lp["all_pass"] and mixed["all_pass"] and objective_equal and throughput_equal
        and max(reconciliation.values()) <= TOL and protected_unchanged
    )
    payload = {
        "schema_version": 1,
        "experiment": "Q1 LP-MILP independent recomputation",
        "status": "PASS" if all_pass else "FAIL",
        "independence": "reads Attachment 1 and experiment trajectory directly; does not import the LP/MILP solver module",
        "lp": lp,
        "milp": mixed,
        "milp_minus_lp": differences,
        "reconciliation": reconciliation,
        "equivalence_flags": {
            "objective_equal_within_tolerance": objective_equal,
            "throughput_equal_within_tolerance": throughput_equal,
            "both_constraint_sets_pass": lp["all_pass"] and mixed["all_pass"],
            "protected_formal_files_unchanged": protected_unchanged,
        },
        "source_sha256": sha256(ATTACHMENT),
        "trajectory_sha256": sha256(TRAJECTORY),
        "summary_sha256": sha256(SUMMARY),
        "protected_formal_files_sha256": protected_actual,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    manifest_path = EXPERIMENT_ROOT / "q1_lp_milp_crosscheck_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["independent_validator_sha256"] = sha256(Path(__file__))
    manifest["outputs_sha256"][str(OUTPUT.relative_to(PROJECT_ROOT))] = sha256(OUTPUT)
    manifest["independent_recompute_status"] = payload["status"]
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": payload["status"], "differences": differences, "lp_max_residual": lp["metrics"]["maximum_constraint_residual_kwh"], "milp_max_residual": mixed["metrics"]["maximum_constraint_residual_kwh"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
