from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DETAIL = ROOT / "results" / "q1_detail.csv"
SUMMARY = ROOT / "results" / "q1_summary.json"
STORAGE = ROOT / "results" / "q1_storage_4h.csv"
SOURCE = ROOT / "source_materials" / "C题" / "附件" / "附件1.xlsx"
OUT = ROOT / "validation" / "q1_independent_recompute.json"

ETA_C = 0.9
ETA_D = 0.9
SOC_MIN = 1200.0
SOC_MAX = 10800.0
FLOW_MAX = 5000.0 / 6.0
TOL = 2e-6


def f(row: dict[str, str], key: str) -> float:
    return float(row[key])


def main() -> None:
    with DETAIL.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    with STORAGE.open("r", encoding="utf-8-sig", newline="") as stream:
        storage_rows = list(csv.DictReader(stream))
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))

    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        source_rows = list(ws.iter_rows(min_row=2, values_only=True))
    finally:
        wb.close()

    balance_residuals = []
    state_residuals = []
    source_differences = []
    soc_values = []
    purchase_total = 0.0
    cost_total = 0.0
    charge_total = 0.0
    discharge_total = 0.0
    simultaneous = 0
    interval_labels = []
    for i, row in enumerate(rows):
        q = f(row, "normal_purchase_kwh")
        pv = f(row, "pv_kwh")
        discharge = f(row, "discharge_bus_kwh")
        load = f(row, "load_kwh")
        charge = f(row, "charge_bus_kwh")
        spill = f(row, "pv_spill_or_surplus_kwh")
        soc_open = f(row, "soc_open_kwh")
        soc_close = f(row, "soc_close_kwh")
        price = f(row, "price_yuan_per_kwh")
        balance_residuals.append(q + pv + discharge - load - charge - spill)
        state_residuals.append(soc_close - soc_open - ETA_C * charge + discharge / ETA_D)
        source_differences.extend([
            abs(f(row, "price_yuan_per_kwh") - float(source_rows[i][1])),
            abs(f(row, "load_kw") - float(source_rows[i][2])),
            abs(f(row, "pv_kw") - float(source_rows[i][3])),
        ])
        soc_values.extend([soc_open, soc_close])
        purchase_total += q
        cost_total += price * q
        charge_total += charge
        discharge_total += discharge
        simultaneous += int(charge > 1e-7 and discharge > 1e-7)
        interval_labels.append(row["interval_label"])

    storage_charge_total = sum(float(row["charge_kwh"]) for row in storage_rows)
    storage_discharge_total = sum(float(row["discharge_kwh"]) for row in storage_rows)
    expected_labels = []
    for i in range(144):
        start = i * 10
        end = (i + 1) * 10
        sf = f"{start // 60:02d}:{start % 60:02d}"
        ef = "24:00" if end == 1440 else f"{end // 60:02d}:{end % 60:02d}"
        expected_labels.append(f"{sf}-{ef}")

    metrics = {
        "detail_row_count": len(rows),
        "source_row_count": len(source_rows),
        "max_source_value_difference": max(source_differences),
        "max_balance_residual_kwh": max(abs(value) for value in balance_residuals),
        "max_state_residual_kwh": max(abs(value) for value in state_residuals),
        "soc_min_kwh": min(soc_values),
        "soc_max_kwh": max(soc_values),
        "soc_open_kwh": f(rows[0], "soc_open_kwh"),
        "soc_close_kwh": f(rows[-1], "soc_close_kwh"),
        "charge_max_kwh": max(f(row, "charge_bus_kwh") for row in rows),
        "discharge_max_kwh": max(f(row, "discharge_bus_kwh") for row in rows),
        "simultaneous_positive_slots": simultaneous,
        "purchase_total_kwh": purchase_total,
        "cost_total_yuan": cost_total,
        "charge_total_kwh": charge_total,
        "discharge_total_kwh": discharge_total,
        "four_hour_charge_total_kwh": storage_charge_total,
        "four_hour_discharge_total_kwh": storage_discharge_total,
        "summary_purchase_difference_kwh": abs(purchase_total - float(summary["optimized"]["purchase_kwh"])),
        "summary_cost_difference_yuan": abs(cost_total - float(summary["optimized"]["cost_yuan"])),
        "storage_charge_difference_kwh": abs(charge_total - storage_charge_total),
        "storage_discharge_difference_kwh": abs(discharge_total - storage_discharge_total),
        "natural_day_label_match": interval_labels == expected_labels,
    }
    checks = {
        "row_count": len(rows) == 144 and len(source_rows) == 144,
        "source_values": metrics["max_source_value_difference"] <= 1e-10,
        "natural_day_labels": metrics["natural_day_label_match"],
        "balance": metrics["max_balance_residual_kwh"] <= TOL,
        "state": metrics["max_state_residual_kwh"] <= TOL,
        "soc_bounds": metrics["soc_min_kwh"] >= SOC_MIN - TOL and metrics["soc_max_kwh"] <= SOC_MAX + TOL,
        "soc_cycle": abs(metrics["soc_open_kwh"] - 6000.0) <= TOL and abs(metrics["soc_close_kwh"] - 6000.0) <= TOL,
        "flow_power": metrics["charge_max_kwh"] <= FLOW_MAX + TOL and metrics["discharge_max_kwh"] <= FLOW_MAX + TOL,
        "mutual_exclusion": simultaneous == 0,
        "summary_reconciliation": metrics["summary_purchase_difference_kwh"] <= TOL and metrics["summary_cost_difference_yuan"] <= TOL,
        "four_hour_reconciliation": metrics["storage_charge_difference_kwh"] <= TOL and metrics["storage_discharge_difference_kwh"] <= TOL,
    }
    payload = {
        "schema_version": 1,
        "question": "Q1",
        "method": "Independent CSV/source recomputation; does not import common_data.py or q1_solve.py.",
        "status": "PASS" if all(checks.values()) else "FAIL",
        "tolerance": TOL,
        "metrics": metrics,
        "checks": checks,
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": payload["status"], "metrics": metrics}, ensure_ascii=False))


if __name__ == "__main__":
    main()
