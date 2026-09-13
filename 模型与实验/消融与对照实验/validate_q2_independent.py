from __future__ import annotations

import json
from datetime import date, datetime, time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
ATTACHMENT2 = ROOT / "source_materials" / "C题" / "附件" / "附件2.xlsx"
DETAIL = ROOT / "results" / "q2_detail.csv"
DAILY = ROOT / "results" / "q2_daily.csv"
STORAGE = ROOT / "results" / "q2_storage_4h.csv"
EVENTS = ROOT / "results" / "q2_emergency_events.csv"
SUMMARY = ROOT / "results" / "q2_summary.json"
OUTPUT = ROOT / "validation" / "q2_independent_recompute.json"
TOL = 1e-6
POS_TOL = 1e-7


def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value)).date()


def load_source() -> tuple[list[date], np.ndarray, np.ndarray]:
    wb = load_workbook(ATTACHMENT2, read_only=True, data_only=True)
    try:
        matrices = []
        dates = None
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            current_dates = [parse_date(row[0]) for row in rows]
            if dates is None:
                dates = current_dates
            elif dates != current_dates:
                raise ValueError("附件2两张表日期轴不一致")
            matrices.append(np.asarray([[float(value) for value in row[1:]] for row in rows], dtype=float))
    finally:
        wb.close()
    assert dates is not None
    return dates, matrices[0], matrices[1]


def fmt_minute(total: int) -> str:
    if total == 1440:
        return "24:00"
    return f"{total // 60:02d}:{total % 60:02d}"


def recompute_events(detail: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for day, group in detail.groupby("date", sort=True):
        emergency = group.sort_values("slot")["emergency_purchase_kwh"].to_numpy(dtype=float)
        start = None
        event_id = 0
        for slot, amount in enumerate(emergency):
            if amount > POS_TOL and start is None:
                start = slot
            if start is not None and (amount <= POS_TOL or slot == 143):
                end = slot if amount <= POS_TOL else slot + 1
                event_id += 1
                rows.append({
                    "date": day,
                    "event_id": event_id,
                    "interval": f"{fmt_minute(start * 10)}-{fmt_minute(end * 10)}",
                    "purchase_kwh": float(np.sum(emergency[start:end])),
                })
                start = None
    return pd.DataFrame(rows, columns=["date", "event_id", "interval", "purchase_kwh"])


def main() -> None:
    source_dates, source_load, source_pv = load_source()
    detail = pd.read_csv(DETAIL, encoding="utf-8-sig")
    daily = pd.read_csv(DAILY, encoding="utf-8-sig")
    storage = pd.read_csv(STORAGE, encoding="utf-8-sig")
    events = pd.read_csv(EVENTS, encoding="utf-8-sig")
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))

    ordered = detail.sort_values(["date", "slot"]).reset_index(drop=True)
    expected_load = source_load[31:].reshape(-1)
    expected_pv = source_pv[31:].reshape(-1)
    load_difference = float(np.max(np.abs(ordered["actual_load_kw"].to_numpy() - expected_load)))
    pv_difference = float(np.max(np.abs(ordered["actual_pv_kw"].to_numpy() - expected_pv)))
    load_energy_difference = float(np.max(np.abs(ordered["load_kwh"].to_numpy() - expected_load / 6.0)))
    pv_energy_difference = float(np.max(np.abs(ordered["pv_kwh"].to_numpy() - expected_pv / 6.0)))

    balance = (
        ordered["normal_purchase_kwh"]
        + ordered["pv_kwh"]
        + ordered["discharge_bus_kwh"]
        + ordered["emergency_purchase_kwh"]
        - ordered["load_kwh"]
        - ordered["charge_bus_kwh"]
        - ordered["surplus_discard_kwh"]
    )
    state = (
        ordered["soc_close_kwh"]
        - ordered["soc_open_kwh"]
        - 0.9 * ordered["charge_bus_kwh"]
        + ordered["discharge_bus_kwh"] / 0.9
    )
    daily_recomputed = ordered.groupby("date", sort=True).agg(
        normal_purchase_kwh=("normal_purchase_kwh", "sum"),
        emergency_purchase_kwh=("emergency_purchase_kwh", "sum"),
        normal_cost_yuan=("normal_cost_yuan", "sum"),
        emergency_cost_yuan=("emergency_cost_yuan", "sum"),
        charge_kwh=("charge_bus_kwh", "sum"),
        discharge_kwh=("discharge_bus_kwh", "sum"),
        surplus_discard_kwh=("surplus_discard_kwh", "sum"),
        soc_open_kwh=("soc_open_kwh", "first"),
        soc_close_kwh=("soc_close_kwh", "last"),
    ).reset_index()
    daily_recomputed["total_cost_yuan"] = daily_recomputed["normal_cost_yuan"] + daily_recomputed["emergency_cost_yuan"]
    daily_sorted = daily.sort_values("date").reset_index(drop=True)
    daily_columns = [
        "normal_purchase_kwh", "emergency_purchase_kwh", "normal_cost_yuan", "emergency_cost_yuan",
        "total_cost_yuan", "charge_kwh", "discharge_kwh", "surplus_discard_kwh", "soc_open_kwh", "soc_close_kwh",
    ]
    max_daily_difference = float(max(
        np.max(np.abs(daily_recomputed[column].to_numpy() - daily_sorted[column].to_numpy()))
        for column in daily_columns
    ))

    block_ids = (ordered["slot"].to_numpy(dtype=int) - 1) // 24 + 1
    temp = ordered.assign(block=block_ids)
    storage_recomputed = temp.groupby(["date", "block"], sort=True).agg(
        charge_kwh=("charge_bus_kwh", "sum"),
        discharge_kwh=("discharge_bus_kwh", "sum"),
    ).reset_index()
    storage_sorted = storage.sort_values(["date", "block"]).reset_index(drop=True)
    max_storage_difference = float(max(
        np.max(np.abs(storage_recomputed[column].to_numpy() - storage_sorted[column].to_numpy()))
        for column in ("charge_kwh", "discharge_kwh")
    ))

    events_recomputed = recompute_events(ordered)
    events_sorted = events[["date", "event_id", "interval", "purchase_kwh"]].sort_values(["date", "event_id"]).reset_index(drop=True)
    events_exact_keys = events_recomputed[["date", "event_id", "interval"]].equals(events_sorted[["date", "event_id", "interval"]])
    max_event_amount_difference = (
        float(np.max(np.abs(events_recomputed["purchase_kwh"].to_numpy() - events_sorted["purchase_kwh"].to_numpy())))
        if len(events_sorted) else 0.0
    )

    dates_from_detail = sorted(ordered["date"].unique().tolist())
    expected_dates = [value.isoformat() for value in source_dates[31:]]
    crossday_gap = float(np.max(np.abs(
        daily_sorted["soc_open_kwh"].to_numpy()[1:] - daily_sorted["soc_close_kwh"].to_numpy()[:-1]
    )))
    actual_simultaneous = int(np.sum(
        (ordered["charge_bus_kwh"].to_numpy() > POS_TOL)
        & (ordered["discharge_bus_kwh"].to_numpy() > POS_TOL)
    ))
    emergency_charge_overlap = int(np.sum(
        (ordered["emergency_purchase_kwh"].to_numpy() > POS_TOL)
        & (ordered["charge_bus_kwh"].to_numpy() > POS_TOL)
    ))
    information_violations = int(np.sum(ordered["train_end_date"] >= ordered["date"]))
    reported = summary["strategies"]["seasonal_q80_main"]
    summary_differences = {
        "normal_purchase_kwh": abs(float(daily_recomputed["normal_purchase_kwh"].sum()) - reported["normal_purchase_kwh"]),
        "emergency_purchase_kwh": abs(float(daily_recomputed["emergency_purchase_kwh"].sum()) - reported["emergency_purchase_kwh"]),
        "total_cost_yuan": abs(float(daily_recomputed["total_cost_yuan"].sum()) - reported["total_cost_yuan"]),
        "final_soc_kwh": abs(float(daily_recomputed.iloc[-1]["soc_close_kwh"]) - reported["soc_final_kwh"]),
    }

    metrics = {
        "detail_rows": int(len(ordered)),
        "daily_rows": int(len(daily_sorted)),
        "storage_rows": int(len(storage_sorted)),
        "event_rows": int(len(events_sorted)),
        "date_axis_exact": dates_from_detail == expected_dates,
        "max_source_load_difference_kw": load_difference,
        "max_source_pv_difference_kw": pv_difference,
        "max_load_energy_difference_kwh": load_energy_difference,
        "max_pv_energy_difference_kwh": pv_energy_difference,
        "max_balance_residual_kwh": float(np.max(np.abs(balance))),
        "max_state_residual_kwh": float(np.max(np.abs(state))),
        "soc_min_kwh": float(min(ordered["soc_open_kwh"].min(), ordered["soc_close_kwh"].min())),
        "soc_max_kwh": float(max(ordered["soc_open_kwh"].max(), ordered["soc_close_kwh"].max())),
        "max_crossday_soc_gap_kwh": crossday_gap,
        "max_daily_recompute_difference": max_daily_difference,
        "max_storage_recompute_difference_kwh": max_storage_difference,
        "event_keys_exact": bool(events_exact_keys),
        "max_event_amount_difference_kwh": max_event_amount_difference,
        "simultaneous_actual_slots": actual_simultaneous,
        "emergency_charge_overlap_slots": emergency_charge_overlap,
        "information_cutoff_violations": information_violations,
        "summary_differences": summary_differences,
    }
    pass_flags = {
        "row_counts": len(ordered) == 334 * 144 and len(daily_sorted) == 334 and len(storage_sorted) == 334 * 6,
        "date_axis": metrics["date_axis_exact"],
        "source_values": max(load_difference, pv_difference, load_energy_difference, pv_energy_difference) <= 1e-9,
        "balance": metrics["max_balance_residual_kwh"] <= TOL,
        "state": metrics["max_state_residual_kwh"] <= TOL,
        "soc_bounds": metrics["soc_min_kwh"] >= 1200 - TOL and metrics["soc_max_kwh"] <= 10800 + TOL,
        "crossday_soc": crossday_gap <= TOL,
        "daily_recompute": max_daily_difference <= TOL,
        "storage_recompute": max_storage_difference <= TOL,
        "events": events_exact_keys and max_event_amount_difference <= TOL,
        "mutual_exclusion": actual_simultaneous == 0,
        "emergency_not_charging": emergency_charge_overlap == 0,
        "information_cutoff": information_violations == 0,
        "summary": max(summary_differences.values()) <= TOL,
    }
    payload = {
        "schema_version": 1,
        "question": "Q2",
        "status": "PASS" if all(pass_flags.values()) else "FAIL",
        "pass_flags": pass_flags,
        "metrics": metrics,
        "independence": "Reads attachment2 and frozen CSV/JSON outputs directly; does not import common_data, dispatch_core, or q2_solve.",
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": payload["status"], "metrics": metrics}, ensure_ascii=False))
    if payload["status"] != "PASS":
        raise RuntimeError("Q2独立复算失败")


if __name__ == "__main__":
    main()
